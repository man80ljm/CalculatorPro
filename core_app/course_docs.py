"""课程基本信息表、毕业要求对应关系表。

从桌面设置窗口抽出的纯文档逻辑，网页与桌面可以共用，且不依赖 PyQt。
"""
import os

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from utils import get_outputs_dir


def _set_cell_border(cell, size=4):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        edge_tag = qn(f"w:{edge}")
        element = borders.find(edge_tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:color"), "000000")


def _write_run(cell, text, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(text or ""))
    run.font.name = "仿宋"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "仿宋")
    run.font.size = Pt(12)
    run.bold = bold
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def export_course_basic_docx(data: dict, output_dir: str | None = None) -> str:
    output_dir = output_dir or get_outputs_dir()
    os.makedirs(output_dir, exist_ok=True)
    output_path = os.path.join(output_dir, "1.课程基本信息表.docx")

    doc = Document()
    table = doc.add_table(rows=4, cols=6)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_width = Cm(14.64 / 6)

    rows = [
        ("课程名称", data.get("course_name", ""), "学分", data.get("credits", ""), "学时", data.get("hours", "")),
        ("课程性质", data.get("course_type", ""), "课程代码", data.get("course_code", ""), "学年学期", data.get("school_year_term", "")),
        ("开课学院", data.get("college", ""), "任课教师", data.get("teacher", ""), "上课专业", data.get("major", "")),
        ("上课班级", data.get("class_name", ""), "上课人数", data.get("student_count", ""), "考核人数", data.get("exam_count", "")),
    ]

    for r_idx, row_vals in enumerate(rows):
        row = table.rows[r_idx]
        row.height_rule = WD_ROW_HEIGHT_RULE.AUTO
        for c_idx in range(6):
            cell = row.cells[c_idx]
            cell.width = col_width
            _write_run(cell, row_vals[c_idx], bold=(c_idx in (0, 2, 4)))
            _set_cell_border(cell, size=4)

    doc.save(output_path)
    return output_path


def export_grad_req_docx(grad_req_map, output_dir: str | None = None) -> str:
    output_dir = output_dir or get_outputs_dir()
    os.makedirs(output_dir, exist_ok=True)
    output_path = os.path.join(output_dir, "3.课程目标与毕业要求的对应关系表.docx")

    doc = Document()
    rows_count = max(1, len(grad_req_map or [])) + 1
    table = doc.add_table(rows=rows_count, cols=3)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    total_cm = 14.64
    col_widths = [2.79, 3.37, total_cm - 2.79 - 3.37]
    headers = ["课程目标", "支撑的毕业要求", "支撑的毕业要求指标点"]

    for c_idx, title in enumerate(headers):
        cell = table.cell(0, c_idx)
        cell.width = Cm(col_widths[c_idx])
        _write_run(cell, title, bold=True)
        _set_cell_border(cell, size=4)

    for r in range(1, rows_count):
        obj_name = f"课程目标{r}"
        requirement = ""
        indicator = ""
        if grad_req_map and r - 1 < len(grad_req_map):
            row = grad_req_map[r - 1] or {}
            obj_name = row.get("objective", obj_name)
            requirement = row.get("requirement", "")
            indicator = row.get("indicator", "")
        for c_idx, val in enumerate([obj_name, requirement, indicator]):
            cell = table.cell(r, c_idx)
            cell.width = Cm(col_widths[c_idx])
            _write_run(cell, val, bold=(c_idx == 0))
            _set_cell_border(cell, size=4)
        table.rows[r].height = Cm(1)
        table.rows[r].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY

    table.rows[0].height = Cm(1)
    table.rows[0].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    doc.save(output_path)
    return output_path
