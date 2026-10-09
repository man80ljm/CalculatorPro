"""5/10/20 人的本地隔离压测。不会读取 .env，不接受线上地址，不调用真实 AI。"""
from __future__ import annotations
import argparse
from concurrent.futures import ThreadPoolExecutor
from contextlib import contextmanager
import io
import json
import math
import os
import re
from pathlib import Path
import secrets
import socket
import sqlite3
import subprocess
import sys
import threading
import time
import zipfile

import requests
from docx import Document
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]


def settings(mode="forward"):
    return {
        "mode": mode,
        "course_open_info": {"course_name": "并发合成课程", "department": "合成学院"},
        "course_basic_info": {"course_name": "并发合成课程", "course_code": "SYNTH001",
                              "course_type": "选修", "credits": "3", "hours": "48", "college": "合成学院"},
        "ratios": {"usual": 0.3, "midterm": 0, "final": 0.7},
        "relation_payload": {"objectives_count": 2, "links": [
            {"name": "平时考核", "ratio": 0.3, "methods": [
                {"name": "平时作业", "supports": {"课程目标1": 0.5, "课程目标2": 0.5}, "subtotal": 1.0}]},
            {"name": "期末考核", "ratio": 0.7, "methods": [
                {"name": "期末考试", "supports": {"课程目标1": 0.4, "课程目标2": 0.6}, "subtotal": 1.0}]}]},
        "grad_req_map": [{"requirement": "合成要求1", "indicator": "指标1"},
                         {"requirement": "合成要求2", "indicator": "指标2"}],
        "course_description": "本地并发测试使用的合成课程。",
        "objective_requirements": ["合成目标1", "合成目标2"],
        "spread_mode": "中跨度（7-13分）", "distribution": "标准正态",
        "noise_config": None, "report_style": "专业", "word_limit": 40,
    }


def grade_register(user_number, count):
    book = Workbook()
    sheet = book.active
    sheet.append(["合成学院2026-2027学年第1学期课程成绩登记表"])
    sheet.append(["课程名称：并发合成课程 课程代码：SYNTH001 课程性质：专业任选课"])
    sheet.append([f"开课学院：合成学院 任课教师：压测老师{user_number} 学分：3 考核方式：考查"])
    sheet.append(["班级", "学号", "姓名", "平时(30%)", "期中(0%)", "期末(70%)", "总评", "备注"])
    for student in range(count):
        usual = 60 + user_number + student % 5
        final = 55 + user_number + student % 7
        sheet.append(["合成班甲" if student % 2 else "合成班乙", f"T{user_number:03}{student:04}",
                      f"压测{user_number}学员{student}", usual, 0, final, usual * 0.3 + final * 0.7, ""])
    sheet.append([f"实考 {count}人 总人数 {count}人"])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


@contextmanager
def isolated_server(directory, ai_delay):
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        port = sock.getsockname()[1]
    environment = dict(os.environ)
    environment["PYTHONUTF8"] = "1"
    environment["NO_PROXY"] = "127.0.0.1,localhost"
    log = (directory / "server.log").open("w", encoding="utf-8")
    process = subprocess.Popen([sys.executable, str(ROOT / "scripts/local_load_server.py"),
                                "--directory", str(directory), "--port", str(port),
                                "--ai-delay", str(ai_delay)], cwd=ROOT, env=environment,
                               stdout=log, stderr=subprocess.STDOUT,
                               creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    base = f"http://127.0.0.1:{port}"
    probe = requests.Session()
    probe.trust_env = False
    try:
        deadline = time.monotonic() + 40
        while time.monotonic() < deadline:
            if process.poll() is not None:
                raise RuntimeError(f"隔离服务启动失败，查看 {directory / 'server.log'}")
            try:
                if probe.get(base + "/healthz", timeout=1).status_code == 200:
                    break
            except requests.RequestException:
                pass
            time.sleep(0.2)
        else:
            raise RuntimeError("隔离服务启动超时")
        yield base
    finally:
        probe.close()
        process.terminate()
        try:
            process.wait(timeout=10)
        except subprocess.TimeoutExpired:
            process.kill()
            process.wait(timeout=5)
        log.close()


class Client:
    def __init__(self, base, number, measurements, mode):
        self.base, self.number, self.measurements = base, number, measurements
        self.mode = ("reverse" if number % 2 == 0 else "forward") if mode == "mixed" else mode
        self.session = requests.Session()
        self.session.trust_env = False
        self.course_id = self.term_id = self.grade_id = None
        self.job_id = None
        self.saw_queue = False
        self.reference = None
        self.grade_bytes = b""

    def request(self, name, method, path, *, expected=200, record=True, **kwargs):
        match = re.match(r"^/api/courses/(\d+)(?:/|$)", path)
        if match and method.upper() not in {"GET", "HEAD"}:
            current = self.session.get(self.base + f"/api/courses/{match[1]}", timeout=30)
            if current.status_code == 200:
                kwargs["headers"] = {**kwargs.get("headers", {}), "X-Course-Revision": current.json()["edit_revision"]}
        started = time.perf_counter()
        try:
            response = self.session.request(method, self.base + path, timeout=120, **kwargs)
            if response.status_code == 409 and response.json().get("code") == "edit_conflict" and match:
                # 压测同账号同时发起多个操作时，使用当前版本重新提交相同的合成输入。
                current = self.session.get(self.base + f"/api/courses/{match[1]}", timeout=30)
                kwargs["headers"]["X-Course-Revision"] = current.json()["edit_revision"]
                response = self.session.request(method, self.base + path, timeout=120, **kwargs)
        except requests.RequestException as exc:
            if record:
                self.measurements.append({"name": name, "user": self.number,
                                          "seconds": time.perf_counter() - started, "status": 0, "ok": False})
            raise AssertionError(f"{name}: 请求失败 {exc.__class__.__name__}") from exc
        if record:
            self.measurements.append({"name": name, "user": self.number, "seconds": time.perf_counter() - started,
                                      "status": response.status_code, "ok": response.status_code == expected})
        if response.status_code != expected:
            raise AssertionError(f"{name}: HTTP {response.status_code}: {response.text[:350]}")
        return response

    def prepare_login(self):
        account = {"username": f"load_teacher_{self.number}", "password": secrets.token_urlsafe(24)}
        for route in ("register", "login"):
            self.request(route, "POST", f"/api/{route}", record=False, json=account)

    def import_course(self, barrier, students):
        barrier.wait(timeout=60)
        course = self.request("建立课程", "POST", "/api/courses", json={"name": "并发合成课程", "settings": settings(self.mode)}).json()
        self.course_id = course["id"]
        imported = self.request("导入成绩", "POST", f"/api/courses/{self.course_id}/grade-register",
                                data={"mode": self.mode},
                                files={"file": ("合成成绩.xlsx", grade_register(self.number, students),
                                                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}).json()
        assert imported["student_count"] == students and imported["exam_count"] == students
        assert imported["class_name"] == "专业任选"
        assert imported["mode"] == self.mode
        self.term_id = imported["term_id"]
        detail = self.request("读取课程", "GET", f"/api/courses/{self.course_id}").json()
        grade = next(row for row in detail["files"] if row["kind"] == "grade")
        self.grade_id = grade["id"]
        self.grade_bytes = self.request("读取成绩", "GET", f"/api/courses/{self.course_id}/files/{self.grade_id}").content

    def compare_summary(self, summary):
        for key in ("student_count", "average_score", "achievement", "term_id"):
            assert summary[key] == self.reference[key], f"老师{self.number}: {key} 与单人结果不一致"

    def validate_archive(self, content):
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            assert archive.testzip() is None
            names = archive.namelist()
            assert len(names) == len(set(names))
            table_name = next(name for name in names if name.endswith("5.基于考核结果的课程目标达成情况评价结果表.docx"))
            document = Document(io.BytesIO(archive.read(table_name)))
            seen = set()
            for table in document.tables:
                if not table.rows or "分目标达成值" not in "".join(cell.text for cell in table.rows[0].cells):
                    continue
                for row in table.rows[1:]:
                    cells = [cell.text.strip() for cell in row.cells]
                    label = cells[0]
                    if label in self.reference["achievement"]:
                        assert round(float(cells[5]), 3) == round(float(self.reference["achievement"][label]), 3)
                        seen.add(label)
                    elif label == "课程目标达成值":
                        assert round(float(cells[6]), 3) == round(float(self.reference["achievement"]["总达成度"]), 3)
            assert len(seen) == 2, "未找到本人的表5目标值"
            source = [name for name in names if name.endswith(".xlsx") and archive.read(name) == self.grade_bytes]
            assert source, "压缩包没有包含本人的原始成绩"

    def workflow(self, barrier, rounds, round_pause):
        barrier.wait(timeout=60)
        for round_number in range(rounds):
            courses = self.request("课程列表", "GET", "/api/courses").json()["courses"]
            assert [row["id"] for row in courses] == [self.course_id]
            summary = self.request("计算", "POST", f"/api/courses/{self.course_id}/calculate").json()
            self.compare_summary(summary)
            archive = self.request("统计表导出", "POST", f"/api/courses/{self.course_id}/export")
            self.validate_archive(archive.content)
            started = time.perf_counter()
            job = self.request("提交报告", "POST", f"/api/courses/{self.course_id}/report-jobs").json()
            self.job_id = job["job_id"]
            deadline = time.monotonic() + 180
            while time.monotonic() < deadline:
                self.saw_queue |= job["stage"] == "queued"
                if job["error"]:
                    raise AssertionError(f"报告失败: {job['error']}")
                if job["done"]:
                    break
                time.sleep(0.3)
                job = self.request("报告进度", "GET", f"/api/report-jobs/{self.job_id}").json()
            else:
                raise AssertionError("报告180秒未完成")
            self.measurements.append({"name": "报告完成（含排队）", "user": self.number,
                                      "seconds": time.perf_counter() - started, "status": 200, "ok": True})
            self.compare_summary(job["summary"])
            assert job["percent"] == 100
            blob = self.request("报告下载", "GET", f"/api/report-jobs/{self.job_id}/download").content
            self.validate_archive(blob)
            saved = self.request("保存的成绩", "GET", f"/api/courses/{self.course_id}/files/{self.grade_id}").content
            assert saved == self.grade_bytes, "成绩文件发生变化"
            if round_number + 1 < rounds:
                time.sleep(round_pause)


def parallel(clients, action):
    barrier = threading.Barrier(len(clients))
    errors = []
    with ThreadPoolExecutor(max_workers=len(clients)) as executor:
        tasks = [(client.number, executor.submit(action, client, barrier)) for client in clients]
        for number, task in tasks:
            try:
                task.result()
            except Exception as exc:
                errors.append({"user": number, "error": str(exc)[:500]})
    return errors


def percentile(values, fraction):
    ordered = sorted(values)
    return round(ordered[max(0, math.ceil(len(ordered) * fraction) - 1)], 3)


def stage(directory, users, students, rounds, ai_delay, mode, round_pause):
    directory.mkdir(parents=True)
    measurements = []
    errors = []
    with isolated_server(directory, ai_delay) as base:
        clients = [Client(base, number, measurements, mode) for number in range(1, users + 1)]
        for client in clients:
            client.prepare_login()
        started = time.perf_counter()
        errors += parallel(clients, lambda client, barrier: client.import_course(barrier, students))
        import_seconds = time.perf_counter() - started
        for client in clients:
            if client.course_id and client.grade_id:
                try:
                    client.reference = client.request("单人基准", "POST", f"/api/courses/{client.course_id}/calculate", record=False).json()
                except Exception as exc:
                    errors.append({"user": client.number, "error": str(exc)[:500]})
        started = time.perf_counter()
        if not errors:
            errors += parallel(clients, lambda client, barrier: client.workflow(barrier, rounds, round_pause))
        elapsed = time.perf_counter() - started
        # 所有压测老师彼此交叉访问课程、成绩、报告：必须全部404。
        privacy_checks = 0
        for index, client in enumerate(clients):
            other = clients[(index + 1) % len(clients)]
            if other is client or not other.course_id:
                continue
            paths = [f"/api/courses/{other.course_id}"]
            if other.grade_id:
                paths.append(f"/api/courses/{other.course_id}/files/{other.grade_id}")
            if other.job_id:
                paths.extend([f"/api/report-jobs/{other.job_id}", f"/api/report-jobs/{other.job_id}/download"])
            for path in paths:
                try:
                    client.request("账号隔离", "GET", path, expected=404)
                    privacy_checks += 1
                except Exception as exc:
                    errors.append({"user": client.number, "error": str(exc)[:500]})
        metrics = clients[0].request("模拟AI计数", "GET", "/__load_metrics", record=False).json()
        if metrics["external_http_attempts"] or metrics["calls"] != users * rounds:
            errors.append({"error": "模拟AI调用次数或外网隔离检查异常", "metrics": metrics})
        # done 字段先更新、事件随后写库，留一点时间验证终态持久化。
        time.sleep(0.2)
        with sqlite3.connect(directory / "app.db") as db:
            events = db.execute("SELECT status, COUNT(*) FROM report_job_events GROUP BY status").fetchall()
        if dict(events).get("success", 0) != users * rounds:
            errors.append({"error": "报告成功事件未完整保存", "events": events})
        for client in clients:
            client.session.close()
    summary = {}
    for name in sorted({row["name"] for row in measurements}):
        group = [row for row in measurements if row["name"] == name]
        values = [row["seconds"] for row in group]
        summary[name] = {"count": len(group), "failed": sum(not row["ok"] for row in group),
                         "p50_seconds": percentile(values, 0.5), "p95_seconds": percentile(values, 0.95),
                         "max_seconds": round(max(values), 3)}
    return {"users": users, "students_per_course": students, "rounds": rounds, "mode": mode, "round_pause_seconds": round_pause,
            "ai_mode": "模拟响应", "ai_delay_seconds": ai_delay, "database": "独立SQLite",
            "import_seconds": round(import_seconds, 3), "workflow_seconds": round(elapsed, 3),
            "queued_users": sum(client.saw_queue for client in clients),
            "privacy_checks": privacy_checks, "ai_metrics": metrics, "errors": errors,
            "summary": summary, "requests": measurements}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--users", type=int, nargs="+", default=[5, 10, 20])
    parser.add_argument("--students", type=int, default=35)
    parser.add_argument("--rounds", type=int, default=1)
    parser.add_argument("--ai-delay", type=float, default=3)
    parser.add_argument("--mode", choices=["forward", "reverse", "mixed"], default="forward")
    parser.add_argument("--round-pause", type=float, default=60,
                        help="重复轮次间隔秒数，默认60秒，避免触发同一IP的报告限流")
    args = parser.parse_args()
    if not all(1 <= value <= 20 for value in args.users) or not 2 <= args.students <= 500 or not 1 <= args.rounds <= 100 or not 0 <= args.ai_delay <= 30 or not 0 <= args.round_pause <= 60:
        parser.error("人数1–20，学生数2–500，轮数1–100，AI模拟等待0–30秒，轮间隔0–60秒")
    run = ROOT / ".local/load-test" / (time.strftime("%Y%m%d-%H%M%S") + "-" + secrets.token_hex(3))
    run.mkdir(parents=True)
    results = []
    for index, users in enumerate(args.users):
        print(f"开始 {users} 人压测（AI模拟、独立库）", flush=True)
        result = stage(run / f"stage-{index + 1}-{users}users", users, args.students, args.rounds, args.ai_delay, args.mode, args.round_pause)
        results.append(result)
        (run / "results.json").write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"{users}人：{len(result['errors'])}项错误，操作耗时{result['workflow_seconds']}秒，排队{result['queued_users']}人", flush=True)
        if result["errors"]:
            print(json.dumps(result["errors"], ensure_ascii=False), flush=True)
    print(f"结果：{run / 'results.json'}", flush=True)
    return 1 if any(result["errors"] for result in results) else 0


if __name__ == "__main__":
    raise SystemExit(main())
