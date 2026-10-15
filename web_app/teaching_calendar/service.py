"""无持久化文件的预览、输入检查和教学日历生成。"""
import copy
import io
import json
import math
import re
import zipfile
from pathlib import Path
from tempfile import TemporaryDirectory

import docx
import openpyxl
import xlrd

from web_app.resource_limits import validate_document
from . import engine
from .conversion import convert_file
from .errors import CalendarError
from .academic_term import TERM_CELL, current_academic_term

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_PAYLOAD_BYTES = 1024 * 1024
MAX_LESSONS = 1000
OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def validate_upload(filename, data, *, outline=False):
    suffix = Path(filename).suffix.lower()
    allowed = {".docx"} if outline else {".xls", ".xlsx"}
    if suffix not in allowed:
        raise CalendarError("教学大纲只接受 .docx。" if outline else "教学日历模板只接受 .xls 或 .xlsx。",
                            code="invalid_file_type")
    if not data:
        raise CalendarError("上传文件是空的。", code="empty_file")
    if len(data) > MAX_FILE_BYTES:
        raise CalendarError("每个上传文件不能超过 5MB。", 413, code="file_too_large")
    validate_document(data, filename)
    if suffix == ".xls":
        if not data.startswith(OLE_MAGIC):
            raise CalendarError("模板内容不是有效的 .xls，请重新另存后上传。", code="invalid_file")
        try:
            workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
            try:
                if not workbook.nsheets:
                    raise ValueError("empty workbook")
                if any(workbook.sheet_by_index(i).nrows > 5000 or workbook.sheet_by_index(i).ncols > 200
                       for i in range(workbook.nsheets)):
                    raise CalendarError("模板超过 5000 行或 200 列，请删除多余空白行列。", 413)
            finally:
                workbook.release_resources()
        except CalendarError:
            raise
        except Exception:
            raise CalendarError("无法读取 .xls 模板，请检查是否损坏或加密。", code="invalid_file") from None
    else:
        required = "word/document.xml" if outline else "xl/workbook.xml"
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if required not in archive.namelist():
                    raise CalendarError("文件内容与扩展名不符，请选择有效的 Word 大纲或 Excel 模板。",
                                        code="invalid_file")
        except zipfile.BadZipFile:
            raise CalendarError("无法读取上传文件，请重新另存后上传。", code="invalid_file") from None


def _text(value, label):
    if value is None:
        return ""
    if isinstance(value, bool) or not isinstance(value, (str, int, float)):
        raise CalendarError(f"{label}格式不正确。", code="invalid_preview")
    value = str(value)
    if len(value) > 10000 or re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", value):
        raise CalendarError(f"{label}过长或包含无法写入 Excel 的字符。", code="invalid_preview")
    return value


def validate_rows(rows):
    if not isinstance(rows, list) or not rows:
        raise CalendarError("没有有效课次，请检查大纲或预览表。", code="empty_lessons")
    if len(rows) > MAX_LESSONS:
        raise CalendarError(f"最多支持 {MAX_LESSONS} 次课。", 413, code="too_many_lessons")
    cleaned = []
    seen = set()
    for index, row in enumerate(rows, 1):
        if not isinstance(row, dict):
            raise CalendarError(f"第 {index} 行格式不正确。", code="invalid_preview")
        item = {key: _text(row.get(key, ""), f"第 {index} 行{key}") for key in engine.COL}
        lesson = item["课次"].strip()
        if not lesson.isascii() or not lesson.isdigit() or not 1 <= int(lesson) <= 99999 or int(lesson) in seen:
            raise CalendarError(f"第 {index} 行课次应为不重复的正整数。", code="invalid_lesson_number")
        item["课次"] = int(lesson)
        seen.add(item["课次"])
        item["学时"] = item["学时"].strip()
        try:
            engine.lesson_hours(item["学时"])
        except CalendarError:
            raise CalendarError(f"第 {index} 行学时应填写有限的非负数字。", code="invalid_hours")
        cleaned.append(item)
    if not math.isfinite(sum(float(r["学时"] or 0) for r in cleaned)):
        raise CalendarError("学时合计超出可处理范围。", code="invalid_hours")
    return cleaned


def validate_payload(raw):
    if not isinstance(raw, str) or len(raw.encode("utf-8")) > MAX_PAYLOAD_BYTES:
        raise CalendarError("预览内容不能超过 1MB。", 413, code="preview_too_large")
    try:
        payload = json.loads(raw)
    except ValueError:
        raise CalendarError("预览内容不是有效的 JSON。", code="invalid_preview") from None
    if not isinstance(payload, dict) or not isinstance(payload.get("教材"), dict):
        raise CalendarError("请提交完整的教材信息和课次表。", code="invalid_preview")
    book = {key: _text(payload["教材"].get(key, ""), key) for key in engine.BOOK_CELLS}
    return validate_rows(payload.get("lessons")), book


def validate_template(ws):
    if ws.max_column < 21 or ws.max_row < engine.FIRST_ROW:
        raise CalendarError("模板结构不匹配：需要 A 至 U 列，课次从第 8 行开始。", code="invalid_template")
    for col, name in ((1, "课次"), (2, "学时")):
        if not any(name in str(ws.cell(r, col).value or "") for r in range(1, engine.FIRST_ROW)):
            raise CalendarError("模板结构不匹配：A 列应为课次，B 列应为学时。", code="invalid_template")
    current = engine.count_lesson_rows(ws)
    if current > MAX_LESSONS:
        raise CalendarError("模板的课次行数过多。", 413, code="invalid_template")
    last = engine.FIRST_ROW + current - 1
    expected = {f"D{r}:N{r}" for r in range(engine.FIRST_ROW, last + 1)}
    expected.update(f"O{r}:T{r}" for r in range(engine.FIRST_ROW, last + 1))
    merged = {str(m) for m in ws.merged_cells.ranges}
    if not expected.issubset(merged):
        raise CalendarError("模板结构不匹配：每次课的教学内容须合并 D:N，教学要求须合并 O:T。",
                            code="invalid_template")
    for area in ws.merged_cells.ranges:
        if TERM_CELL in area and ws.cell(area.min_row, area.min_col).coordinate != TERM_CELL:
            raise CalendarError("模板第二行不能填写学年学期，请使用教务系统原始模板。", code="invalid_template")
        if area.min_row <= last and area.max_row >= engine.FIRST_ROW and str(area) not in expected:
            raise CalendarError("模板课次区域存在不支持的合并单元格，请使用教务系统原始模板。",
                                code="invalid_template")
        for ref in engine.BOOK_CELLS.values():
            if ref in area and ref != ws.cell(area.min_row, area.min_col).coordinate:
                raise CalendarError("模板教材填写位置与 C5、K5、Q5、T5 不匹配。", code="invalid_template")


def _open_template(filename, data, directory):
    if Path(filename).suffix.lower() == ".xls":
        source = directory / "template.xls"
        source.write_bytes(data)
        data = convert_file(source, "xlsx", directory).read_bytes()
        validate_document(data, "template.xlsx")
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data))
    except Exception:
        raise CalendarError("无法打开教学日历模板，请检查是否损坏或加密。", code="invalid_template") from None
    try:
        validate_template(workbook.active)
        # 表头总学时有公式时使用原模板缓存值，但保留其他公式和样式。
        total = workbook.active[engine.TOTAL_HOURS_CELL]
        if total.data_type == "f":
            cached = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
            try:
                workbook.active._calendar_cached_total = cached.active[engine.TOTAL_HOURS_CELL].value
            finally:
                cached.close()
        return workbook
    except Exception:
        workbook.close()
        raise


def _resize(ws, total):
    """保持原行调整算法，补齐 openpyxl 不自动移动的行高和打印区域元数据。"""
    last = engine.FIRST_ROW + engine.count_lesson_rows(ws) - 1
    difference = total - engine.count_lesson_rows(ws)
    dimensions = {r: copy.copy(d) for r, d in ws.row_dimensions.items() if r > last}
    print_ranges = [copy.copy(area) for area in ws._print_area.ranges]
    if difference:
        for r in dimensions:
            del ws.row_dimensions[r]
    engine.set_lesson_rows(ws, total)
    if difference:
        for r in list(ws.row_dimensions):
            if r > last + difference:
                del ws.row_dimensions[r]
        for r, dimension in dimensions.items():
            dimension.index = r + difference
            ws.row_dimensions[r + difference] = dimension
        for area in print_ranges:
            if area.min_row > last:
                area.shift(row_shift=difference)
            elif area.max_row > last:
                area.max_row += difference
        if print_ranges:
            ws.print_area = [area.coord for area in print_ranges]


def _summary(ws, rows):
    total_cell = ws[engine.TOTAL_HOURS_CELL]
    original = total_cell.value
    try:
        if total_cell.data_type == "f":
            total_cell.value = getattr(ws, "_calendar_cached_total", None)
        summary = engine.report(ws, rows)
    finally:
        total_cell.value = original
    total = summary["template_total_hours"]
    if total is not None and (not math.isfinite(total) or total < 0):
        summary.update(template_total_hours=None, hours_match=None)
    if summary["template_total_hours"] is None:
        summary["warnings"].append("未能读取模板 S4 的总学时，请人工核对学时合计。")
    return summary


def preview(outline_name, outline_bytes, template_name, template_bytes, use_ai=False):
    validate_upload(outline_name, outline_bytes, outline=True)
    validate_upload(template_name, template_bytes)
    with TemporaryDirectory(prefix="calculatorpro-calendar-") as temporary:
        workbook = _open_template(template_name, template_bytes, Path(temporary))
        try:
            try:
                # DOCX 只打开一次，两个读取函数共用解析后的文档。
                document = docx.Document(io.BytesIO(outline_bytes))
            except Exception:
                raise CalendarError("无法读取教学大纲，请检查 Word 文件是否损坏。", code="invalid_outline") from None
            outline = engine.read_outline(document)
            book_text = engine.read_textbook(document)
            rows, book, warnings = engine.build_rows(outline, book_text, use_ai=False)
            rows = validate_rows(rows)
            if use_ai:
                rows, book, warnings = engine.build_rows(outline, book_text, use_ai=True)
                rows = validate_rows(rows)
            _resize(workbook.active, len(rows))
            summary = _summary(workbook.active, rows)
            warnings.extend(summary.pop("warnings"))
            return {"教材": book, "lessons": rows, "textbook_source": book_text,
                    "mode": "ai" if use_ai else "rule", "warnings": warnings,
                    "academic_term": current_academic_term(), **summary}
        finally:
            workbook.close()


def generate(template_name, template_bytes, raw_payload, output_format, *, course_info=None, progress=None):
    if output_format not in {"xls", "xlsx"}:
        raise CalendarError("下载格式只支持 xls 或 xlsx。", code="invalid_format")
    validate_upload(template_name, template_bytes)
    rows, book = validate_payload(raw_payload)
    notify = progress or (lambda percent, message: None)
    notify(15, "正在准备教学日历模板…")
    with TemporaryDirectory(prefix="calculatorpro-calendar-") as temporary:
        root = Path(temporary)
        workbook = _open_template(template_name, template_bytes, root)
        try:
            ws = workbook.active
            notify(35, "正在按大纲调整课次和表格布局…")
            _resize(ws, len(rows))
            notify(50, "正在填写教材、教学内容和学时…")
            academic_term = current_academic_term()
            # 仅更新第二行文字，保留学校模板的 A2:U2 合并和格式。
            ws[TERM_CELL] = academic_term["template_text"]
            if course_info is not None:
                # 固定空白模板的课程信息来自已保存课程，数值总学时用于核对。
                for ref, key in (("C3", "college"), ("C4", "course_name")):
                    if course_info.get(key):
                        ws[ref].value = _text(course_info[key], key)
                        ws[ref].data_type = "s"
                for ref, key in (("G3", "classes"), ("M3", "teacher")):
                    if key in course_info:
                        ws[ref].value = _text(course_info[key], key) or None
                        ws[ref].data_type = "s"
                hours = str(course_info.get("hours") or "").strip()
                ws[engine.TOTAL_HOURS_CELL] = (engine.lesson_hours(hours) if hours else
                                               sum(float(row["学时"] or 0) for row in rows))
                ws[engine.TOTAL_HOURS_CELL].number_format = "General"
            # 原 fill_book 仅写非空值；先清空四格，保证页面删除教材值也能生效。
            for ref in engine.BOOK_CELLS.values():
                ws[ref].value = None
            engine.fill_book(ws, book)
            engine.fill(ws, rows)
            # 仅文字列标记文本，避免 = 开头的教学文本成为公式；学时保持数字类型。
            for index in range(len(rows)):
                for col in engine.COL.values():
                    if col != engine.COL["学时"]:
                        ws.cell(engine.FIRST_ROW + index, col).data_type = "s"
            for ref in engine.BOOK_CELLS.values():
                ws[ref].data_type = "s"
            notify(65, "正在检查学时、合并单元格和签字栏…")
            summary = _summary(ws, rows)
            summary["academic_term"] = academic_term
            if summary["missing_merges"]:
                raise CalendarError("模板合并单元格检查失败，未生成文件。", code="invalid_template")
            output = root / "calendar.xlsx"
            workbook.save(output)
            notify(80, "正在生成 .xls 文件…" if output_format == "xls" else "正在生成文件…")
            if output_format == "xls":
                output = convert_file(output, "xls", root)
            content = output.read_bytes()
            if output_format == "xls" and not content.startswith(OLE_MAGIC):
                raise CalendarError("生成的 .xls 格式检查失败，请重试或联系管理员。", 502,
                                    code="conversion_failed")
            notify(95, "文件已生成，正在准备下载…")
            base = Path(template_name.replace("\\", "/")).stem[:100]
            filename = (f"{base}_{academic_term['year_start']}-{academic_term['year_end']}_"
                        f"第{academic_term['semester']}学期_已填写.{output_format}")
            return filename, content, summary
        finally:
            workbook.close()


def generate_for_course(course_name, settings, output_format="xls", *, progress=None, teaching_assignment=None):
    """无需文件上传或 AI：使用该课程最新大纲和内置学校模板直接下载。"""
    from .source import normalize_calendar_data
    calendar = normalize_calendar_data(settings.get("teaching_calendar"))
    if calendar is None:
        raise CalendarError("这门课程尚未保存大纲课次明细，请通过“从大纲建课”用大纲更新该课程后再下载。",
                            code="calendar_source_missing")
    if calendar.get("error"):
        raise CalendarError(calendar["error"], code="calendar_source_unavailable")
    if progress:
        progress(10, "正在读取已保存的 AI 整理内容…" if calendar.get("mode") == "ai" else
                     "正在读取已保存的课程和大纲信息…")
    basic = settings.get("course_basic_info") or {}
    template = Path(__file__).with_name("templates") / "teaching_calendar.xls"
    try:
        template_bytes = template.read_bytes()
    except OSError:
        raise CalendarError("服务器教学日历模板不可用，请联系管理员。", 503,
                            code="calendar_template_missing") from None
    course_info = {"course_name": basic.get("course_name") or course_name,
                   "college": basic.get("college") or (settings.get("course_open_info") or {}).get("department"),
                   "hours": basic.get("hours")}
    if teaching_assignment is not None:
        course_info.update(teacher=teaching_assignment["teacher"], classes="、".join(teaching_assignment["classes"]))
    safe_name = re.sub(r'[\\/:*?"<>|\r\n]', "_", course_name)[:80] or "课程"
    filename, content, summary = generate(f"{safe_name}_教学日历.xls", template_bytes,
                    json.dumps(calendar, ensure_ascii=False), output_format, course_info=course_info, progress=progress)
    if teaching_assignment is not None:
        summary["teaching_assignment"] = teaching_assignment
        if teaching_assignment["status"] not in {"ready", "manual"}: summary["warnings"].append(teaching_assignment["message"])
    return filename, content, summary
