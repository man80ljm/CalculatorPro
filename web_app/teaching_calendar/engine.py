"""复用「大纲生成教学日历.py」的解析、规则拆分及模板填写算法。

保留原脚本的字段、提示词和行调整顺序；以异常和返回值代替命令行退出/输出。
"""
import copy
import json
import math
import os
import re

import requests
from docx.oxml.ns import qn
from openpyxl.utils import get_column_letter

from core_app.ai_report import deepseek_chat_url, deepseek_model
from web_app.deepseek_pool import PoolError, external_slot
from .errors import CalendarError

FIRST_ROW = 8
COL = {"课次": 1, "学时": 2, "教材章节": 3, "教学内容": 4, "教学要求": 15, "备注": 21}
MERGES = [(4, 14), (15, 20)]
TOTAL_HOURS_CELL = "S4"
BOOK_CELLS = {"教材名称": "C5", "版次": "K5", "编者": "Q5", "出版社": "T5"}
CN_NUM = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5, "六": 6, "七": 7, "八": 8, "九": 9, "十": 10}

SYSTEM_PROMPT = """你是高校教务助理，负责把课程教学大纲整理成教务系统的“教学日历”。
对每一次课，根据给出的“章目名称”和“教学目标及教学内容”，生成两个字段：
- 教学内容：本次课讲授的具体知识点，用“1.…；2.…；3.…”的格式，简洁书面，不超过 80 字。
- 教学要求：学生应达到的要求，以“掌握/理解/了解/能够”开头，用“1.…；2.…”的格式，不超过 60 字。
只依据大纲原文整理，不要编造大纲里没有的内容，不要写课程思政。
另外，把输入中的“教材选用”拆成：教材名称、版次（统一写成“第N版”，阿拉伯数字）、编者、出版社；原文没有的项填空字符串。
严格输出 JSON：
{"教材":{"教材名称":"…","版次":"…","编者":"…","出版社":"…"},
 "lessons":[{"课次":1,"教学内容":"…","教学要求":"…"}, …]}
课次与输入一一对应。"""


def cell_text(tc):
    paras = []
    for p in tc.iter(qn("w:p")):
        t = "".join(x.text or "" for x in p.iter(qn("w:t"))).strip()
        if t:
            paras.append(t)
    return "\n".join(paras)


def read_textbook(document):
    """接受已经打开的 DOCX，与 read_outline 共用同一份文档。"""
    for tbl in document.element.body.iter(qn("w:tbl")):
        for tr in tbl.iter(qn("w:tr")):
            cells = [cell_text(tc) for tc in tr.findall(qn("w:tc"))]
            for i, c in enumerate(cells[:-1]):
                if c.replace(" ", "") in ("教材选用", "教材名称", "教材", "选用教材") and cells[i + 1]:
                    return cells[i + 1].strip()
    return ""


def rule_split_book(text):
    book = {"教材名称": "", "版次": "", "编者": "", "出版社": ""}
    if not text:
        return book
    m = re.search(r"第\s*([0-9一二三四五六七八九十]+)\s*版", text)
    if m:
        n = m.group(1)
        book["版次"] = f"第{CN_NUM.get(n, n)}版"
    parts = [p.strip(" （）()") for p in re.split(r"[，,．.。；;：:]", text) if p.strip(" （）()")]
    rest = []
    for p in parts:
        if "出版社" in p:
            book["出版社"] = p
        elif re.search(r"版$|^\d{4}年?$", p):
            continue
        else:
            rest.append(re.sub(r"[（(]?第\s*[0-9一二三四五六七八九十]+\s*版[）)]?", "", p).strip())
    rest = [p for p in rest if p]
    if rest:
        book["教材名称"] = max(rest, key=len)
        others = [p for p in rest if p != book["教材名称"]]
        if others:
            book["编者"] = min(others, key=len)
    return book


def read_outline(document):
    for tbl in document.element.body.iter(qn("w:tbl")):
        rows = [[cell_text(tc) for tc in tr.findall(qn("w:tc"))] for tr in tbl.iter(qn("w:tr"))]
        if not rows:
            continue
        header = [h.replace("\n", "") for h in rows[0]]
        if "课次" in header and "学时" in header:
            lessons = []
            for r in rows[1:]:
                item = {header[i]: (r[i] if i < len(r) else "") for i in range(len(header))}
                if str(item.get("课次", "")).strip().isdigit():
                    lessons.append(item)
            return lessons
    raise CalendarError("没有在大纲里找到包含「课次」「学时」的教学内容表格。", code="outline_table_missing")


def pick(item, *keys):
    for k in keys:
        for name, v in item.items():
            if k in name:
                return v.strip()
    return ""


def rule_split(text):
    goal = re.search(r"教学目标[:：](.*?)(?=教学内容[:：]|$)", text, re.S)
    cont = re.search(r"教学内容[:：](.*)", text, re.S)
    g = goal.group(1).strip(" ；;。\n") if goal else ""
    c = cont.group(1).strip(" ；;。\n") if cont else text.strip()
    return c, g


def parse_ai_result(data):
    """沿用原脚本的 AI JSON 校验，供独立整理及建课合并响应共用。"""
    if not isinstance(data, dict) or not isinstance(data.get("lessons", []), list):
        raise ValueError("invalid AI result")
    ai = {}
    for item in data.get("lessons", []):
        number = int(item["课次"])
        if number in ai or any(not isinstance(item.get(k, ""), str) for k in ("教学内容", "教学要求")):
            raise ValueError("invalid lesson")
        ai[number] = item
    book = data.get("教材") or {}
    if not isinstance(book, dict):
        raise ValueError("invalid textbook")
    return ai, book


def call_deepseek(lessons_brief, api_key):
    payload = {
        "model": deepseek_model(),
        "messages": [
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": json.dumps(lessons_brief, ensure_ascii=False)},
        ],
        "response_format": {"type": "json_object"},
        "temperature": 0.3,
    }
    try:
        # 只使用本模块读取的 DEEPSEEK_API_KEY，但与其他 AI 功能共用全站并发预算。
        with external_slot(wait_timeout=10):
            resp = requests.post(
                deepseek_chat_url(),
                headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
                json=payload,
                timeout=(10, 120),
            )
            resp.raise_for_status()
            data = json.loads(resp.json()["choices"][0]["message"]["content"])
            return parse_ai_result(data)
    except requests.Timeout:
        raise CalendarError("DeepSeek 整理超时，请重试或改用规则模式。", 504,
                            code="ai_timeout", fallback_available=True) from None
    except PoolError:
        raise CalendarError("AI 服务正忙，请稍后重试或改用规则模式。", 503,
                            code="ai_busy", fallback_available=True) from None
    except (requests.RequestException, ValueError, KeyError, TypeError, IndexError, AttributeError, OverflowError):
        # 不将响应体、请求头、密钥或上游异常原文返回页面。
        raise CalendarError("DeepSeek 调用失败或返回内容无法读取，请重试或改用规则模式。", 502,
                            code="ai_failed", fallback_available=True) from None


def build_rows(outline, book_text, use_ai, *, ai_call=None):
    book = rule_split_book(book_text)
    rows = []
    warnings = []
    for it in outline:
        detail = pick(it, "教学目标", "教学内容")
        content, req = rule_split(detail)
        rows.append({
            "课次": int(pick(it, "课次")),
            "学时": pick(it, "学时"),
            "教材章节": pick(it, "章目", "章节"),
            "教学内容": content,
            "教学要求": req,
            "备注": pick(it, "教学方式"),
            "_原文": detail,
        })
    if use_ai:
        key = os.environ.get("DEEPSEEK_API_KEY", "").strip()
        if ai_call is None and not key:
            raise CalendarError("服务器未配置 DEEPSEEK_API_KEY，请改用规则模式。", 503,
                                code="ai_not_configured", fallback_available=True)
        brief = {
            "教材选用": book_text,
            "lessons": [{"课次": r["课次"], "章目名称": r["教材章节"], "教学目标及教学内容": r["_原文"]} for r in rows],
        }
        ai, ai_book = ai_call(brief) if ai_call is not None else call_deepseek(brief, key)
        missing = []
        for r in rows:
            a = ai.get(r["课次"])
            if a and a.get("教学内容") and a.get("教学要求"):
                r["教学内容"], r["教学要求"] = a["教学内容"].strip(), a["教学要求"].strip()
            else:
                missing.append(r["课次"])
        if missing:
            warnings.append(f"第 {missing} 次课 AI 没有返回完整内容，已按规则拆分填入。")
        for k in book:
            if str(ai_book.get(k, "")).strip():
                book[k] = str(ai_book[k]).strip()
    for r in rows:
        r.pop("_原文", None)
    return rows, book, warnings


def count_lesson_rows(ws):
    # 学校空白模板没有课次编号；用原有的两组合并识别占位行。
    if not str(ws.cell(FIRST_ROW, 1).value or "").strip():
        merged = {str(m) for m in ws.merged_cells.ranges}
        row = FIRST_ROW
        while f"D{row}:N{row}" in merged and f"O{row}:T{row}" in merged:
            row += 1
        if row > FIRST_ROW:
            return row - FIRST_ROW
    r = FIRST_ROW
    while str(ws.cell(r + 1, 1).value or "").strip().isdigit():
        r += 1
    return r - FIRST_ROW + 1


def set_lesson_rows(ws, total):
    """原脚本的行调整算法：拆开全部合并，再移动签字栏和恢复合并。"""
    current = count_lesson_rows(ws)
    last = FIRST_ROW + current - 1
    diff = total - current
    if diff == 0:
        return
    merged = list(ws.merged_cells.ranges)
    for m in merged:
        ws.unmerge_cells(str(m))
    height = ws.row_dimensions[last].height
    if diff > 0:
        ws.insert_rows(last + 1, diff)
    else:
        ws.delete_rows(last + diff + 1, -diff)
    new_last = last + diff
    for m in merged:
        if m.min_row > last:
            m.shift(row_shift=diff)
        elif m.min_row > new_last:
            continue
        ws.merge_cells(str(m))
    for r in range(last + 1, new_last + 1):
        for c in range(1, ws.max_column + 1):
            src = ws.cell(last, c)
            if src.has_style:
                ws.cell(r, c)._style = copy.copy(src._style)
        ws.row_dimensions[r].height = height
        for c1, c2 in MERGES:
            ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)


def fill_book(ws, book):
    for k, ref in BOOK_CELLS.items():
        if book.get(k):
            ws[ref].value = book[k]
    return "教材信息：" + "，".join(f"{k}={book.get(k) or '（空）'}" for k in BOOK_CELLS)


def lesson_hours(value):
    """空学时写空单元格；整数写 int，小数写 float，拒绝非数字和非有限值。"""
    text = "" if value is None else str(value).strip()
    if not text:
        return None
    if not re.fullmatch(r"[+-]?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[eE][+-]?[0-9]+)?", text):
        raise CalendarError("学时应填写有限的非负数字。", code="invalid_hours")
    number = float(text)
    if not math.isfinite(number) or number < 0:
        raise CalendarError("学时应填写有限的非负数字。", code="invalid_hours")
    if "." in text or not number.is_integer():
        return number
    return int(text) if "e" not in text.lower() else int(number)


def fill(ws, rows):
    for i, r in enumerate(rows):
        row = FIRST_ROW + i
        for field, col in COL.items():
            v = r.get(field, "")
            cell = ws.cell(row, col)
            if field == "学时":
                cell.value = lesson_hours(v)
                cell.number_format = "General"
            else:
                cell.value = str(v) if v != "" else None
                cell.number_format = "@"
        longest = max(len(str(r.get("教学内容", ""))) / 38, len(str(r.get("教学要求", ""))) / 20, 1)
        ws.row_dimensions[row].height = max(14, 13 * int(longest + 0.999))
        for col in (COL["教学内容"], COL["教学要求"], COL["教材章节"]):
            al = copy.copy(ws.cell(row, col).alignment)
            al.wrap_text = True
            al.vertical = "center"
            ws.cell(row, col).alignment = al


def report(ws, rows):
    hours = sum(float(r["学时"] or 0) for r in rows)
    warnings = []
    header_total = None
    try:
        header_total = float(str(ws[TOTAL_HOURS_CELL].value).strip())
        if abs(header_total - hours) > 1e-6:
            warnings.append(f"学时合计 {hours:g} 与模板表头总学时 {header_total:g} 不一致，请核对。")
    except (TypeError, ValueError):
        pass
    bad = []
    for i in range(len(rows)):
        r = FIRST_ROW + i
        for c1, c2 in MERGES:
            ref = f"{get_column_letter(c1)}{r}:{get_column_letter(c2)}{r}"
            if ref not in {str(m) for m in ws.merged_cells.ranges}:
                bad.append(ref)
    return {"lesson_count": len(rows), "total_hours": hours, "template_total_hours": header_total,
            "hours_match": None if header_total is None else abs(header_total - hours) <= 1e-6,
            "missing_merges": bad, "warnings": warnings}
