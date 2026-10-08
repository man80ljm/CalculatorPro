"""学校成绩登记表：不调用 AI。xlsx / xls / 文字 PDF。"""
from __future__ import annotations

import io
import re
import unicodedata
from typing import Any

from openpyxl import Workbook, load_workbook

from web_app.service import ServiceError
from web_app.syllabus.extract_text import DOC_MESSAGE, SCAN_MESSAGE, suffix_of

_TERM = re.compile(r"(\d{4})\s*[-–—]\s*(\d{4})\s*学年\s*第\s*([12])\s*学期")
_LABELED = {
    "course_name": re.compile(r"课程名称\s*[：:]\s*(\S+)"),
    "course_code": re.compile(r"课程代码\s*[：:]\s*([A-Za-z0-9]+)"),
    "course_type": re.compile(r"课程性质\s*[：:]\s*(\S+)"),
    "college": re.compile(r"开课学院\s*[：:]\s*(\S+)"),
    "teacher": re.compile(r"任课教师\s*[：:]\s*(\S+)"),
    "credits": re.compile(r"学\s*分\s*[：:]\s*(\d+(?:\.\d+)?)"),
    "assess": re.compile(r"考核方式\s*[：:]\s*(\S+)"),
}
_EXAM = re.compile(r"实考\s*(\d+)\s*人")
_TOTAL = re.compile(r"总人数\s*(\d+)\s*人")
_PERCENT = re.compile(r"(\d+(?:\.\d+)?)\s*[%％]")
_PERCENT_PAREN = re.compile(r"[（(]\s*\d+(?:\.\d+)?\s*[%％]\s*[)）]")
_STOP = ("成绩统计", "平均成绩", "实考", "缓考", "分数段", "90分", "优秀", "良好", "及格", "不及格")
_BUCKET_LABEL = {"usual": "平时", "midterm": "期中", "final": "期末"}


def _clean(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).replace("\u3000", " ").split())


def _rows_from_xlsx(data: bytes) -> list[list[str]]:
    try:
        workbook = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:
        raise ServiceError("无法读取这份 xlsx 成绩登记表。") from exc
    sheet = workbook.active
    rows = []
    for row in sheet.iter_rows(max_row=sheet.max_row, max_col=sheet.max_column, values_only=True):
        rows.append([_clean(cell) for cell in row])
    return rows


def _rows_from_xls(data: bytes) -> list[list[str]]:
    import xlrd

    try:
        book = xlrd.open_workbook(file_contents=data)
    except Exception as exc:
        raise ServiceError("无法读取这份 xls 成绩登记表。") from exc
    sheet = book.sheet_by_index(0)
    rows = []
    for index in range(sheet.nrows):
        rows.append([_clean(sheet.cell_value(index, col)) for col in range(sheet.ncols)])
    return rows


def _rows_and_prose_from_pdf(data: bytes) -> tuple[list[list[str]], str]:
    import pdfplumber

    try:
        pdf = pdfplumber.open(io.BytesIO(data))
    except Exception as exc:
        raise ServiceError(SCAN_MESSAGE) from exc
    prose_parts: list[str] = []
    rows: list[list[str]] = []
    try:
        for page in pdf.pages:
            prose_parts.append(page.extract_text() or "")
            tables = page.extract_tables() or []
            if not tables:
                continue
            for table in tables:
                for row in table:
                    rows.append([_clean(cell) for cell in row])
    finally:
        pdf.close()
    prose = "\n".join(prose_parts)
    if len("".join(prose.split())) < 20 and not rows:
        raise ServiceError(SCAN_MESSAGE)
    if not any(_header_map(row) for row in rows):
        text_rows = []
        for line in prose.splitlines():
            pieces = [piece for piece in line.split() if piece]
            if pieces:
                text_rows.append(pieces)
        if any(_header_map(row) for row in text_rows):
            rows = text_rows
    if not rows:
        for line in prose.splitlines():
            rows.append([_clean(line)])
    return rows, prose


def column_label(value: Any) -> str:
    """去空白，去掉括号里的百分比，供老师看见的列名。"""
    text = _clean(value)
    text = _PERCENT_PAREN.sub("", text)
    return re.sub(r"\s+", "", text)


def normalize_header(value: Any) -> str:
    """列名和考核方式共用：全角半角统一后再比。"""
    return unicodedata.normalize("NFKC", column_label(value))


def _link_kind(norm: str) -> str | None:
    for kind, token in (("usual", "平时"), ("midterm", "期中"), ("final", "期末")):
        if norm == token or norm in {token + "成绩", token + "考核", token + "分数"}:
            return kind
    if norm in {"总评", "总评成绩", "总成绩"}:
        return "overall"
    return None


def _header_map(row: list[str]) -> dict[str, Any] | None:
    joined = "".join(row)
    if "姓名" not in joined:
        return None
    if "学号" not in joined and "班级" not in joined:
        return None
    mapping: dict[str, Any] = {
        "class": None,
        "sid": None,
        "name": None,
        "usual": None,
        "midterm": None,
        "final": None,
        "overall": None,
        "extras": [],
    }
    for index, cell in enumerate(row):
        if not cell:
            continue
        norm = normalize_header(cell)
        if not norm:
            continue
        if mapping["class"] is None and "班级" in cell and "学号" not in cell:
            mapping["class"] = index
            continue
        if mapping["sid"] is None and "学号" in cell:
            mapping["sid"] = index
            continue
        if mapping["name"] is None and "姓名" in cell:
            mapping["name"] = index
            continue
        if norm == "备注" or norm.startswith("备注"):
            continue
        kind = _link_kind(norm)
        if kind and mapping[kind] is None:
            mapping[kind] = index
            continue
        if kind:
            continue
        mapping["extras"].append({"index": index, "name": norm, "label": column_label(cell) or norm})
    if mapping["name"] is None:
        return None
    has_score = any(mapping[key] is not None for key in ("usual", "midterm", "final", "overall")) or mapping["extras"]
    if not has_score:
        return None
    return mapping


def _percent(cell: str) -> float | None:
    match = _PERCENT.search(cell or "")
    if not match:
        return None
    return float(match.group(1)) / 100.0


def _number(cell: str) -> float | None:
    text = (cell or "").strip().replace("%", "")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _metadata(prose: str) -> dict[str, str]:
    found = {key: "" for key in ("course_name", "course_code", "course_type", "college", "teacher", "credits", "assess")}
    for key, pattern in _LABELED.items():
        match = pattern.search(prose or "")
        if match:
            found[key] = match.group(1).strip()
    term = _TERM.search(prose or "")
    found["school_year_term"] = ""
    found["year_start"] = ""
    found["year_end"] = ""
    found["semester"] = ""
    if term:
        found["year_start"], found["year_end"], found["semester"] = term.group(1), term.group(2), term.group(3)
        found["school_year_term"] = f"{term.group(1)}-{term.group(2)}学年第{term.group(3)}学期"
    exam = _EXAM.search(prose or "")
    total = _TOTAL.search(prose or "")
    found["exam_count_text"] = exam.group(1) if exam else ""
    found["total_count_text"] = total.group(1) if total else ""
    return found


def parse_register(filename: str, data: bytes) -> dict:
    suffix = suffix_of(filename)
    if suffix == ".doc":
        raise ServiceError(DOC_MESSAGE)
    if suffix == ".xlsx":
        rows = _rows_from_xlsx(data)
        prose = "\n".join(" ".join(cell for cell in row if cell) for row in rows)
    elif suffix == ".xls":
        rows = _rows_from_xls(data)
        prose = "\n".join(" ".join(cell for cell in row if cell) for row in rows)
    elif suffix == ".pdf":
        rows, prose = _rows_and_prose_from_pdf(data)
    else:
        raise ServiceError("成绩登记表只接受 .xlsx、.xls 或带文字的 PDF。")

    header_index = None
    mapping = None
    for index, row in enumerate(rows):
        mapping = _header_map(row)
        if mapping:
            header_index = index
            break
    if header_index is None or mapping is None:
        raise ServiceError("成绩登记表无法识别：缺少表头（需要班级或学号、姓名，以及平时、期中、期末、总评或考核方式列）。")

    header = rows[header_index]
    percents = {
        "usual": _percent(header[mapping["usual"]]) if mapping["usual"] is not None else 0.0,
        "midterm": _percent(header[mapping["midterm"]]) if mapping["midterm"] is not None else 0.0,
        "final": _percent(header[mapping["final"]]) if mapping["final"] is not None else 0.0,
    }
    for key, value in list(percents.items()):
        if value is None:
            percents[key] = 0.0

    def cell_at(row: list[str], index: int | None) -> str:
        if index is None or index >= len(row):
            return ""
        return row[index]

    pending_extras: list[list[tuple[float | None, bool]]] = []
    students = []
    classes: list[str] = []
    for row in rows[header_index + 1 :]:
        blob = "".join(row)
        if any(token in blob for token in _STOP):
            break
        name = cell_at(row, mapping["name"])
        if not name or name in {"姓名", "合计"}:
            if students:
                break
            continue
        class_name = cell_at(row, mapping["class"])
        if class_name and class_name not in classes:
            classes.append(class_name)
        extra_cells: list[tuple[float | None, bool]] = []
        for meta in mapping["extras"]:
            raw = cell_at(row, meta["index"])
            number = _number(raw)
            extra_cells.append((number, bool(raw) and number is None))
        pending_extras.append(extra_cells)
        students.append(
            {
                "name": name,
                "class_name": class_name,
                "usual": _number(cell_at(row, mapping["usual"])) if mapping["usual"] is not None else None,
                "midterm": _number(cell_at(row, mapping["midterm"])) if mapping["midterm"] is not None else None,
                "final": _number(cell_at(row, mapping["final"])) if mapping["final"] is not None else None,
                "overall": _number(cell_at(row, mapping["overall"])) if mapping["overall"] is not None else None,
            }
        )
    if not students:
        raise ServiceError("成绩登记表无法识别：表头下面没有学生成绩行。")

    extra_columns = []
    for column_index, meta in enumerate(mapping["extras"]):
        values = [item[column_index][0] for item in pending_extras]
        texts = sum(1 for item in pending_extras if item[column_index][1])
        numbers = sum(1 for value in values if value is not None)
        if texts and numbers == 0:
            continue
        extra_columns.append(
            {
                "id": len(extra_columns),
                "name": meta["name"],
                "label": meta["label"],
                "values": values,
            }
        )
    for student, cells in zip(students, pending_extras):
        student["extras"] = {}
    for column in extra_columns:
        for student, value in zip(students, column["values"]):
            student["extras"][column["label"]] = value

    meta = _metadata(prose)
    exam_count = int(meta["exam_count_text"]) if meta["exam_count_text"] else len(students)
    return {
        "course_name": meta["course_name"],
        "course_code": meta["course_code"],
        "course_type": meta["course_type"],
        "college": meta["college"],
        "teacher": meta["teacher"],
        "credits": meta["credits"],
        "assess": meta["assess"],
        "school_year_term": meta["school_year_term"],
        "year_start": meta["year_start"],
        "year_end": meta["year_end"],
        "semester": meta["semester"],
        "classes": classes,
        "student_count": len(students),
        "exam_count": exam_count,
        "percents": percents,
        "students": students,
        "link_columns": {
            "usual": mapping["usual"] is not None,
            "midterm": mapping["midterm"] is not None,
            "final": mapping["final"] is not None,
            "overall": mapping["overall"] is not None,
        },
        "extra_columns": extra_columns,
    }


def limit_register_to_class(parsed: dict, class_name: str) -> dict:
    """一个学期只收一个班。额外考核列和学生行一起裁掉，避免错位。"""
    chosen = (class_name or "").strip()
    if not chosen:
        return parsed
    students = list(parsed.get("students") or [])
    kept = [index for index, student in enumerate(students) if (student.get("class_name") or "") == chosen]
    if not kept:
        raise ServiceError("登记表里没有这个班级的成绩")
    data = dict(parsed)
    data["students"] = [students[index] for index in kept]
    columns = []
    for column in parsed.get("extra_columns") or []:
        values = list(column.get("values") or [])
        copied = dict(column)
        copied["values"] = [values[index] if index < len(values) else None for index in kept]
        columns.append(copied)
    data["extra_columns"] = columns
    data["classes"] = [chosen]
    data["student_count"] = len(kept)
    if len(kept) == len(students):
        data["exam_count"] = parsed.get("exam_count") or len(kept)
    else:
        data["exam_count"] = len(kept)
    return data


def markdown_for_ai(parsed: dict) -> str:
    """只保留表头、人数和班级。考核方式只出现列名，不包含姓名、学号和学生分数。"""
    classes = "、".join(parsed.get("classes") or [])
    percents = parsed.get("percents") or {}
    headers = ["班级", "学号", "姓名"]
    for column in parsed.get("extra_columns") or []:
        label = str(column.get("label") or column.get("name") or "").replace("|", "")
        if label:
            headers.append(label)
    headers.append(f"平时({_pct(percents.get('usual'))})")
    headers.append(f"期中({_pct(percents.get('midterm'))})")
    headers.append(f"期末({_pct(percents.get('final'))})")
    if (parsed.get("link_columns") or {}).get("overall"):
        headers.append("总评")
    lines = [
        parsed.get("school_year_term") or "",
        f"课程名称：{parsed.get('course_name') or ''} 课程代码：{parsed.get('course_code') or ''} 课程性质：{parsed.get('course_type') or ''}",
        f"开课学院：{parsed.get('college') or ''} 任课教师：{parsed.get('teacher') or ''} 学分：{parsed.get('credits') or ''}",
        "[表格1]",
        "| " + " | ".join(headers) + " |",
        f"（学生名单已省略：共 {parsed.get('student_count') or 0} 行学生记录；班级取值：{classes}）",
        "[/表格1]",
        f"实考 {parsed.get('exam_count') or 0}人 总人数 {parsed.get('student_count') or 0}人",
    ]
    return "\n".join(line for line in lines if line)


def _pct(value: float | None) -> str:
    if value is None:
        return ""
    return f"{round(float(value) * 100)}%"


def _as_float(value: Any, default: float) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _link_bucket(name: str) -> str | None:
    if "平时" in name:
        return "usual"
    if "期中" in name:
        return "midterm"
    if "期末" in name:
        return "final"
    return None


def _match_rank(method_norm: str, column_norm: str) -> int:
    if not method_norm or not column_norm:
        return 0
    if method_norm == column_norm:
        return 3
    shorter, longer = (
        (method_norm, column_norm) if len(method_norm) <= len(column_norm) else (column_norm, method_norm)
    )
    if len(shorter) < 2:
        return 0
    if shorter in longer:
        return 2
    return 0


def _score_targets(links: list[dict]) -> list[dict]:
    targets = []
    for link in links or []:
        if not isinstance(link, dict):
            continue
        ratio = _as_float(link.get("ratio"), 0.0)
        methods = link.get("methods") if isinstance(link.get("methods"), list) else []
        weighted = []
        for method in methods:
            if not isinstance(method, dict):
                continue
            name = str(method.get("name") or "").strip()
            if not name:
                continue
            weight = _as_float(method.get("subtotal"), 1.0)
            if weight > 1e-9 and ratio > 1e-9:
                weighted.append(name)
        single = len(weighted) == 1
        link_name = str(link.get("name") or "").strip()
        for name in weighted:
            targets.append(
                {
                    "link": link_name,
                    "method": name,
                    "norm": normalize_header(name),
                    "single": single,
                    "bucket": _link_bucket(link_name) if single else None,
                }
            )
    return targets


def _candidate_ranks(target: dict, columns: list[dict], buckets: dict[str, dict]) -> dict[str, int]:
    ranks: dict[str, int] = {}
    for column in columns:
        rank = _match_rank(target["norm"], column["norm"])
        if rank and (column["id"] not in ranks or rank > ranks[column["id"]]):
            ranks[column["id"]] = rank
    bucket_name = target.get("bucket")
    if target.get("single") and bucket_name and bucket_name in buckets:
        column = buckets[bucket_name]
        ranks.setdefault(column["id"], 1)
    return ranks


def _best_assignments(targets: list[dict], columns: list[dict], buckets: dict[str, dict]) -> list[tuple]:
    options = [_candidate_ranks(target, columns, buckets) for target in targets]
    found: list[tuple[int, int, tuple]] = []
    limit = {"visited": 0}

    def walk(index: int, used: set[str], chosen: list[str | None], coverage: int, rank_sum: int) -> None:
        limit["visited"] += 1
        if limit["visited"] > 100000:
            return
        if index == len(targets):
            found.append((coverage, rank_sum, tuple(chosen)))
            return
        walk(index + 1, used, chosen + [None], coverage, rank_sum)
        for column_id, rank in options[index].items():
            if column_id in used:
                continue
            used.add(column_id)
            walk(index + 1, used, chosen + [column_id], coverage + 1, rank_sum + rank)
            used.remove(column_id)

    walk(0, set(), [], 0, 0)
    if not found or limit["visited"] > 100000:
        return []
    found.sort(key=lambda item: (item[0], item[1]), reverse=True)
    best_coverage, best_rank, _assignment = found[0]
    return [item for item in found if item[0] == best_coverage and item[1] == best_rank]


def detect_register_mode(parsed: dict, links: list[dict]) -> dict:
    """用关系表判断登记表该走正向还是逆向。不调用模型。"""
    students = parsed.get("students") or []
    columns = []
    by_id: dict[str, dict] = {}
    for column in parsed.get("extra_columns") or []:
        item = {
            "id": f"col:{column.get('id')}",
            "norm": column.get("name") or "",
            "label": column.get("label") or column.get("name") or "",
            "kind": "column",
            "values": list(column.get("values") or []),
        }
        columns.append(item)
        by_id[item["id"]] = item
    present = parsed.get("link_columns") or {}
    buckets: dict[str, dict] = {}
    for bucket, label in _BUCKET_LABEL.items():
        if not present.get(bucket):
            continue
        item = {
            "id": f"bucket:{bucket}",
            "norm": label,
            "label": label,
            "kind": "bucket",
            "bucket": bucket,
            "values": [student.get(bucket) for student in students],
        }
        buckets[bucket] = item
        by_id[item["id"]] = item

    targets = _score_targets(links if isinstance(links, list) else [])
    best = _best_assignments(targets, columns, buckets) if targets else []
    assigned: dict[int, dict] = {}
    stable = len(best) == 1
    if stable:
        for index, column_id in enumerate(best[0][2]):
            if column_id:
                assigned[index] = by_id[column_id]

    matched: dict[str, str] = {}
    sources: dict[str, dict] = {}
    missing: list[str] = []
    ambiguous: list[str] = []
    for index, target in enumerate(targets):
        column = assigned.get(index)
        key = target["method"]
        if key in matched:
            key = f"{target['link']}/{target['method']}"
        if column is not None:
            matched[key] = column["label"]
            sources[f"{target['link']}||{target['method']}"] = column
            continue
        if _candidate_ranks(target, columns, buckets):
            ambiguous.append(target["method"])
        else:
            missing.append(target["method"])

    has_link = any(present.get(bucket) for bucket in _BUCKET_LABEL)
    has_extra = bool(columns)
    only_overall = bool(present.get("overall")) and not has_link and not has_extra
    if not targets:
        mode = "reverse"
        reason = "关系表里没有权重大于 0 的考核方式，按环节分数导入"
    elif stable and not missing and not ambiguous:
        mode = "forward"
        pairs = [f"{method}←{label}" for method, label in matched.items()]
        reason = "考核方式已对应登记表列：" + "、".join(pairs)
    else:
        mode = "reverse"
        if only_overall:
            uncovered = "、".join(missing or ambiguous or [target["method"] for target in targets])
            reason = "登记表只有总评，没有平时、期中、期末或考核方式列，已把总评当作各环节同分"
            if uncovered:
                reason += f"（未覆盖：{uncovered}）"
        else:
            parts = []
            if missing:
                parts.append("缺少考核方式列：" + "、".join(missing))
            if ambiguous:
                parts.append("无法唯一匹配：" + "、".join(ambiguous))
            if matched:
                parts.append("已匹配：" + "、".join(f"{method}←{label}" for method, label in matched.items()))
            if not parts:
                parts.append("登记表只有平时、期中、期末环节分数，无法对应全部考核方式")
            reason = "；".join(parts)
    return {
        "mode": mode,
        "reason": reason,
        "matched": matched,
        "missing": missing,
        "ambiguous": ambiguous,
        "sources": sources,
    }


def forward_workbook(parsed: dict, links: list[dict], detection: dict) -> tuple[bytes, list[str]]:
    """写成正向模板同款二级表头，供 process_forward_grades 直接读取。"""
    notes: list[str] = []
    warned: set[str] = set()
    book = Workbook()
    sheet = book.active
    sheet.title = "正向成绩模板"
    sheet.cell(row=1, column=1, value="姓名")
    sheet.cell(row=2, column=1, value="姓名")
    sheet.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
    sources = detection.get("sources") or {}
    column = 2
    plan: list[tuple[str, dict | None]] = []
    for link in links or []:
        if not isinstance(link, dict):
            continue
        methods = link.get("methods") if isinstance(link.get("methods"), list) else []
        if not methods:
            methods = [{"name": "无"}]
        link_name = str(link.get("name") or "").strip()
        start = column
        for method in methods:
            method_name = str((method or {}).get("name") or "无").strip() or "无"
            sheet.cell(row=2, column=column, value=method_name)
            source = sources.get(f"{link_name}||{method_name}")
            plan.append((method_name, source))
            column += 1
        end = column - 1
        for index in range(start, end + 1):
            sheet.cell(row=1, column=index, value=link_name)
        if end > start:
            sheet.merge_cells(start_row=1, start_column=start, end_row=1, end_column=end)

    students = parsed.get("students") or []
    for row_index, student in enumerate(students):
        sheet.cell(row=3 + row_index, column=1, value=student.get("name") or "")
        for offset, (method_name, source) in enumerate(plan):
            value = 0
            if source is not None:
                values = source.get("values") or []
                raw = values[row_index] if row_index < len(values) else None
                if raw is None:
                    label = source.get("label") or method_name
                    if label not in warned:
                        warned.add(label)
                        notes.append(f"{label}有空值，已按 0 分计")
                    value = 0
                else:
                    value = raw
            sheet.cell(row=3 + row_index, column=2 + offset, value=value)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue(), notes


def reverse_workbook(parsed: dict, links: list[dict]) -> tuple[bytes, list[str]]:
    """转成逆向计算能读的表：姓名 + 关系表里的环节名。"""
    notes: list[str] = []
    columns = ["姓名"]
    buckets: list[tuple[str, str]] = []
    safe_links = [link for link in (links or []) if isinstance(link, dict)]
    link_columns = parsed.get("link_columns") if isinstance(parsed.get("link_columns"), dict) else None
    for bucket, label in (("usual", "平时"), ("midterm", "期中"), ("final", "期末")):
        if link_columns is not None and not link_columns.get(bucket):
            continue
        matched = [link for link in safe_links if label in str(link.get("name") or "")]
        percent = float((parsed.get("percents") or {}).get(bucket) or 0)
        has_score = any(item.get(bucket) not in (None, 0) for item in parsed.get("students") or [])
        if not matched:
            if percent > 1e-9 or has_score:
                notes.append(f"登记表的{label}列在关系表里没有名称含“{label}”的环节")
            continue
        if len(matched) > 1:
            notes.append(f"关系表有多个{label}环节，分数写入“{matched[0]['name']}”")
        columns.append(str(matched[0]["name"]))
        buckets.append((bucket, str(matched[0]["name"])))

    book = Workbook()
    sheet = book.active
    sheet.title = "逆向成绩"
    if not buckets and (parsed.get("link_columns") or {}).get("overall"):
        link_names = [str(link.get("name") or "").strip() for link in safe_links]
        link_names = [name for name in link_names if name]
        notes.append("登记表只有总评，已按各环节同分写入逆向成绩")
        columns = ["姓名"] + link_names
        sheet.append(columns)
        for student in parsed.get("students") or []:
            score = student.get("overall")
            sheet.append([student.get("name") or ""] + [0 if score is None else score for _name in link_names])
        buffer = io.BytesIO()
        book.save(buffer)
        return buffer.getvalue(), notes

    sheet.append(columns)
    for student in parsed.get("students") or []:
        row = [student.get("name") or ""]
        values = {name: student.get(bucket) for bucket, name in buckets}
        for column in columns[1:]:
            score = values.get(column)
            row.append(0 if score is None else score)
        sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue(), notes


def register_course_message(settings: dict, course_name: str, parsed: dict) -> str:
    """课程名称或课程代码对不上时，返回要老师确认的那句话。对得上则返回空字符串。"""
    basic = settings.get("course_basic_info") if isinstance(settings.get("course_basic_info"), dict) else {}
    opened = settings.get("course_open_info") if isinstance(settings.get("course_open_info"), dict) else {}
    name = str(basic.get("course_name") or opened.get("course_name") or course_name or "").strip()
    code = str(basic.get("course_code") or "").strip()
    reg_name = str(parsed.get("course_name") or "").strip()
    reg_code = str(parsed.get("course_code") or "").strip()
    name_diff = bool(reg_name and name and reg_name != name)
    code_diff = bool(reg_code and code and reg_code != code)
    if not name_diff and not code_diff:
        return ""
    left_name = reg_name or "未识别课程"
    right_name = name or "未命名课程"
    left_code = reg_code or "无课程代码"
    right_code = code or "无课程代码"
    return f"登记表是《{left_name}》（{left_code}），当前课程是《{right_name}》（{right_code}），确定要导入吗？"


def headcount_warnings(
    enrolled: int, exam: int, file_count: int, file_exam: int | None = None, source: str = "登记表"
) -> list[str]:
    """上课人数或考核人数和成绩人数不一致时给出提醒，不阻止导入或计算。"""
    notes = []
    try:
        enrolled = int(enrolled or 0)
    except (TypeError, ValueError):
        enrolled = 0
    try:
        exam = int(exam or 0)
    except (TypeError, ValueError):
        exam = 0
    try:
        file_count = int(file_count or 0)
    except (TypeError, ValueError):
        file_count = 0
    compared_exam = file_exam if file_exam else file_count
    try:
        compared_exam = int(compared_exam or 0)
    except (TypeError, ValueError):
        compared_exam = 0
    if enrolled and file_count and enrolled != file_count:
        notes.append(f"本学期上课人数是 {enrolled}，{source}有 {file_count} 人，人数不一致。")
    if exam and compared_exam and exam != compared_exam:
        notes.append(f"本学期考核人数是 {exam}，{source}实考 {compared_exam} 人，人数不一致。")
    return notes


def ratio_warnings(percents: dict, ratios: dict) -> list[str]:
    warnings = []
    labels = (("usual", "平时"), ("midterm", "期中"), ("final", "期末"))
    for key, label in labels:
        left = float((percents or {}).get(key) or 0)
        right = float((ratios or {}).get(key) or 0)
        if abs(left - right) > 0.015:
            warnings.append(f"{label}占比不一致：登记表 {round(left * 100)}%，关系表 {round(right * 100)}%")
    return warnings
