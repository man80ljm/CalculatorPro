"""毕业要求矩阵：合成 docx 的几种版式，以及本地样例大纲。"""
import io
import json
import os

import pytest
from docx import Document

from core_app.course_docs import export_grad_req_docx
from tests.test_syllabus_rules import _ai
from web_app.syllabus.extract_text import docx_to_markdown
from web_app.syllabus.grad_matrix import NAME_ONLY_NOTE
from web_app.syllabus.rules import build_draft

SAMPLE_SYLLABUS = os.environ.get(
    "CALC_SAMPLE_SYLLABUS_DOCX",
    "/home/box/agent-data/agents/7034809e-f119-43b6-8ff6-3505cccdb393/attachments/"
    "f50d80ae25a0667002c6fb819fbe6f902b80fd26e67df0572760544e273ba029.docx",
)


def _docx(paragraphs, rows) -> bytes:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            table.rows[row_index].cells[col_index].text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _draft_from_docx(data: bytes, objectives=1):
    text = docx_to_markdown(data)
    ai = _ai()
    ai["objectives"] = [{"index": index + 1, "text": f"目标{index + 1}"} for index in range(objectives)]
    ai["objectives_status"] = "已填"
    ai["grad_req_map"] = [
        {
            "objective": "课程目标1",
            "indicator": {"value": "模型乱填", "status": "已填"},
            "requirement": {"value": "模型乱填", "status": "已填"},
            "strength": "L",
        }
    ]
    return build_draft(ai, text, "", None)


def test_name_only_matrix_fills_indicator_and_strength():
    data = _docx(
        ["（二）课程目标与毕业要求的矩阵关系"],
        [
            ["课程目标 毕业要求指标点", "课程目标1", "课程目标2", "课程目标3", "课程目标4", "课程目标5"],
            ["理论知识", "H", "", "", "", ""],
            ["专业视野", "", "M", "", "", ""],
            ["知识应用", "", "", "H", "", ""],
            ["专业技能", "", "", "", "H", ""],
            ["创新创业", "", "", "", "", "M"],
        ],
    )
    draft = _draft_from_docx(data, 5)
    expected = [
        ("理论知识", "H"),
        ("专业视野", "M"),
        ("知识应用", "H"),
        ("专业技能", "H"),
        ("创新创业", "M"),
    ]
    assert len(draft["grad_req_map"]) == 5
    for item, (name, strength) in zip(draft["grad_req_map"], expected):
        assert item["requirement"]["value"] == name
        assert item["requirement"]["status"] == "已填"
        assert item["requirement"]["source"] == "大纲·矩阵行标签"
        assert item["requirement"]["reason"] == ""
        assert item["requirement"]["note"] == NAME_ONLY_NOTE
        assert item["indicator"]["value"] == name
        assert item["indicator"]["status"] == "已填"
        assert item["indicator"]["source"] == "大纲·矩阵行标签"
        assert item["indicator"]["note"] == NAME_ONLY_NOTE
        assert item["strength"] == strength
        assert item["strength_source"] == "大纲·矩阵"
    longer = _draft_from_docx(data, 6)
    missing = longer["grad_req_map"][5]
    assert missing["requirement"]["status"] == "需手填"
    assert missing["indicator"]["status"] == "需手填"
    assert missing["requirement"]["reason"] == "大纲矩阵里课程目标6没有标 H/M/L"
    pdf_markdown = docx_to_markdown(data).replace("[表格1]", "[表格p1-1]").replace("[/表格1]", "[/表格p1-1]")
    again = build_draft(_ai(), pdf_markdown, "", None)
    assert again["grad_req_map"][0]["indicator"]["value"] == "理论知识"
    assert again["grad_req_map"][1]["strength"] == "M"


def test_numbered_matrix_uses_codes_and_definitions():
    data = _docx(
        [
            "毕业要求2：问题分析",
            "毕业要求3：工程知识",
            "指标点2-1：能够设计实验方案",
            "指标点3.2：能够综合运用专业知识",
            "（二）课程目标与毕业要求的矩阵关系",
        ],
        [
            ["毕业要求", "指标点", "课程目标1", "课程目标2"],
            ["毕业要求3", "3.2 专业技能", "H", ""],
            ["毕业要求2", "指标点 2-1", "", "M"],
        ],
    )
    draft = _draft_from_docx(data, 2)
    first, second = draft["grad_req_map"]
    assert first["requirement"]["status"] == "已填"
    assert first["requirement"]["value"] == "毕业要求3：工程知识"
    assert first["requirement"]["source"] == "大纲·矩阵行标签；大纲·毕业要求定义"
    assert first["indicator"]["value"] == "3.2 专业技能：能够综合运用专业知识"
    assert first["indicator"]["source"] == "大纲·矩阵行标签；大纲·指标点定义"
    assert first["strength"] == "H"
    assert "note" not in first["requirement"]
    assert "note" not in first["indicator"]
    assert second["requirement"]["value"] == "毕业要求2：问题分析"
    assert second["indicator"]["value"] == "2-1 能够设计实验方案"
    assert second["strength"] == "M"


def test_one_column_keeps_multiple_strengths():
    data = _docx(
        ["（二）课程目标与毕业要求的矩阵关系"],
        [
            ["毕业要求指标点", "课程目标1", "课程目标2"],
            ["理论知识", "H/M", ""],
            ["专业视野", "L", "H"],
            ["知识应用", "", "M"],
        ],
    )
    draft = _draft_from_docx(data, 2)
    first, second = draft["grad_req_map"]
    assert first["indicator"]["value"] == "理论知识；专业视野"
    assert first["indicator"]["status"] == "已填"
    assert first["indicator"]["source"] == "大纲·矩阵行标签"
    assert first["indicator"]["note"] == NAME_ONLY_NOTE
    assert first["requirement"]["value"] == "理论知识；专业视野"
    assert first["requirement"]["status"] == "已填"
    assert first["requirement"]["note"] == NAME_ONLY_NOTE
    assert first["strength"] == "H、M；L"
    assert second["indicator"]["value"] == "专业视野；知识应用"
    assert second["requirement"]["value"] == "专业视野；知识应用"
    assert second["strength"] == "H；M"


def test_syllabus_without_matrix_keeps_manual_reasons():
    data = _docx(["三、课程目标", "课程目标1：能完成作品"], [["考核环节", "课程目标1"], ["平时考核", "100%"]])
    text = docx_to_markdown(data)
    ai = _ai()
    ai["grad_req_map"] = []
    ai["objectives"] = [{"index": 1, "text": "能完成作品"}]
    draft = build_draft(ai, text, "", None)
    grad = draft["grad_req_map"][0]
    assert grad["indicator"]["status"] == "需手填"
    assert grad["indicator"]["reason"] == "大纲没有课程目标与毕业要求的矩阵"
    assert grad["requirement"]["status"] == "需手填"
    assert grad["requirement"]["reason"] == "大纲没有课程目标与毕业要求的矩阵"
    assert grad["strength"] == ""


def test_grad_export_shows_strength(tmp_path):
    path = export_grad_req_docx(
        [{"objective": "课程目标1", "requirement": "毕业要求3：工程知识", "indicator": "3.2 专业技能", "strength": "H"}],
        output_dir=str(tmp_path),
    )
    document = Document(path)
    header = [cell.text.strip() for cell in document.tables[0].rows[0].cells]
    body = [cell.text.strip() for cell in document.tables[0].rows[1].cells]
    assert header == ["课程目标", "支撑的毕业要求", "支撑的毕业要求指标点", "支撑强度"]
    assert body[3] == "H"


def test_sample_syllabus_matrix_when_present():
    if not os.path.isfile(SAMPLE_SYLLABUS):
        pytest.skip("样例大纲不存在")
    draft = _draft_from_docx(open(SAMPLE_SYLLABUS, "rb").read(), 5)
    expected = [
        ("理论知识", "H"),
        ("专业视野", "M"),
        ("知识应用", "H"),
        ("专业技能", "H"),
        ("创新创业", "M"),
    ]
    assert len(draft["grad_req_map"]) == 5
    for item, (name, strength) in zip(draft["grad_req_map"], expected):
        assert item["requirement"]["value"] == name
        assert item["requirement"]["status"] == "已填"
        assert item["requirement"]["source"] == "大纲·矩阵行标签"
        assert item["requirement"]["note"] == NAME_ONLY_NOTE
        assert item["indicator"]["value"] == name
        assert item["indicator"]["status"] == "已填"
        assert item["indicator"]["source"] == "大纲·矩阵行标签"
        assert item["indicator"]["note"] == NAME_ONLY_NOTE
        assert item["strength"] == strength
        assert item["strength_source"] == "大纲·矩阵"


def test_full_syllabus_matrix_stays_local_and_is_not_sent(monkeypatch):
    monkeypatch.setenv("DEEPSEEK_API_KEY", "sk-test")
    sent = []

    def fake_complete(messages, **_kwargs):
        sent.append(messages)
        return {
            "content": json.dumps(_ai(), ensure_ascii=False),
            "model": "deepseek-flash",
            "prompt_tokens": 3,
            "completion_tokens": 2,
        }

    monkeypatch.setattr("web_app.syllabus.service.complete", fake_complete)
    document = Document()
    document.add_paragraph("一、基本信息")
    document.add_paragraph("适用专业：数字媒体艺术")
    roster = document.add_table(rows=2, cols=2)
    roster.rows[0].cells[0].text = "学号"
    roster.rows[0].cells[1].text = "姓名"
    roster.rows[1].cells[0].text = "8800112299"
    roster.rows[1].cells[1].text = "合成生甲"
    document.add_paragraph("四、教学内容")
    document.add_paragraph("不应送出的全文标记句ABCDEF")
    matrix = document.add_table(rows=2, cols=2)
    matrix.rows[0].cells[0].text = "毕业要求指标点"
    matrix.rows[0].cells[1].text = "课程目标1"
    matrix.rows[1].cells[0].text = "理论知识"
    matrix.rows[1].cells[1].text = "H"
    buffer = io.BytesIO()
    document.save(buffer)
    from web_app.syllabus.service import extract_draft

    draft = extract_draft("大纲.docx", buffer.getvalue())
    assert len(sent) == 1
    user = next(item["content"] for item in sent[0] if item["role"] == "user")
    assert "数字媒体艺术" in user
    assert "共 1 行学生记录" in user
    assert "合成生甲" not in user
    assert "8800112299" not in user
    assert "不应送出的全文标记句ABCDEF" not in user
    assert "理论知识" not in user
    filled = draft["grad_req_map"][0]
    assert filled["requirement"]["value"] == "理论知识"
    assert filled["requirement"]["status"] == "已填"
    assert filled["indicator"]["value"] == "理论知识"
    assert filled["strength"] == "H"
