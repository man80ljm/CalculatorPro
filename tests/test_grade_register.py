"""成绩登记表：合成的两种版式、xls，以及占比警告不阻止导入。"""
import io
import json

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Table

from tests.test_web_flow import PASSWORD, _settings
from web_app.grade_register import detect_register_mode, markdown_for_ai, parse_register
from web_app.limiter import reset_limiters
from web_app.service import run_calculation


def _xlsx_like_handcraft(
    path_or_buf,
    count=35,
    course_name="创意手作",
    course_code="BJ2230005",
    single_class="",
    term="2024-2025学年第1学期",
    teacher="章老师",
):
    book = Workbook()
    sheet = book.active
    title = "广东第二师范学院课程成绩登记表" if not term else f"广东第二师范学院{term}课程成绩登记表"
    sheet["A1"] = title
    sheet["A2"] = f"课程名称：{course_name}                           课程代码：{course_code}                          课程性质：专业任选课"
    sheet["A3"] = f"开课学院：美术学院                           任课教师：{teacher}                               学分：3  考核方式：考查"
    sheet["A5"] = "班级"
    sheet["B5"] = "学号"
    sheet["C5"] = ""
    sheet["D5"] = "姓名"
    sheet["E5"] = "平时(30%)"
    sheet["F5"] = "期中(0%)"
    sheet["G5"] = "期末(70%)"
    sheet["H5"] = "总评"
    sheet["I5"] = "备注"
    for index in range(count):
        row = 6 + index
        klass = single_class or ("23美术教育A" if index < 20 else "23数字媒体")
        sheet.cell(row, 1, klass)
        sheet.cell(row, 2, f"230000{index:05d}")
        sheet.cell(row, 4, f"学生{index}")
        sheet.cell(row, 5, 80)
        sheet.cell(row, 6, 0)
        sheet.cell(row, 7, 90)
    sheet.cell(6 + count + 1, 1, "成绩统计")
    sheet.cell(6 + count + 2, 1, f"实考 {count}人 总人数 {count}人")
    book.save(path_or_buf)


def _pdf_like_engine(count=31) -> bytes:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import Paragraph
    from reportlab.lib.styles import ParagraphStyle

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    buffer = io.BytesIO()
    document = SimpleDocTemplate(buffer, pagesize=A4)
    style = ParagraphStyle("cn", fontName="STSong-Light", fontSize=8, leading=11)
    story = [
        Paragraph("广东第二师范学院2025-2026学年第2学期课程成绩登记表", style),
        Paragraph("课程名称：2D游戏引擎 课程代码：BJ2330188 课程性质：专业限选课", style),
        Paragraph("开课学院：美术学院 任课教师：黄老师 学分：3 考核方式：考查", style),
    ]
    header = ["班级", "学号", "姓名", "平时(30%)", "期中(0%)", "期末(70%)", "总评", "备注"]
    data = [[Paragraph(cell, style) for cell in header]]
    for index in range(count):
        row = ["24数字媒体", f"240000{index:05d}", f"学员{index}", "88", "0", "86", "86.6", ""]
        data.append([Paragraph(cell, style) for cell in row])
    story.append(Table(data, colWidths=[70, 78, 48, 58, 50, 58, 36, 36]))
    story.append(Paragraph(f"实考 {count}人 总人数 {count}人", style))
    document.build(story)
    return buffer.getvalue()


def _to_xls(xlsx_bytes: bytes) -> bytes:
    import xlwt

    workbook = load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    sheet = workbook.active
    book = xlwt.Workbook()
    out = book.add_sheet("Sheet1")
    for row in sheet.iter_rows(max_row=sheet.max_row, max_col=sheet.max_column):
        for cell in row:
            if cell.value is not None:
                out.write(cell.row - 1, cell.column - 1, cell.value)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def test_synthetic_formats_and_zero_midterm():
    buffer = io.BytesIO()
    _xlsx_like_handcraft(buffer, 35)
    parsed = parse_register("创意手作.xlsx", buffer.getvalue())
    assert parsed["student_count"] == 35
    assert parsed["exam_count"] == 35
    assert parsed["percents"]["midterm"] == 0
    assert parsed["teacher"] == "章老师"
    assert parsed["course_code"] == "BJ2230005"
    assert "23美术教育A" in parsed["classes"]
    assert "23数字媒体" in parsed["classes"]
    assert all(item["midterm"] == 0 for item in parsed["students"])

    xls = parse_register("创意手作.xls", _to_xls(buffer.getvalue()))
    assert xls["student_count"] == 35

    pdf = parse_register("引擎.pdf", _pdf_like_engine(31))
    assert pdf["student_count"] == 31
    assert pdf["classes"] == ["24数字媒体"]
    assert pdf["school_year_term"] == "2025-2026学年第2学期"
    assert pdf["course_type"] == "专业限选课"
    assert pdf["percents"]["usual"] == pytest.approx(0.3)
    assert pdf["percents"]["final"] == pytest.approx(0.7)


def test_real_samples_when_configured():
    import os

    pdf = os.environ.get("CALC_SAMPLE_REGISTER_PDF", "")
    xlsx = os.environ.get("CALC_SAMPLE_REGISTER_XLSX", "")
    if not pdf or not xlsx:
        pytest.skip("未设置真实样例路径")
    assert parse_register("real.pdf", open(pdf, "rb").read())["student_count"] == 31
    assert parse_register("real.xlsx", open(xlsx, "rb").read())["student_count"] == 35


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("SECRET_KEY", "test-secret-key")
    monkeypatch.setenv("ALLOW_SQLITE", "1")
    monkeypatch.setenv("DATABASE_URL", "sqlite:///" + (tmp_path / "app.db").as_posix())
    monkeypatch.setenv("UPLOAD_DIR", str(tmp_path / "uploads"))
    monkeypatch.setenv("COOKIE_SECURE", "false")
    monkeypatch.setenv("DEEPSEEK_API_KEY", "")
    reset_limiters()
    from web_app.app import create_app

    with TestClient(create_app()) as test_client:
        assert test_client.post("/api/register", json={"username": "teacher", "password": PASSWORD}).status_code == 200
        assert test_client.post("/api/login", json={"username": "teacher", "password": PASSWORD}).status_code == 200
        yield test_client


def test_import_warns_on_ratio_mismatch_and_does_not_overwrite(client):
    settings = _settings(mode="forward")
    settings["ratios"] = {"usual": 0.4, "midterm": 0.0, "final": 0.6}
    settings["relation_payload"]["links"][0]["ratio"] = 0.4
    settings["relation_payload"]["links"][1]["ratio"] = 0.6
    created = client.post("/api/courses", json={"name": "导入课", "settings": settings})
    assert created.status_code == 200, created.text
    course_id = created.json()["id"]
    buffer = io.BytesIO()
    _xlsx_like_handcraft(buffer, 4)
    # 单班，避免多班级警告淹没占比警告。把班级改成同一个需要重建。这里接受多班警告，并检查占比。
    response = client.post(
        f"/api/courses/{course_id}/grade-register?confirm=1",
        files={"file": ("reg.xlsx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["student_count"] == 4
    assert body["mode"] == "forward"
    assert any("平时占比不一致" in item for item in body["warnings"])
    assert body["conflicts"] == []
    course = client.get(f"/api/courses/{course_id}").json()
    assert course["current_term"]["teacher"] == "章老师"
    assert course["current_term"]["class_name"] == "23美术教育A"
    assert course["current_term"]["major"] == "美术教育"
    assert course["current_term"]["year_start"] == "2024"
    assert course["current_term"]["semester"] == "1"
    untouched = next(item for item in course["terms"] if item["class_name"] == "软工2201")
    assert untouched["teacher"] == "王老师"
    assert course["settings"]["mode"] == "forward"
    grade = next(item for item in course["files"] if item["kind"] == "grade")
    downloaded = client.get(f"/api/courses/{course_id}/files/{grade['id']}")
    sheet = load_workbook(io.BytesIO(downloaded.content)).active
    headers = [sheet.cell(1, col).value for col in range(1, sheet.max_column + 1)]
    assert headers[0] == "姓名"
    assert "平时考核" in headers
    assert "期末考核" in headers


def _named_settings(name, code, students=0, exam=0):
    settings = _settings(mode="forward")
    settings["course_basic_info"] = dict(settings["course_basic_info"])
    settings["course_open_info"] = dict(settings["course_open_info"])
    settings["course_basic_info"]["course_name"] = name
    settings["course_open_info"]["course_name"] = name
    settings["course_basic_info"]["course_code"] = code
    settings["student_count"] = students
    settings["course_basic_info"]["student_count"] = str(students or "")
    settings["course_basic_info"]["exam_count"] = str(exam or "")
    return settings


def _register_bytes(count, course_name, course_code, term="2024-2025学年第1学期", teacher="章老师"):
    buffer = io.BytesIO()
    _xlsx_like_handcraft(
        buffer,
        count,
        course_name=course_name,
        course_code=course_code,
        single_class="24数字媒体",
        term=term,
        teacher=teacher,
    )
    return buffer.getvalue()


def test_register_course_mismatch_needs_confirm(client):
    created = client.post("/api/courses", json={"name": "2D游戏引擎", "settings": _named_settings("2D游戏引擎", "BJ2330188", 31, 31)})
    course_id = created.json()["id"]
    payload = _register_bytes(4, "创意手作", "BJ2230005")
    blocked = client.post(
        f"/api/courses/{course_id}/grade-register",
        files={"file": ("reg.xlsx", payload, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert blocked.status_code == 409
    assert blocked.json()["detail"] == "登记表是《创意手作》（BJ2230005），当前课程是《2D游戏引擎》（BJ2330188），确定要导入吗？"
    assert client.get(f"/api/courses/{course_id}").json()["files"] == []
    allowed = client.post(
        f"/api/courses/{course_id}/grade-register",
        data={"confirm": "1"},
        files={"file": ("reg.xlsx", payload, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert allowed.status_code == 200, allowed.text
    assert allowed.json()["student_count"] == 4


def test_headcount_mismatch_warns_on_import_and_calculate(client):
    created = client.post("/api/courses", json={"name": "人数课", "settings": _named_settings("人数课", "CS101", 10, 9)})
    course_id = created.json()["id"]
    term_id = created.json()["current_term_id"]
    patched = client.patch(
        f"/api/courses/{course_id}/terms/{term_id}",
        json={"class_name": "24数字媒体", "school_year_term": "2024-2025学年第1学期"},
    )
    assert patched.status_code == 200, patched.text
    payload = _register_bytes(4, "人数课", "CS101")
    blocked = client.post(
        f"/api/courses/{course_id}/grade-register",
        files={"file": ("reg.xlsx", payload, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert blocked.status_code == 409
    assert "将替换该学期" in blocked.json()["detail"]
    imported = client.post(
        f"/api/courses/{course_id}/grade-register",
        data={"confirm": "1"},
        files={"file": ("reg.xlsx", payload, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert imported.status_code == 200, imported.text
    warnings = imported.json()["warnings"]
    assert any("上课人数是 10" in item and "4 人" in item for item in warnings)
    assert any("考核人数是 9" in item for item in warnings)
    assert imported.json()["conflicts"] == []
    course = client.get(f"/api/courses/{course_id}").json()
    assert course["current_term"]["student_count"] == 4
    assert course["current_term"]["exam_count"] == 4
    calculated = client.post(f"/api/courses/{course_id}/calculate")
    assert calculated.status_code == 200, calculated.text
    body = calculated.json()
    assert body["student_count"] == 4
    assert not any("人数不一致" in item for item in body["warnings"])


def test_import_does_not_copy_another_terms_exam_count(client):
    created = client.post(
        "/api/courses",
        json={"name": "2D游戏引擎", "settings": _named_settings("2D游戏引擎", "BJ2330188", 31, 31)},
    ).json()
    course_id = created["id"]
    small = client.post(
        f"/api/courses/{course_id}/grade-register",
        files={"file": ("b.xlsx", _register_bytes(8, "2D游戏引擎", "BJ2330188"), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert small.status_code == 200, small.text
    small_term = small.json()["term_id"]
    large = client.post(
        f"/api/courses/{course_id}/grade-register",
        files={
            "file": (
                "a.xlsx",
                _register_bytes(31, "2D游戏引擎", "BJ2330188", term="2025-2026学年第2学期"),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert large.status_code == 200, large.text
    assert large.json()["term_id"] != small_term
    assert large.json()["student_count"] == 31
    selected = client.post(f"/api/courses/{course_id}/terms/{small_term}/select")
    assert selected.status_code == 200, selected.text
    course = selected.json()
    assert course["current_term"]["exam_count"] == 8
    assert course["current_term"]["student_count"] == 8
    assert "31" not in " ".join(small.json()["warnings"] + small.json()["conflicts"])
    calculated = client.post(f"/api/courses/{course_id}/calculate")
    assert calculated.status_code == 200, calculated.text
    assert calculated.json()["student_count"] == 8
    assert not any("考核人数" in item for item in calculated.json()["warnings"])


def test_new_term_starts_empty_and_import_does_not_reuse_headcount(client):
    created = client.post(
        "/api/courses",
        json={"name": "2D游戏引擎", "settings": _named_settings("2D游戏引擎", "BJ2330188", 31, 31)},
    ).json()
    course_id = created["id"]
    first_id = created["current_term_id"]
    first = created["current_term"]
    assert first["student_count"] == 31
    assert first["exam_count"] == 31
    assert first["teacher"] == "王老师"
    assert first["class_name"] == "软工2201"
    stored = created["settings"]
    assert "student_count" not in stored
    basic = stored["course_basic_info"]
    assert basic.get("course_name") == "2D游戏引擎"
    for key in ("student_count", "exam_count", "class_name", "school_year_term", "teacher"):
        assert not basic.get(key)
    opened_info = stored["course_open_info"]
    for key in ("year_start", "year_end", "semester", "teacher"):
        assert not opened_info.get(key)

    opened = client.post(f"/api/courses/{course_id}/terms", json={})
    assert opened.status_code == 200, opened.text
    blank = opened.json()["current_term"]
    assert blank["id"] != first_id
    assert blank["student_count"] == 0
    assert blank["exam_count"] == 0
    assert blank["teacher"] == ""
    assert blank["class_name"] == ""
    assert blank["school_year_term"] == ""
    assert blank["year_start"] == ""
    assert blank["year_end"] == ""
    assert blank["semester"] == ""
    assert "student_count" not in opened.json()["settings"]

    imported = client.post(
        f"/api/courses/{course_id}/grade-register",
        files={"file": ("b.xlsx", _register_bytes(8, "2D游戏引擎", "BJ2330188"), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert imported.status_code == 200, imported.text
    body = imported.json()
    text = " ".join(body["warnings"] + body["conflicts"])
    assert "31" not in text
    assert "人数不一致" not in text
    current = client.get(f"/api/courses/{course_id}").json()
    assert current["current_term"]["student_count"] == 8
    assert "student_count" not in current["settings"]
    assert not current["settings"]["course_basic_info"].get("student_count")

    selected = client.post(f"/api/courses/{course_id}/terms/{first_id}/select")
    assert selected.status_code == 200, selected.text
    original = selected.json()["current_term"]
    assert original["student_count"] == 31
    assert original["exam_count"] == 31
    assert original["teacher"] == "王老师"
    assert original["class_name"] == "软工2201"


def test_courses_and_terms_keep_their_own_grade_counts(client):
    engine = client.post(
        "/api/courses",
        json={"name": "2D游戏引擎", "settings": _named_settings("2D游戏引擎", "BJ2330188", 0, 0)},
    ).json()
    craft = client.post(
        "/api/courses",
        json={"name": "创意手作", "settings": _named_settings("创意手作", "BJ2230005", 0, 0)},
    ).json()
    first = client.post(
        f"/api/courses/{engine['id']}/grade-register",
        files={"file": ("a.xlsx", _register_bytes(3, "2D游戏引擎", "BJ2330188"), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert first.status_code == 200, first.text
    second = client.post(
        f"/api/courses/{engine['id']}/grade-register",
        files={
            "file": (
                "b.xlsx",
                _register_bytes(2, "2D游戏引擎", "BJ2330188", term="2023-2024学年第2学期"),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert second.status_code == 200, second.text
    assert second.json()["term_id"] != first.json()["term_id"]
    other = client.post(
        f"/api/courses/{craft['id']}/grade-register",
        files={"file": ("c.xlsx", _register_bytes(5, "创意手作", "BJ2230005"), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert other.status_code == 200, other.text
    back = client.post(f"/api/courses/{engine['id']}/terms/{first.json()['term_id']}/select")
    assert back.status_code == 200
    first_calc = client.post(f"/api/courses/{engine['id']}/calculate")
    assert first_calc.status_code == 200, first_calc.text
    assert first_calc.json()["student_count"] == 3
    client.post(f"/api/courses/{engine['id']}/terms/{second.json()['term_id']}/select")
    second_calc = client.post(f"/api/courses/{engine['id']}/calculate")
    assert second_calc.status_code == 200, second_calc.text
    assert second_calc.json()["student_count"] == 2
    other_calc = client.post(f"/api/courses/{craft['id']}/calculate")
    assert other_calc.status_code == 200, other_calc.text
    assert other_calc.json()["student_count"] == 5
    again = client.post(f"/api/courses/{engine['id']}/terms/{first.json()['term_id']}/select")
    assert client.post(f"/api/courses/{engine['id']}/calculate").json()["student_count"] == 3
    assert again.status_code == 200


CRAFT_LINKS = [
    {
        "name": "平时考核",
        "ratio": 0.3,
        "methods": [
            {"name": "课堂参与", "supports": {"课程目标1": 0.2, "课程目标2": 0.2}, "subtotal": 0.4},
            {"name": "作业", "supports": {"课程目标1": 0.3, "课程目标2": 0.3}, "subtotal": 0.6},
        ],
    },
    {
        "name": "期末考核",
        "ratio": 0.7,
        "methods": [
            {"name": "考查（设计作品）", "supports": {"课程目标1": 0.4, "课程目标2": 0.1}, "subtotal": 0.5},
            {"name": "课程考核情况表", "supports": {"课程目标1": 0.1, "课程目标2": 0.4}, "subtotal": 0.5},
        ],
    },
]


def _craft_settings():
    settings = _settings(mode="forward")
    settings["relation_payload"] = {
        "objectives_count": 2,
        "links": CRAFT_LINKS,
        "objectives_total_weights": {"课程目标1": 0.5, "课程目标2": 0.5},
        "total_sum": 1.0,
    }
    settings["ratios"] = {"usual": 0.3, "midterm": 0.0, "final": 0.7}
    settings["course_basic_info"] = dict(settings["course_basic_info"])
    settings["course_open_info"] = dict(settings["course_open_info"])
    settings["course_basic_info"]["course_name"] = "创意手作"
    settings["course_open_info"]["course_name"] = "创意手作"
    settings["course_basic_info"]["course_code"] = "BJ2230005"
    settings["student_count"] = 0
    settings["course_basic_info"]["student_count"] = ""
    settings["course_basic_info"]["exam_count"] = ""
    return settings


def _register_xlsx(headers, students, course_name="创意手作", course_code="BJ2230005", term="2024-2025学年第1学期", teacher="章老师"):
    book = Workbook()
    sheet = book.active
    title = "广东第二师范学院课程成绩登记表" if not term else f"广东第二师范学院{term}课程成绩登记表"
    sheet["A1"] = title
    sheet["A2"] = f"课程名称：{course_name}                           课程代码：{course_code}                          课程性质：专业任选课"
    sheet["A3"] = f"开课学院：美术学院                           任课教师：{teacher}                               学分：3  考核方式：考查"
    for col, header in enumerate(headers, start=1):
        sheet.cell(5, col, header)
    for row_index, student in enumerate(students):
        for col, value in enumerate(student, start=1):
            if value is not None:
                sheet.cell(6 + row_index, col, value)
    count = len(students)
    sheet.cell(6 + count + 1, 1, "成绩统计")
    sheet.cell(6 + count + 2, 1, f"实考 {count}人 总人数 {count}人")
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_FORWARD_HEADERS = [
    "班级",
    "学号",
    "姓名",
    "课堂参与",
    "作业",
    "考查（设计作品）",
    "课程考核情况表",
    "平时(30%)",
    "期中(0%)",
    "期末(70%)",
    "总评",
]
_FORWARD_ROWS = [
    ["24数字媒体", "99000001", "学员甲", 80, 90, 70, 60, 85, 0, 75, 78],
    ["24数字媒体", "99000002", "学员乙", 100, 80, 90, 80, 88, 0, 82, 84],
]


def _post_register(client, course_id, payload, filename="reg.xlsx", mode="", term_id=None, confirm=False, **fields):
    url = f"/api/courses/{course_id}/grade-register"
    if term_id is not None:
        url = f"/api/courses/{course_id}/terms/{term_id}/grade-register"
    if mode:
        url += f"?mode={mode}"
    data = {key: value for key, value in fields.items() if value not in (None, "")}
    if confirm:
        data["confirm"] = "1"
    kwargs = {"files": {"file": (filename, payload, _XLSX)}}
    if data:
        kwargs["data"] = data
    return client.post(url, **kwargs)


def _grade_sheet(client, course_id):
    course = client.get(f"/api/courses/{course_id}").json()
    grades = [item for item in course["files"] if item["kind"] == "grade"]
    assert len(grades) == 1
    downloaded = client.get(f"/api/courses/{course_id}/files/{grades[0]['id']}")
    assert downloaded.status_code == 200
    return grades[0], load_workbook(io.BytesIO(downloaded.content)).active


def _hand_attainment(method_means):
    """与正向计算里分目标达成值同一公式。"""
    objectives = ["课程目标1", "课程目标2"]
    total_weight = 0.0
    total_actual = 0.0
    found = {}
    for index, key in enumerate(objectives):
        weight_sum = 0.0
        actual_sum = 0.0
        for link in CRAFT_LINKS:
            ratio = float(link["ratio"])
            support_sum = 0.0
            score_sum = 0.0
            for method in link["methods"]:
                weight = float(method["supports"][key])
                support_sum += weight
                score_sum += method_means[method["name"]] * weight
            weight_sum += ratio * 100.0 * support_sum
            actual_sum += ratio * score_sum
        found[f"课程目标{index + 1}"] = round(actual_sum / weight_sum, 3)
        total_weight += weight_sum
        total_actual += actual_sum
    found["总达成度"] = round(total_actual / total_weight, 3)
    return found


def _template_bytes(score_rows):
    book = Workbook()
    sheet = book.active
    sheet.cell(1, 1, "姓名")
    sheet.cell(2, 1, "姓名")
    methods = []
    for link in CRAFT_LINKS:
        for method in link["methods"]:
            methods.append((link["name"], method["name"]))
    for index, (link_name, method_name) in enumerate(methods, start=2):
        sheet.cell(1, index, link_name)
        sheet.cell(2, index, method_name)
    for row_index, (name, scores) in enumerate(score_rows, start=3):
        sheet.cell(row_index, 1, name)
        for col, score in enumerate(scores, start=2):
            sheet.cell(row_index, col, score)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def test_forward_register_matches_template_attainment(client):
    created = client.post("/api/courses", json={"name": "创意手作", "settings": _craft_settings()})
    assert created.status_code == 200, created.text
    course_id = created.json()["id"]
    term_id = created.json()["current_term_id"]
    link_only = _register_xlsx(
        ["班级", "学号", "姓名", "平时(30%)", "期中(0%)", "期末(70%)", "总评"],
        [["24数字媒体", "99000001", "学员甲", 80, 0, 90, 87]],
    )
    reverse_import = _post_register(client, course_id, link_only)
    assert reverse_import.status_code == 200, reverse_import.text
    assert reverse_import.json()["mode"] == "reverse"
    assert "课堂参与" in reverse_import.json()["reason"]
    assert "课程考核情况表" in reverse_import.json()["reason"]

    payload = _register_xlsx(_FORWARD_HEADERS, _FORWARD_ROWS)
    imported = _post_register(client, course_id, payload, term_id=term_id, confirm=True)
    assert imported.status_code == 200, imported.text
    body = imported.json()
    assert body["mode"] == "forward"
    assert body["student_count"] == 2
    assert body["detection"]["mode"] == "forward"
    for label in ("课堂参与", "作业", "设计作品", "课程考核情况表"):
        assert label in body["reason"]
    assert body["detection"]["matched"]["课堂参与"] == "课堂参与"
    assert body["detection"]["matched"]["作业"] == "作业"
    course = client.get(f"/api/courses/{course_id}").json()
    assert course["settings"]["mode"] == "forward"
    assert course["current_term"]["mode"] == "forward"
    grade, sheet = _grade_sheet(client, course_id)
    assert grade["original_name"] == "成绩登记表-正向.xlsx"
    assert sheet.cell(1, 1).value == "姓名"
    assert [sheet.cell(2, col).value for col in range(2, 6)] == [
        "课堂参与",
        "作业",
        "考查（设计作品）",
        "课程考核情况表",
    ]
    assert sheet.cell(3, 2).value == 80
    assert sheet.cell(4, 5).value == 80
    calculated = client.post(f"/api/courses/{course_id}/calculate")
    assert calculated.status_code == 200, calculated.text
    achievement = calculated.json()["achievement"]
    assert calculated.json()["student_count"] == 2
    assert calculated.json()["mode"] == "forward"
    means = {"课堂参与": 90, "作业": 85, "考查（设计作品）": 80, "课程考核情况表": 70}
    expected = _hand_attainment(means)
    for key, value in expected.items():
        assert achievement[key] == pytest.approx(value, abs=0.001)
    direct = run_calculation(
        _template_bytes([("学员甲", [80, 90, 70, 60]), ("学员乙", [100, 80, 90, 80])]),
        None,
        course["settings"],
    )
    for key, value in achievement.items():
        assert direct["achievement"][key] == pytest.approx(value, abs=1e-9)
    assert direct["student_count"] == 2


def test_overall_only_register_is_reverse(client):
    created = client.post("/api/courses", json={"name": "创意手作", "settings": _craft_settings()}).json()
    payload = _register_xlsx(
        ["班级", "学号", "姓名", "总评"],
        [["24数字媒体", "99000003", "学员丙", 86], ["24数字媒体", "99000004", "学员丁", 90]],
    )
    imported = _post_register(client, created["id"], payload)
    assert imported.status_code == 200, imported.text
    body = imported.json()
    assert body["mode"] == "reverse"
    assert "总评" in body["reason"]
    assert "同分" in body["reason"]
    _grade, sheet = _grade_sheet(client, created["id"])
    assert sheet.cell(1, 1).value == "姓名"
    assert sheet.cell(1, 2).value == "平时考核"
    assert sheet.cell(1, 3).value == "期末考核"
    assert sheet.cell(2, 2).value == 86
    assert sheet.cell(2, 3).value == 86
    assert sheet.cell(3, 2).value == 90
    calculated = client.post(f"/api/courses/{created['id']}/calculate")
    assert calculated.status_code == 200, calculated.text
    assert calculated.json()["mode"] == "reverse"
    assert calculated.json()["student_count"] == 2


def test_partial_columns_reverse_unless_forced(client):
    created = client.post("/api/courses", json={"name": "创意手作", "settings": _craft_settings()}).json()
    course_id = created["id"]
    payload = _register_xlsx(
        ["班级", "学号", "姓名", "课堂参与", "作业", "平时(30%)", "期中(0%)", "期末(70%)", "总评"],
        [["24数字媒体", "99000005", "学员戊", 70, 80, 75, 0, 88, 84]],
    )
    imported = _post_register(client, course_id, payload)
    assert imported.status_code == 200, imported.text
    body = imported.json()
    assert body["mode"] == "reverse"
    assert "考查（设计作品）" in body["reason"]
    assert "课程考核情况表" in body["reason"]
    forced = _post_register(client, course_id, payload, mode="forward", confirm=True)
    assert forced.status_code == 400
    assert "考查（设计作品）" in forced.json()["detail"]
    assert "课程考核情况表" in forced.json()["detail"]
    assert len([item for item in client.get(f"/api/courses/{course_id}").json()["files"] if item["kind"] == "grade"]) == 1
    reversed_import = _post_register(client, course_id, payload, mode="reverse", confirm=True)
    assert reversed_import.status_code == 200, reversed_import.text
    assert reversed_import.json()["mode"] == "reverse"
    _grade, sheet = _grade_sheet(client, course_id)
    assert sheet.cell(1, 2).value == "平时考核"
    assert sheet.cell(2, 2).value == 75
    assert sheet.cell(2, 3).value == 88


def test_single_method_link_uses_stage_column(client):
    settings = _craft_settings()
    settings["relation_payload"] = {
        "objectives_count": 2,
        "links": [
            {
                "name": "平时考核",
                "ratio": 0.3,
                "methods": [{"name": "课堂表现", "supports": {"课程目标1": 0.5, "课程目标2": 0.5}, "subtotal": 1}],
            },
            {
                "name": "期末考核",
                "ratio": 0.7,
                "methods": [{"name": "期末作品", "supports": {"课程目标1": 0.4, "课程目标2": 0.6}, "subtotal": 1}],
            },
        ],
    }
    created = client.post("/api/courses", json={"name": "创意手作", "settings": settings}).json()
    payload = _register_xlsx(
        ["班级", "学号", "姓名", "平时(30%)", "期中(0%)", "期末(70%)", "总评"],
        [
            ["24数字媒体", "99000006", "学员己", 82, 0, 91, 88],
            ["24数字媒体", "99000007", "学员庚", None, 0, 77, 70],
        ],
    )
    imported = _post_register(client, created["id"], payload)
    assert imported.status_code == 200, imported.text
    body = imported.json()
    assert body["mode"] == "forward"
    assert body["detection"]["matched"]["课堂表现"] == "平时"
    assert body["detection"]["matched"]["期末作品"] == "期末"
    assert any("平时有空值" in item for item in body["warnings"])
    _grade, sheet = _grade_sheet(client, created["id"])
    assert sheet.cell(2, 2).value == "课堂表现"
    assert sheet.cell(2, 3).value == "期末作品"
    assert sheet.cell(3, 2).value == 82
    assert sheet.cell(4, 2).value == 0
    calculated = client.post(f"/api/courses/{created['id']}/calculate")
    assert calculated.status_code == 200, calculated.text
    assert calculated.json()["student_count"] == 2
    assert calculated.json()["mode"] == "forward"


def test_column_cleaning_containment_and_ambiguity():
    cleaned = _register_xlsx(
        ["班级", "学号", "姓名", "作业（20％）", "平时(30%)", "期末(70%)"],
        [["24数字媒体", "99000011", "学员丙", 77, 50, 88]],
    )
    parsed = parse_register("清洗.xlsx", cleaned)
    assert parsed["extra_columns"][0]["label"] == "作业"
    detection = detect_register_mode(
        parsed,
        [{"name": "平时考核", "ratio": 1, "methods": [{"name": "作业", "subtotal": 1, "supports": {"课程目标1": 1}}]}],
    )
    assert detection["mode"] == "forward"
    assert detection["matched"]["作业"] == "作业"
    text = markdown_for_ai(parsed)
    assert "作业" in text
    assert "学员丙" not in text
    assert "99000011" not in text
    assert "77" not in text
    assert "50" not in text

    contained = _register_xlsx(
        ["班级", "学号", "姓名", "课堂参与", "作业", "设计作品", "课程考核情况表", "平时(30%)", "期末(70%)"],
        [["24数字媒体", "99000012", "学员壬", 80, 80, 80, 80, 80, 80]],
    )
    contained_parsed = parse_register("包含.xlsx", contained)
    contained_mode = detect_register_mode(contained_parsed, CRAFT_LINKS)
    assert contained_mode["mode"] == "forward"
    assert contained_mode["matched"]["考查（设计作品）"] == "设计作品"

    ambiguous = _register_xlsx(
        ["班级", "学号", "姓名", "设计作品", "作品展示"],
        [["24数字媒体", "99000013", "学员癸", 80, 90]],
    )
    ambiguous_mode = detect_register_mode(
        parse_register("歧义.xlsx", ambiguous),
        [{"name": "期末考核", "ratio": 1, "methods": [{"name": "作品", "subtotal": 1, "supports": {"课程目标1": 1}}]}],
    )
    assert ambiguous_mode["mode"] == "reverse"
    assert "无法唯一匹配" in ambiguous_mode["reason"]
    assert "作品" in ambiguous_mode["reason"]


def test_from_syllabus_forward_register_can_calculate(client):
    grid = [
        ["考核环节", "占比", "考核方式", "课程目标1", "课程目标2", "小计"],
        ["平时考核", "30%", "课堂参与", "20%", "20%", "40%"],
        ["", "", "作业", "30%", "30%", "60%"],
        ["期末考核", "70%", "考查（设计作品）", "40%", "10%", "50%"],
        ["", "", "课程考核情况表", "10%", "40%", "50%"],
    ]
    payload = _register_xlsx(_FORWARD_HEADERS, _FORWARD_ROWS)
    created = client.post(
        "/api/courses/from-syllabus",
        data={
            "payload": json.dumps(
                {
                    "fields": {"course_name": "创意手作", "course_code": "BJ2230005"},
                    "objectives": ["能完成作品", "能分析作品"],
                    "description": "简介",
                    "grad_req_map": [
                        {"requirement": "毕业要求1", "indicator": "指标1", "strength": "H"},
                        {"requirement": "毕业要求2", "indicator": "指标2", "strength": "M"},
                    ],
                    "relation_grid": grid,
                },
                ensure_ascii=False,
            )
        },
        files={"register": ("正向登记.xlsx", payload, _XLSX)},
    )
    assert created.status_code == 200, created.text
    body = created.json()
    assert body["grade_import"]["mode"] == "forward"
    assert body["current_term"]["mode"] == "forward"
    assert "课堂参与" in body["grade_import"]["reason"]
    calculated = client.post(f"/api/courses/{body['id']}/calculate")
    assert calculated.status_code == 200, calculated.text
    assert calculated.json()["mode"] == "forward"
    assert calculated.json()["student_count"] == 2
    assert calculated.json()["term_id"] == body["current_term_id"]


def test_calculate_returns_the_current_term_result(client):
    created = client.post("/api/courses", json={"name": "创意手作", "settings": _craft_settings()})
    assert created.status_code == 200, created.text
    course_id = created.json()["id"]
    first_term = created.json()["current_term_id"]
    link_only = _register_xlsx(
        ["班级", "学号", "姓名", "平时(30%)", "期中(0%)", "期末(70%)", "总评"],
        [["24数字媒体", "99000001", "学员甲", 60, 0, 70, 67]],
    )
    reverse_import = _post_register(client, course_id, link_only)
    assert reverse_import.status_code == 200, reverse_import.text
    assert reverse_import.json()["mode"] == "reverse"
    imported_term = reverse_import.json()["term_id"]
    assert imported_term != first_term
    first = client.post(f"/api/courses/{course_id}/calculate")
    assert first.status_code == 200, first.text
    first_body = first.json()
    assert first_body["term_id"] == imported_term
    assert first_body["mode"] == "reverse"
    assert first_body["student_count"] == 1
    assert "总达成度" in first_body["achievement"]

    forward_import = _post_register(
        client,
        course_id,
        _register_xlsx(_FORWARD_HEADERS, _FORWARD_ROWS, term="2025-2026学年第2学期"),
    )
    assert forward_import.status_code == 200, forward_import.text
    assert forward_import.json()["mode"] == "forward"
    second_term = forward_import.json()["term_id"]
    assert second_term != imported_term
    second = client.post(f"/api/courses/{course_id}/calculate")
    assert second.status_code == 200, second.text
    second_body = second.json()
    assert second_body["term_id"] == second_term
    assert second_body["mode"] == "forward"
    assert second_body["student_count"] == 2
    assert second_body["achievement"]["总达成度"] != first_body["achievement"]["总达成度"]

    back = client.post(f"/api/courses/{course_id}/terms/{imported_term}/select")
    assert back.status_code == 200, back.text
    assert back.json()["current_term"]["mode"] == "reverse"
    again = client.post(f"/api/courses/{course_id}/calculate")
    assert again.status_code == 200, again.text
    again_body = again.json()
    assert again_body["term_id"] == imported_term
    assert again_body["mode"] == "reverse"
    assert again_body["student_count"] == 1
    assert again_body["achievement"] == first_body["achievement"]


def _assert_no_term_identity(settings):
    basic = settings.get("course_basic_info") or {}
    opened = settings.get("course_open_info") or {}
    for key in ("teacher", "major", "class_name", "school_year_term", "student_count", "exam_count"):
        assert not basic.get(key)
    for key in ("year_start", "year_end", "semester", "term", "teacher"):
        assert not opened.get(key)
    assert "student_count" not in settings


def test_register_creates_replaces_and_adds_another_term(client, monkeypatch):
    def fail_ai(*_args, **_kwargs):
        raise AssertionError("导入成绩不得调用模型")

    monkeypatch.setattr("web_app.syllabus.client.complete", fail_ai)
    grid = [
        ["考核环节", "占比", "考核方式", "课程目标1", "课程目标2", "小计"],
        ["平时考核", "30%", "课堂参与", "20%", "20%", "40%"],
        ["", "", "作业", "30%", "30%", "60%"],
        ["期末考核", "70%", "考查（设计作品）", "40%", "10%", "50%"],
        ["", "", "课程考核情况表", "10%", "40%", "50%"],
    ]
    created = client.post(
        "/api/courses/from-syllabus",
        json={
            "fields": {
                "course_name": "创意手作",
                "course_code": "BJ2230005",
                "teacher": "不该留下",
                "major": "不该留下",
                "class_name": "不该留下",
                "year_start": "1999",
                "year_end": "2000",
                "semester": "2",
                "student_count": "99",
            },
            "objectives": ["能完成作品", "能分析作品"],
            "description": "简介",
            "grad_req_map": [
                {"requirement": "毕业要求1", "indicator": "指标1", "strength": "H"},
                {"requirement": "毕业要求2", "indicator": "指标2", "strength": "M"},
            ],
            "relation_grid": grid,
        },
    )
    assert created.status_code == 200, created.text
    course_id = created.json()["id"]
    assert created.json()["current_term"] is None
    assert created.json()["terms"] == []
    _assert_no_term_identity(created.json()["settings"])

    payload = _register_xlsx(_FORWARD_HEADERS, _FORWARD_ROWS, teacher="章老师")
    imported = _post_register(client, course_id, payload)
    assert imported.status_code == 200, imported.text
    assert imported.json()["mode"] == "forward"
    assert imported.json()["created_term"] is True
    posted = json.dumps(imported.json(), ensure_ascii=False)
    assert "学员甲" not in posted
    assert "99000001" not in posted
    course = client.get(f"/api/courses/{course_id}").json()
    term = course["current_term"]
    assert term["teacher"] == "章老师"
    assert term["class_name"] == "24数字媒体"
    assert term["major"] == "数字媒体"
    assert term["year_start"] == "2024"
    assert term["year_end"] == "2025"
    assert term["semester"] == "1"
    assert term["school_year_term"] == "2024-2025学年第1学期"
    assert term["student_count"] == 2
    _assert_no_term_identity(course["settings"])

    changed = _register_xlsx(_FORWARD_HEADERS, _FORWARD_ROWS, teacher="李老师")
    blocked = _post_register(client, course_id, changed)
    assert blocked.status_code == 409
    assert blocked.json()["code"] == "replace"
    assert blocked.json()["detail"] == "将替换该学期（2024-2025学年第1学期 · 24数字媒体）的成绩，确定吗？"
    replaced = _post_register(client, course_id, changed, confirm=True)
    assert replaced.status_code == 200, replaced.text
    assert replaced.json()["conflicts"] == []
    assert replaced.json()["created_term"] is False
    assert client.get(f"/api/courses/{course_id}").json()["current_term"]["teacher"] == "李老师"

    score = 82.5
    forward_rows = [
        ["24数字媒体", "99000021", "学员甲", score, score, score, score, score, 0, score, score],
        ["24数字媒体", "99000022", "学员乙", score, score, score, score, score, 0, score, score],
    ]
    other = _register_xlsx(_FORWARD_HEADERS, forward_rows, term="2025-2026学年第2学期", teacher="周老师")
    second = _post_register(client, course_id, other)
    assert second.status_code == 200, second.text
    assert second.json()["mode"] == "forward"
    assert second.json()["created_term"] is True
    assert second.json()["term_id"] != term["id"]
    calculated = client.post(f"/api/courses/{course_id}/calculate")
    assert calculated.status_code == 200, calculated.text
    assert calculated.json()["mode"] == "forward"
    assert calculated.json()["achievement"]["总达成度"] == pytest.approx(0.825, abs=0.001)
    assert "99000021" not in json.dumps(calculated.json(), ensure_ascii=False)

    bare = _register_xlsx(
        ["班级", "学号", "姓名", "平时(30%)", "期末(70%)", "总评"],
        [["22数字媒体", "99000031", "学员丙", 80, 90, 88]],
        term="",
    )
    missing = _post_register(client, course_id, bare)
    assert missing.status_code == 409
    assert missing.json()["code"] == "need_input"
    assert "term" in missing.json()["needs"]
    filled = _post_register(client, course_id, bare, year_start="2022", year_end="2023", semester="1")
    assert filled.status_code == 200, filled.text
    stored = client.get(f"/api/courses/{course_id}").json()["current_term"]
    assert stored["year_start"] == "2022"
    assert stored["year_end"] == "2023"
    assert stored["semester"] == "1"
    assert stored["class_name"] == "22数字媒体"
    assert len(client.get(f"/api/courses/{course_id}").json()["terms"]) == 3


def test_multi_class_import_requires_one_class(client):
    created = client.post(
        "/api/courses",
        json={"name": "创意手作", "settings": _named_settings("创意手作", "BJ2230005", 0, 0)},
    ).json()
    buffer = io.BytesIO()
    _xlsx_like_handcraft(buffer, 35)
    payload = buffer.getvalue()
    blocked = _post_register(client, created["id"], payload)
    assert blocked.status_code == 409
    body = blocked.json()
    assert body["code"] == "need_input"
    assert "class" in body["needs"]
    assert set(body["classes"]) == {"23美术教育A", "23数字媒体"}
    chosen = _post_register(client, created["id"], payload, class_name="23数字媒体")
    assert chosen.status_code == 200, chosen.text
    assert chosen.json()["student_count"] == 15
    course = client.get(f"/api/courses/{created['id']}").json()
    assert course["current_term"]["class_name"] == "23数字媒体"
    assert course["current_term"]["major"] == "数字媒体"
    assert course["current_term"]["student_count"] == 15


def test_opening_fields_save_onto_the_current_term_only(client):
    grid = [
        ["考核环节", "占比", "考核方式", "课程目标1", "小计"],
        ["平时考核", "100%", "作业", "100%", "100%"],
    ]
    created = client.post(
        "/api/courses/from-syllabus",
        json={
            "fields": {"course_name": "创意手作", "course_code": "BJ2230005", "college": "美术学院"},
            "objectives": ["能完成作品"],
            "description": "简介",
            "grad_req_map": [{"requirement": "", "indicator": "理论知识", "strength": "H"}],
            "relation_grid": grid,
        },
    )
    assert created.status_code == 200, created.text
    course_id = created.json()["id"]
    untouched = client.patch(
        f"/api/courses/{course_id}",
        json={"settings": {"course_description": "只改简介", "course_basic_info": {"course_name": "创意手作"}, "course_open_info": {"course_name": "创意手作", "department": "美术学院"}}},
    )
    assert untouched.status_code == 200, untouched.text
    assert untouched.json()["terms"] == []

    first = _post_register(
        client,
        course_id,
        _register_xlsx(
            ["班级", "学号", "姓名", "平时(100%)", "总评"],
            [["24数字媒体", "99000041", "学员甲", 80, 80]],
            term="2024-2025学年第1学期",
        ),
    )
    assert first.status_code == 200, first.text
    second = _post_register(
        client,
        course_id,
        _register_xlsx(
            ["班级", "学号", "姓名", "平时(100%)", "总评"],
            [["24数字媒体", "99000042", "学员乙", 90, 90]],
            term="2025-2026学年第2学期",
            teacher="周老师",
        ),
    )
    assert second.status_code == 200, second.text
    assert second.json()["term_id"] != first.json()["term_id"]
    current = client.get(f"/api/courses/{course_id}").json()
    settings = current["settings"]
    settings["course_open_info"] = {
        "course_name": "创意手作",
        "department": "设计学院",
        "year_start": "2025",
        "year_end": "2026",
        "semester": "2",
        "teacher": "新老师",
    }
    basic = dict(settings["course_basic_info"])
    basic.update(
        {
            "teacher": "新老师",
            "class_name": "25数字媒体",
            "major": "数字媒体艺术",
            "school_year_term": "旧标签",
            "student_count": "12",
            "exam_count": "11",
        }
    )
    settings["course_basic_info"] = basic
    settings["student_count"] = 12
    saved = client.patch(f"/api/courses/{course_id}", json={"settings": settings})
    assert saved.status_code == 200, saved.text
    body = saved.json()
    term = body["current_term"]
    assert term["id"] == second.json()["term_id"]
    assert term["teacher"] == "新老师"
    assert term["class_name"] == "25数字媒体"
    assert term["major"] == "数字媒体艺术"
    assert term["year_start"] == "2025"
    assert term["year_end"] == "2026"
    assert term["semester"] == "2"
    assert term["school_year_term"] == "2025-2026学年第2学期"
    assert term["student_count"] == 12
    assert term["exam_count"] == 11
    assert body["settings"]["course_open_info"]["department"] == "设计学院"
    assert body["settings"]["course_open_info"]["course_name"] == "创意手作"
    _assert_no_term_identity(body["settings"])
    other = next(item for item in body["terms"] if item["id"] == first.json()["term_id"])
    assert other["teacher"] == "章老师"
    assert other["class_name"] == "24数字媒体"
    assert other["year_start"] == "2024"
    assert other["semester"] == "1"
    assert other["student_count"] == 1
