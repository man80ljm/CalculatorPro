"""资料按学期和生成批次下载，验证版本不会串档、历史包不会随新输入改变。"""
import io
import zipfile
from urllib.parse import unquote

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select

from tests.test_web_flow import app, client, PASSWORD, _course, _settings, _fill_forward_template
from web_app.db import CourseFile, FileBatch, session_scope
from web_app.download_names import archive_filename
from web_app.storage import save_blob


@pytest.mark.parametrize("args, expected", [
    (("2D游戏引擎", "2026", "2027", "1"), "2D游戏引擎_2026-2027学年_第1学期.zip"),
    (("测试课程", "2025", "2026", "2"), "测试课程_2025-2026学年_第2学期.zip"),
    (("测试课程", "", "", "", "2025-2026学年第2学期"), "测试课程_2025-2026学年_第2学期.zip"),
    (("测试课程", "", "", "", "2025-2026-1"), "测试课程_2025-2026学年_第1学期.zip"),
    (("测试课程", "", "", ""), "测试课程_未填写学期.zip"),
    (("课程/名称:*", "../2026", "2027", "1"), "课程_名称___未填写学期.zip"),
])
def test_short_archive_names(args, expected):
    assert archive_filename(*args) == expected


def _prepare(client):
    course_id = _course(client, _settings(), name="合成资料测试")
    response = client.post(f"/api/courses/{course_id}/template")
    assert response.status_code == 200
    grades = _fill_forward_template(response.content)
    _upload(client, course_id, grades)
    return course_id, grades


def _upload(client, course_id, content):
    response = client.post(f"/api/courses/{course_id}/files", data={"kind": "grade"},
                           files={"file": ("合成成绩.xlsx", content)})
    assert response.status_code == 200, response.text


def _catalog(client, course_id):
    response = client.get(f"/api/courses/{course_id}/materials")
    assert response.status_code == 200, response.text
    return response.json()


def test_new_versions_keep_inputs_and_output_bytes(client):
    course_id, original = _prepare(client)
    first = client.post(f"/api/courses/{course_id}/export")
    assert first.status_code == 200, first.text
    catalog = _catalog(client, course_id)
    term = catalog["terms"][0]
    version = term["versions"][0]
    assert len(term["versions"]) == 1
    assert version["legacy"] is False
    assert "合成资料测试_2024-2025学年_第1学期.zip" in unquote(first.headers["content-disposition"])
    archive = zipfile.ZipFile(io.BytesIO(first.content))
    assert archive.read("导入资料/本学期成绩表.xlsx") == original
    assert not any(name.endswith(".zip") for name in archive.namelist())
    for file in version["files"]:
        response = client.get(f"/api/courses/{course_id}/files/{file['id']}")
        assert response.status_code == 200
        assert archive.read(file["original_name"]) == response.content

    # 更换输入后生成第二版；第一版 ZIP 及其单个文件保持不变。
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(original))
    book.active.cell(3, 3, 60)
    changed = io.BytesIO()
    book.save(changed)
    _upload(client, course_id, changed.getvalue())
    second = client.post(f"/api/courses/{course_id}/export")
    assert second.status_code == 200, second.text
    versions = _catalog(client, course_id)["terms"][0]["versions"]
    assert len(versions) == 2
    assert {file["id"] for file in versions[0]["files"]}.isdisjoint(file["id"] for file in versions[1]["files"])
    latest = zipfile.ZipFile(io.BytesIO(second.content))
    assert latest.read("导入资料/本学期成绩表.xlsx") == changed.getvalue()
    old_download = client.get(f"/api/courses/{course_id}/files/{version['archive']['id']}")
    assert old_download.content == first.content


def test_materials_browsing_and_downloading_do_not_select_term(client):
    course_id, grades = _prepare(client)
    assert client.post(f"/api/courses/{course_id}/calculate").status_code == 200
    old_term = _catalog(client, course_id)["terms"][0]
    new = client.post(f"/api/courses/{course_id}/terms", json={
        "year_start": "2026", "year_end": "2027", "semester": "1", "class_name": "合成新班",
        "student_count": 2, "exam_count": 2,
    })
    assert new.status_code == 200, new.text
    new_term = new.json()["current_term_id"]
    _upload(client, course_id, grades)
    assert client.post(f"/api/courses/{course_id}/export").status_code == 200
    terms = _catalog(client, course_id)["terms"]
    assert [term["id"] for term in terms] == [new_term, old_term["id"]]
    for term in terms:
        for version in term["versions"]:
            assert all(file["term_id"] == term["id"] for file in version["files"])
            response = client.get(f"/api/courses/{course_id}/files/{version['archive']['id']}")
            assert response.status_code == 200
            assert term["download_name"] in unquote(response.headers["content-disposition"])
    assert client.get(f"/api/courses/{course_id}").json()["current_term_id"] == new_term
    assert client.delete(f"/api/courses/{course_id}/terms/{old_term['id']}?confirm=1").status_code == 200
    assert len(_catalog(client, course_id)["terms"]) == 1
    with session_scope() as db:
        assert not list(db.scalars(select(FileBatch).where(FileBatch.term_id == old_term["id"])))


def test_legacy_packages_and_single_files_remain_downloadable(client):
    course_id, _ = _prepare(client)
    current = client.get(f"/api/courses/{course_id}").json()["current_term_id"]
    archive_bytes = io.BytesIO()
    with zipfile.ZipFile(archive_bytes, "w") as archive:
        archive.writestr("旧报告.docx", b"old document")
    user_id = client.get("/api/me").json()["id"]
    with session_scope() as db:
        single = save_blob(db, user_id, course_id, "旧报告.docx", b"old document", "report", current)
        package = save_blob(db, user_id, course_id, "旧名称AI分析报告.zip", archive_bytes.getvalue(), "report", current)
    term = _catalog(client, course_id)["terms"][0]
    assert term["versions"][0]["legacy"] is True
    assert term["other_files"][0]["id"] == single.id
    response = client.get(f"/api/courses/{course_id}/files/{package.id}")
    assert response.content == archive_bytes.getvalue()
    assert term["download_name"] in unquote(response.headers["content-disposition"])


def test_materials_are_private_and_course_deletion_cleans_batches(client, app):
    course_id, _ = _prepare(client)
    assert client.post(f"/api/courses/{course_id}/export").status_code == 200
    version = _catalog(client, course_id)["terms"][0]["versions"][0]
    with TestClient(app) as other:
        assert other.post("/api/register", json={"username": "other-user", "password": PASSWORD}).status_code == 200
        assert other.post("/api/login", json={"username": "other-user", "password": PASSWORD}).status_code == 200
        assert other.get(f"/api/courses/{course_id}/materials").status_code == 404
        assert other.get(f"/api/courses/{course_id}/files/{version['archive']['id']}").status_code == 404
    assert client.post("/api/logout").status_code == 200
    assert client.post("/api/login", json={"username": "teacher", "password": PASSWORD}).status_code == 200
    assert _catalog(client, course_id)["terms"][0]["versions"][0] == version
    assert client.delete(f"/api/courses/{course_id}").status_code == 200
    with session_scope() as db:
        assert not list(db.scalars(select(FileBatch).where(FileBatch.course_id == course_id)))
        assert not list(db.scalars(select(CourseFile).where(CourseFile.course_id == course_id)))
