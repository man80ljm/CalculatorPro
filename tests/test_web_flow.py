"""网页版全流程：模板、上传、计算、导出，AI 调用打桩。"""
import io
import json
import os
import sys
import zipfile
from pathlib import Path
from urllib.parse import unquote

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

os.environ.setdefault("APP_PASSWORD", "test-pass")

from web_app.app import create_app  # noqa: E402
from web_app.limiter import ai_limiter, login_limiter  # noqa: E402


SAMPLE_RELATION = {
    "objectives_count": 2,
    "links": [
        {
            "name": "平时考核",
            "ratio": 0.3,
            "methods": [
                {
                    "name": "平时作业",
                    "supports": {"课程目标1": 0.5, "课程目标2": 0.5},
                    "subtotal": 1.0,
                }
            ],
        },
        {
            "name": "期末考核",
            "ratio": 0.7,
            "methods": [
                {
                    "name": "期末考试",
                    "supports": {"课程目标1": 0.4, "课程目标2": 0.6},
                    "subtotal": 1.0,
                }
            ],
        },
    ],
    "objectives_total_weights": {"课程目标1": 0.43, "课程目标2": 0.57},
    "total_sum": 1.0,
}


def _settings(mode="forward", **extra):
    data = {
        "mode": mode,
        "student_count": 3,
        "course_open_info": {
            "year_start": "2024",
            "year_end": "2025",
            "semester": "1",
            "term": "1",
            "course_name": "测试课程",
            "department": "计算机学院",
            "teacher": "王老师",
        },
        "course_basic_info": {
            "course_name": "测试课程",
            "credits": "3",
            "hours": "48",
            "course_type": "专业课",
            "course_code": "CS101",
            "school_year_term": "2024-2025-1",
            "college": "计算机学院",
            "teacher": "王老师",
            "major": "软件工程",
            "class_name": "软工2201",
            "student_count": "2",
            "exam_count": "2",
        },
        "ratios": {"usual": 0.3, "midterm": 0, "final": 0.7},
        "relation_payload": SAMPLE_RELATION,
        "grad_req_map": [
            {"objective": "课程目标1", "requirement": "毕业要求1", "indicator": "指标点1.1"},
            {"objective": "课程目标2", "requirement": "毕业要求2", "indicator": "指标点2.2"},
        ],
        "course_description": "这是一门用于测试的程序设计课程。",
        "objective_requirements": ["能完成基本程序设计", "能分析简单问题"],
        "spread_mode": "中跨度（7-13分）",
        "distribution": "标准正态",
        "noise_config": None,
        "report_style": "专业",
        "word_limit": 80,
    }
    data.update(extra)
    return data


@pytest.fixture
def client():
    login_limiter.reset()
    ai_limiter.reset()
    app = create_app()
    with TestClient(app) as test_client:
        response = test_client.post("/api/login", json={"password": "test-pass"})
        assert response.status_code == 200, response.text
        assert "cp_session" in response.cookies
        yield test_client


def test_healthz_is_public_and_pages_require_login():
    app = create_app()
    with TestClient(app) as anonymous:
        assert anonymous.get("/healthz").json() == {"status": "ok"}
        page = anonymous.get("/", follow_redirects=False)
        assert page.status_code == 303
        assert page.headers["location"] == "/login"
        assert anonymous.get("/api/calculate").status_code == 401
        assert anonymous.get("/static/app.js", follow_redirects=False).status_code == 303
        login = anonymous.get("/login")
        assert login.status_code == 200
        assert "访问密码" in login.text
        assert "api_key" not in login.text.lower()
        assert "deepseek" not in login.text.lower()


def test_wrong_password_and_missing_password(monkeypatch):
    login_limiter.reset()
    app = create_app()
    with TestClient(app) as anonymous:
        bad = anonymous.post("/api/login", json={"password": "nope"})
        assert bad.status_code == 401
    monkeypatch.setenv("APP_PASSWORD", "")
    with TestClient(create_app()) as anonymous:
        blocked = anonymous.post("/api/login", json={"password": "anything"})
        assert blocked.status_code == 503
        assert "APP_PASSWORD" in blocked.json()["detail"]


def test_web_stack_does_not_import_pyqt():
    assert "PyQt6" not in sys.modules
    import web_app.service  # noqa: F401

    assert "PyQt6" not in sys.modules
    assert "PyQt6.QtWidgets" not in sys.modules


def _cell_value(sheet, row, col):
    value = sheet.cell(row, col).value
    if value is not None:
        return value
    for merged in sheet.merged_cells.ranges:
        if merged.min_row <= row <= merged.max_row and merged.min_col <= col <= merged.max_col:
            return sheet.cell(merged.min_row, merged.min_col).value
    return None


def _fill_forward_template(template_bytes: bytes) -> bytes:
    workbook = load_workbook(io.BytesIO(template_bytes))
    sheet = workbook.active
    headers = [_cell_value(sheet, 2, col) for col in range(1, sheet.max_column + 1)]
    assert headers[0] == "姓名"
    assert "平时作业" in headers
    assert "期末考试" in headers
    usual_col = headers.index("平时作业") + 1
    final_col = headers.index("期末考试") + 1
    rows = [("张三", 86, 92), ("李四", 74, 63)]
    for offset, (name, usual, final) in enumerate(rows):
        row = 3 + offset
        sheet.cell(row, 1, name)
        sheet.cell(row, usual_col, usual)
        sheet.cell(row, final_col, final)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _post_file(client, path, settings, excel_bytes, filename="grades.xlsx"):
    return client.post(
        path,
        data={"settings": json.dumps(settings, ensure_ascii=False)},
        files={"file": (filename, excel_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )


def test_forward_template_calculate_export_and_mocked_ai(client, monkeypatch):
    outputs = ROOT / "outputs"
    before = {path.name for path in outputs.iterdir()} if outputs.exists() else set()
    template = client.post(
        "/api/template",
        data={"settings": json.dumps(_settings(), ensure_ascii=False)},
    )
    assert template.status_code == 200, template.text
    assert "正向成绩模板" in unquote(template.headers["content-disposition"])
    excel_bytes = _fill_forward_template(template.content)

    calculated = _post_file(client, "/api/calculate", _settings(), excel_bytes)
    assert calculated.status_code == 200, calculated.text
    body = calculated.json()
    assert body["course_name"] == "测试课程"
    assert body["student_count"] == 2
    assert body["average_score"] > 0
    assert "总达成度" in body["achievement"]
    assert any(name.endswith(".xlsx") for name in body["files"])
    assert any("课程成绩统计表" in name for name in body["files"])
    assert any("评价结果" in name for name in body["files"])

    exported = _post_file(client, "/api/export", _settings(), excel_bytes)
    assert exported.status_code == 200, exported.text
    assert exported.headers["content-type"].startswith("application/zip")
    archive = zipfile.ZipFile(io.BytesIO(exported.content))
    names = archive.namelist()
    assert any(name.endswith(".xlsx") for name in names)
    assert any(name.endswith("2.课程成绩统计表.docx") for name in names)
    assert any(name.endswith("5.基于考核结果的课程目标达成情况评价结果表.docx") for name in names)
    assert any(name.endswith("1.课程基本信息表.docx") for name in names)
    assert any(name.endswith("4.课程考核与课程目标对应关系表.docx") for name in names)
    detail_name = next(name for name in names if name.endswith("成绩明细.xlsx"))
    detail = load_workbook(io.BytesIO(archive.read(detail_name)))
    assert detail.active["A2"].value == "张三"
    after = {path.name for path in outputs.iterdir()} if outputs.exists() else set()
    assert after == before

    monkeypatch.setenv("DEEPSEEK_API_KEY", "")
    missing = _post_file(client, "/api/ai-report", _settings(), excel_bytes)
    assert missing.status_code == 400
    assert "DEEPSEEK_API_KEY" in missing.json()["detail"]
    status = client.get("/api/ai-status").json()
    assert status["enabled"] is False

    calls = []

    class _Response:
        text = ""

        def raise_for_status(self):
            return None

        def json(self):
            return {
                "choices": [
                    {
                        "message": {
                            "content": "本轮课程目标达成情况总体平稳，考核覆盖了主要能力点，后续将补充针对性练习。"
                        }
                    }
                ]
            }

    def fake_post(url, headers=None, json=None, timeout=None):
        calls.append({"url": url, "headers": headers, "json": json, "timeout": timeout})
        return _Response()

    monkeypatch.setenv("DEEPSEEK_API_KEY", "sk-test-secret")
    monkeypatch.setattr("core_app.ai_report.requests.post", fake_post)
    report = _post_file(client, "/api/ai-report", _settings(), excel_bytes)
    assert report.status_code == 200, report.text
    assert calls, "AI 接口没有被调用"
    assert "api.deepseek.com" in calls[0]["url"]
    assert calls[0]["json"]["model"] == "deepseek-chat"
    assert calls[0]["headers"]["Authorization"] == "Bearer sk-test-secret"
    report_zip = zipfile.ZipFile(io.BytesIO(report.content))
    report_names = report_zip.namelist()
    assert any(name.startswith("6.课程目标达成情况分析") for name in report_names)
    assert any("达成度报告" in name for name in report_names)
    assert b"sk-test-secret" not in report.content
    enabled = client.get("/api/ai-status").json()
    assert enabled["enabled"] is True
    assert "sk-test-secret" not in enabled["message"]


def test_reverse_export(client):
    template = client.post(
        "/api/template",
        data={"settings": json.dumps(_settings(mode="reverse"), ensure_ascii=False)},
    )
    assert template.status_code == 200, template.text
    workbook = load_workbook(io.BytesIO(template.content))
    sheet = workbook.active
    headers = [sheet.cell(1, col).value for col in range(1, sheet.max_column + 1)]
    assert headers[0] == "姓名"
    sheet.cell(2, 1, "赵六")
    sheet.cell(2, headers.index("平时考核") + 1, 80)
    sheet.cell(2, headers.index("期末考核") + 1, 88)
    sheet.cell(3, 1, "钱七")
    sheet.cell(3, headers.index("平时考核") + 1, 70)
    sheet.cell(3, headers.index("期末考核") + 1, 76)
    buffer = io.BytesIO()
    workbook.save(buffer)
    exported = _post_file(client, "/api/export", _settings(mode="reverse"), buffer.getvalue())
    assert exported.status_code == 200, exported.text
    names = zipfile.ZipFile(io.BytesIO(exported.content)).namelist()
    assert any("逆向" in name and name.endswith(".xlsx") for name in names)
    assert any(name.endswith(".docx") for name in names)


def test_outputs_override_is_isolated():
    import tempfile
    import threading
    import time

    from utils import get_outputs_dir, override_outputs_dir

    results = {}

    def worker(index):
        with tempfile.TemporaryDirectory() as directory:
            with override_outputs_dir(directory):
                results[index] = get_outputs_dir()
                time.sleep(0.05)
                assert get_outputs_dir() == results[index]

    threads = [threading.Thread(target=worker, args=(index,)) for index in range(8)]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join()
    assert len(set(results.values())) == 8


def test_login_rate_limit():
    login_limiter.reset()
    with TestClient(create_app()) as anonymous:
        statuses = [
            anonymous.post("/api/login", json={"password": "bad"}).status_code
            for _ in range(9)
        ]
    login_limiter.reset()
    assert statuses[:8] == [401] * 8
    assert statuses[8] == 429


def test_upload_limit(client):
    huge = b"PK" + b"x" * (10 * 1024 * 1024)
    response = _post_file(client, "/api/calculate", _settings(), huge)
    assert response.status_code == 400
    assert "10MB" in response.json()["detail"]
