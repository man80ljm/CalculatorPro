"""Calendar/lesson metadata use timetable data, with no grade-register fallback."""
import copy
import io
import json
import uuid
from datetime import timedelta
from unittest.mock import Mock

import docx
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select
from sqlalchemy.orm import Session

from tests.test_web_flow import app, client, PASSWORD
from tests.test_course_calendar import create_course, calendar_data, cached_template, converted_template, workbook_from_download
from tests.test_lesson_plans import syllabus, wait_job
from web_app.db import Course, ScheduleDocument, ScheduleRecord, Term, TeachingAssignmentSelection, get_engine, utcnow
from web_app.lesson_plans import engine, service
from web_app.teaching import course_assignments as assignments


def add_schedule(name, rows, *, semester=None, start=None, newer=0):
    period=assignments.academic_period()
    semester=str(semester or period["semester"]); start=str(start or period["year_start"])
    identity=uuid.uuid4().hex
    with Session(get_engine()) as db:
        db.add(ScheduleDocument(id=identity,filename="合成课表.pdf",digest=uuid.uuid4().hex,
            model="offline",term=json.dumps({"academicYear":f"{start}-{int(start)+1}","semester":int(semester)}),
            created_at=utcnow()+timedelta(minutes=newer)))
        db.flush()
        for teacher,classes,*issues in rows:
            db.add(ScheduleRecord(id=uuid.uuid4().hex,document_id=identity,term_id=f"{start}-{semester}",course_name=name,
                payload=json.dumps({"teachers":[teacher],"classes":classes,"issues":issues[0] if issues else [],
                    "sourceText":"不可向用户返回的课表原文", "room":"内部教室", "reviewed":False},ensure_ascii=False)))
        db.commit()
    return identity


def decoy_grades(course):
    with Session(get_engine()) as db:
        row=db.get(Course,course["id"])
        settings=json.loads(row.settings_json)
        settings["course_basic_info"].update(teacher="成绩单教师",class_name="成绩单班级")
        settings.setdefault("course_open_info",{}).update(teacher="成绩单教师",class_name="成绩单班级")
        row.settings_json=json.dumps(settings,ensure_ascii=False)
        term=Term(user_id=row.user_id,course_id=row.id,year_start="2021",year_end="2022",semester="2",
                  teacher="历史成绩单教师",class_name="历史成绩单班级",is_current=1)
        db.add(term);db.commit()
        return term.id


def make_plan(client,course):
    response=client.post("/api/lesson-plans",files={"outline":("合成.docx",syllabus())},data={"course_id":str(course["id"])})
    assert response.status_code==200,response.text
    return response.json()


def text_in_docx(data):
    document=docx.Document(io.BytesIO(data))
    return "\n".join([p.text for p in document.paragraphs]+[c.text for t in document.tables for r in t.rows for c in r.cells])


def test_calendar_uses_current_timetable_teacher_classes_and_preserves_numeric_hours(client,cached_template,tmp_path,monkeypatch):
    course=create_course(client,calendar_data(),name="合成教案课程")
    term_id=decoy_grades(course)
    add_schedule("合成教案课程",[("课表教师",["数媒A班"]),("课表教师",["数媒A班","数媒B班"])])
    add_schedule("合成教案课程",[("历史课表教师",["历史课表班级"])],start="2021",semester="2")
    monkeypatch.setattr(engine,"deepseek",Mock(side_effect=AssertionError("metadata lookup must not call AI")))
    response=client.post(f"/api/courses/{course['id']}/teaching-calendar")
    assert response.status_code==200,response.text[:500]
    workbook=workbook_from_download(response,tmp_path); ws=workbook.active
    assert ws["M3"].value=="课表教师" and ws["G3"].value=="数媒A班、数媒B班"
    assert ws["B8"].data_type=="n" and ws["B8"].number_format=="General" and ws["B8"].value==2
    assert ws["A8"].data_type=="s" and ws["C5"].value=="程序设计基础"
    assert "G3:H3" in {str(r) for r in ws.merged_cells.ranges}
    assert "M3:Q3" in {str(r) for r in ws.merged_cells.ranges}
    assert any("签名" in str(cell.value) for row in ws.iter_rows() for cell in row)
    workbook.close()
    with Session(get_engine()) as db:
        grade=db.get(Term,term_id)
        assert grade.teacher=="历史成绩单教师" and grade.class_name=="历史成绩单班级"
        assert json.loads(db.get(Course,course["id"]).settings_json)["course_basic_info"]["teacher"]=="成绩单教师"


def test_exact_course_period_and_latest_import_do_not_mix_outdated_or_other_course(client):
    course=create_course(client,calendar_data(),name="2D游戏引擎")
    add_schedule("2D游戏引擎",[("旧课表教师",["旧班级"])])
    add_schedule("２Ｄ 游戏引擎",[("最新教师",["新班级"])],newer=2)
    add_schedule("2D游戏引擎基础",[("另一门课教师",["另一班级"])],newer=3)
    result=client.get(f"/api/courses/{course['id']}/teaching-assignment").json()
    assert result["status"]=="ready" and result["teacher"]=="最新教师" and result["classes"]==["新班级"]
    assert "sourceText" not in json.dumps(result) and "不可向用户" not in json.dumps(result)
    assert "room" not in json.dumps(result) and "filename" not in json.dumps(result)


def test_missing_or_unverified_timetable_never_uses_grade_teacher_or_outline_responsible(client):
    course=create_course(client,calendar_data(),name="合成教案课程"); decoy_grades(course)
    plan=make_plan(client,course)
    assert plan["teaching_assignment"]["status"]=="missing"
    assert plan["context"]["授课教师"]==plan["context"]["授课对象"]==""
    add_schedule("合成教案课程",[("疑似教师",["疑似班级"],["教师姓名需与原文复核"])])
    result=client.get(f"/api/lesson-plans/{plan['id']}").json()
    assert result["teaching_assignment"]["status"]=="review_required"
    assert result["context"]["授课教师"]==result["context"]["授课对象"]==""


def test_lesson_choice_is_validated_remembered_and_shared_with_calendar(client,monkeypatch):
    course=create_course(client,calendar_data(),name="合成教案课程")
    add_schedule("合成教案课程",[("教师甲",["A班"]),("教师乙",["B班"])])
    plan=make_plan(client,course)
    assert plan["teaching_assignment"]["status"]=="selection_required"
    download=client.post(f"/api/courses/{course['id']}/teaching-calendar?progress=1")
    assert download.status_code==409 and download.json()["code"]=="teaching_assignment_required"
    assert client.post(f"/api/lesson-plans/{plan['id']}/export",json={"revision":plan["revision"]}).status_code==409
    context=copy.deepcopy(plan["context"])
    context["授课教师"]="成绩单教师"; context["授课对象"]="伪造班级"
    context["lessons"][0]["教学目标"]=["已编辑的目标必须保留"]
    url=f"/api/lesson-plans/{plan['id']}"
    assert client.patch(url,json={"revision":plan["revision"],"context":context,"teacher_choice":"未任课教师"}).status_code==400
    saved=client.patch(url,json={"revision":plan["revision"],"context":context,"teacher_choice":"教师乙"})
    assert saved.status_code==200,saved.text
    current=saved.json()
    assert current["context"]["授课教师"]=="教师乙" and current["context"]["授课对象"]=="B班"
    assert current["context"]["lessons"][0]["教学目标"]==context["lessons"][0]["教学目标"]
    selected=client.get(f"/api/courses/{course['id']}/teaching-assignment").json()
    assert selected["teacher"]=="教师乙" and selected["classes"]==["B班"]
    assert client.get(url).json()["context"]["授课教师"]=="教师乙"
    with Session(get_engine()) as db:
        assert len(list(db.scalars(select(TeachingAssignmentSelection))))==1


def test_choice_is_per_user_and_per_semester_and_owner_protected(client,app):
    course=create_course(client,calendar_data(),name="合成教案课程")
    add_schedule("合成教案课程",[("教师甲",["A班"]),("教师乙",["B班"])])
    period=assignments.academic_period(); url=f"/api/courses/{course['id']}/teaching-assignment"
    body={"teacher":"教师乙",**{key:period[key] for key in ("year_start","year_end","semester")}}
    assert client.post(url,json=body,headers={"Origin":"https://other.example"}).status_code==403
    assert client.post(url,json=body).json()["teacher"]=="教师乙"
    other_sem="2" if period["semester"]=="1" else "1"
    add_schedule("合成教案课程",[("教师甲",["第二学期A班"]),("教师乙",["第二学期B班"])],semester=other_sem)
    other=client.get(url,params={"year_start":period["year_start"],"year_end":period["year_end"],"semester":other_sem}).json()
    assert other["status"]=="selection_required" and other["teacher"]==""
    with TestClient(app) as anonymous: assert anonymous.get(url).status_code==401
    assert client.post("/api/register",json={"invite_code":"TST123","username":"assignment-other","password":PASSWORD}).status_code==200
    assert client.post("/api/login",json={"username":"assignment-other","password":PASSWORD}).status_code==200
    assert client.get(url).status_code==404 and client.post(url,json=body).status_code==404
    own=create_course(client,calendar_data(),name="合成教案课程")
    assert client.get(f"/api/courses/{own['id']}/teaching-assignment").json()["status"]=="selection_required"


def test_word_teacher_classes_refresh_and_stale_download_is_rejected(client,monkeypatch):
    course=create_course(client,calendar_data(),name="合成教案课程")
    add_schedule("合成教案课程",[("课表教师",["课表班级"])])
    plan=make_plan(client,course)
    assert plan["context"]["授课教师"]=="课表教师" and plan["context"]["授课对象"]=="课表班级"
    original=service.make_document
    monkeypatch.setattr(service,"make_document",lambda context,check=True:original(context,check=False))
    monkeypatch.setattr(engine,"deepseek",Mock(side_effect=AssertionError("export must not call AI")))
    def export(current):
        response=client.post(f"/api/lesson-plans/{plan['id']}/export",json={"revision":current["revision"]})
        assert response.status_code==202,response.text
        assert wait_job(client,response.json())["status"]=="success"
    export(plan)
    data=client.get(f"/api/lesson-plans/{plan['id']}/download")
    assert data.status_code==200
    text=text_in_docx(data.content)
    assert "课表教师" in text and "课表班级" in text and "合成教师" not in text and "{{" not in text
    add_schedule("合成教案课程",[("最新教师",["最新班级"])],newer=5)
    current=client.get(f"/api/lesson-plans/{plan['id']}").json()
    assert current["context"]["授课教师"]=="最新教师" and current["download_url"] is None
    assert current["context"]["lessons"]==plan["context"]["lessons"]
    assert client.get(f"/api/lesson-plans/{plan['id']}/download").status_code==409
    export(current)
    updated=client.get(f"/api/lesson-plans/{plan['id']}/download")
    assert updated.status_code==200 and "最新教师" in text_in_docx(updated.content) and "最新班级" in text_in_docx(updated.content)


def test_lesson_semester_query_and_save_refresh_class_without_changing_lesson_content(client):
    course=create_course(client,calendar_data(),name="合成教案课程")
    period=assignments.academic_period(); other_sem="2" if period["semester"]=="1" else "1"
    add_schedule("合成教案课程",[("当前教师",["当前班级"])])
    add_schedule("合成教案课程",[("另一学期教师",["另一学期班级"])],semester=other_sem)
    plan=make_plan(client,course); context=copy.deepcopy(plan["context"])
    context["学期"]=other_sem; context["授课教师"]="前端伪造教师"; context["授课对象"]="前端伪造班级"
    url=f"/api/lesson-plans/{plan['id']}"
    query=client.get(url+"/teaching-assignment",params={"year_start":period["year_start"],"year_end":period["year_end"],"semester":other_sem})
    assert query.json()["teacher"]=="另一学期教师"
    assert client.get(url+"/teaching-assignment",params={"year_start":"invalid"}).status_code==400
    saved=client.patch(url,json={"revision":plan["revision"],"context":context}).json()
    assert saved["context"]["授课教师"]=="另一学期教师" and saved["context"]["授课对象"]=="另一学期班级"
    assert saved["context"]["lessons"]==plan["context"]["lessons"]


def test_active_job_retry_does_not_revise_its_snapshot_after_timetable_updates(client,monkeypatch):
    import threading
    course=create_course(client,calendar_data(),name="合成教案课程")
    add_schedule("合成教案课程",[("提交时教师",["提交时班级"])])
    plan=make_plan(client,course); hold=threading.Event()
    def document(context,check=True):
        assert hold.wait(6)
        return b"PK-offline",{1:1,2:1},[]
    monkeypatch.setattr(service,"make_document",document)
    first=client.post(f"/api/lesson-plans/{plan['id']}/check-pages",json={"revision":plan["revision"]})
    try:
        assert first.status_code==202,first.text
        add_schedule("合成教案课程",[("更新后教师",["更新后班级"])],newer=5)
        current=client.get(f"/api/lesson-plans/{plan['id']}").json()
        retry=client.post(f"/api/lesson-plans/{plan['id']}/check-pages",json={"revision":current["revision"]})
        assert retry.status_code==202,retry.text
        assert retry.json()["job_id"]==first.json()["job_id"]
    finally: hold.set()
    assert wait_job(client,first.json())["status"]=="success"
    current=client.get(f"/api/lesson-plans/{plan['id']}").json()
    assert current["context"]["授课教师"]=="更新后教师" and not current["pages_current"]
