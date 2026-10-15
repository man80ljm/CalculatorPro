"""Regressions for course-plan imports, model numeric fields and cover-only dates."""
import copy
import io
import json
from unittest.mock import Mock
import docx
import pytest
from tests.test_syllabus_completion import legacy,answer
from tests.test_lesson_plans import app,client,create_plan,wait_job
from web_app.syllabus import compatibility,completion,client as ai
from web_app.lesson_plans import engine,service,source
from web_app.service import ServiceError


def test_ai_numeric_fields_and_outline_numbers_are_normalized():
    analysis=compatibility.load("legacy.docx",legacy())
    result=answer(analysis)
    result["fields"]["hours"]=4
    result["fields"]["credits"]=1.0
    result["outline"]=copy.deepcopy(analysis["outline"])
    for row in result["outline"]["lessons"]:
        row["课次"]=int(row["课次"]);row["学时"]=float(row["学时"])
    result["outline"]["考核"]["平时比例"]=30
    result["outline"]["考核"]["期末比例"]=70
    draft,outline=completion.merge(analysis,result)
    assert draft["completion"]["ok"]
    assert outline["lessons"][0]["课次"]=="1" and outline["lessons"][0]["学时"]=="2"
    for bad in (True,float("nan"),float("inf")):
        broken=copy.deepcopy(outline);broken["lessons"][0]["学时"]=bad
        with pytest.raises(ServiceError):compatibility.normalize_outline(broken)


def test_outline_format_error_keeps_valid_ai_sections(monkeypatch):
    analysis=compatibility.load("legacy.docx",legacy())
    result=answer(analysis)
    result["outline"]=copy.deepcopy(analysis["outline"])
    result["outline"]["lessons"][0]["学时"]={"invalid":"2"}
    call=Mock(return_value={"content":json.dumps(result,ensure_ascii=False)})
    monkeypatch.setattr(ai,"complete",call)
    draft,outline=completion.run(analysis)
    assert call.call_count==2 and not draft["completion"]["ok"]
    assert draft["objectives"][-1]=="能够完成并评价造型设计作品"
    assert draft["relation"]["ok"] and len(outline["lessons"])==2
    assert any("第 1 课次的学时" in error for error in draft["completion"]["errors"])


def test_course_lesson_plan_recovers_schedule_and_has_cover_date_once(tmp_path):
    context=service.initial_context(source.parse_bytes("legacy.docx",legacy()))
    context.update(年=2026,月=10,日=14)
    path=tmp_path/"课程教案.docx"
    engine.render(engine.DEFAULT_TEMPLATE,context,str(path))
    document=docx.Document(path)
    dates=[p for p in document.paragraphs if "2026 年" in p.text and "10月" in p.text]
    assert len(dates)==1
    paragraphs=[p.text for p in document.paragraphs]
    assert paragraphs.index(dates[0].text)<next(i for i,p in enumerate(paragraphs) if "第 1 课次" in p)
    result=compatibility.load("课程教案.docx",path.read_bytes())
    assert result["draft"]["fields"]["hours"]["value"]=="4"
    assert len(result["outline"]["lessons"])==2
    assert len(result["outline"]["学习资料"])==len(result["outline"]["学习网站"])==1
    assert any("课程教案" in issue for issue in result["issues"])
    assert result["outline"]["lessons"][0]["章目名称"]=="第一章 基础知识"


def test_old_word_invalidated_by_template_change_without_losing_content(client,monkeypatch):
    plan=create_plan(client)
    original_context=copy.deepcopy(plan["context"])
    monkeypatch.setattr(service,"make_document",lambda context,check=True:(b"PK-template-test",{1:1,2:1},[]))
    url="/api/lesson-plans/"+plan["id"]
    job=client.post(url+"/export",json={"revision":plan["revision"]})
    assert job.status_code==202
    assert wait_job(client,job.json())["status"]=="success"
    assert client.get(url).json()["download_url"]
    monkeypatch.setattr(service,"template_digest",lambda:"updated-cover-date-template")
    reopened=client.get(url).json()
    assert reopened["download_url"] is None
    assert reopened["context"]["lessons"]==original_context["lessons"]
    assert client.get(url+"/download").status_code==409
    job=client.post(url+"/export",json={"revision":reopened["revision"]})
    result=wait_job(client,job.json())
    assert result["status"]=="success" and result["api_calls"]==0
    assert client.get(url).json()["download_url"]
