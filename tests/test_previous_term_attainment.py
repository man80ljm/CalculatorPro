"""导入新学期成绩时，自动带上更早学期已经算出的分目标达成度。"""
import io
import json
import re
from pathlib import Path
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from core_app.ai_report import PREVIOUS_MISSING
from tests.test_grade_register import _FORWARD_HEADERS, _FORWARD_ROWS, _XLSX, _post_register, _register_xlsx
from tests.test_web_flow import PASSWORD
from web_app.db import Term
from web_app.limiter import reset_limiters
from web_app.previous_attainment import (
    AUTO_PREVIOUS_NAME,
    achievement_to_previous_xlsx,
    find_prior_term_with_achievement,
    normalize_achievement,
    term_sort_key,
)
from web_app.terms import apply_term_fields

ROOT = Path(__file__).resolve().parents[1]
HINT = "导入新学期成绩时会自动读取上一学期达成度；也可手动上传覆盖。没有上一学期或不曾计算则为 —"
TERM_A = "2023-2024学年第1学期"
TERM_B = "2024-2025学年第1学期"
TERM_C = "2025-2026学年第1学期"
_GRID = [
    ["考核环节", "占比", "考核方式", "课程目标1", "课程目标2", "小计"],
    ["平时考核", "30%", "课堂参与", "20%", "20%", "40%"],
    ["", "", "作业", "30%", "30%", "60%"],
    ["期末考核", "70%", "考查（设计作品）", "40%", "10%", "50%"],
    ["", "", "课程考核情况表", "10%", "40%", "50%"],
]


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("SECRET_KEY", "test-secret-key")
    monkeypatch.setenv("ALLOW_SQLITE", "1")
    monkeypatch.setenv("DATABASE_URL", "sqlite:///" + (tmp_path / "app.db").as_posix())
    monkeypatch.setenv("UPLOAD_DIR", str(tmp_path / "uploads"))
    monkeypatch.setenv("COOKIE_SECURE", "false")
    monkeypatch.setenv("DEEPSEEK_API_KEY", "")
    reset_limiters()
    from web_app.app import create_app

    with TestClient(create_app()) as test_client:
        assert test_client.post("/api/register", json={"username": "teacher", "password": PASSWORD}).status_code == 200
        assert test_client.post("/api/login", json={"username": "teacher", "password": PASSWORD}).status_code == 200
        yield test_client


def test_normalize_sort_key_and_previous_sheet_roundtrip(tmp_path):
    dated = SimpleNamespace(id=3, year_start="2024", year_end="2025", semester="1")
    later = SimpleNamespace(id=1, year_start="2024", year_end="2025", semester="2")
    same = SimpleNamespace(id=99, year_start="2024", year_end="2025", semester="1")
    missing = SimpleNamespace(id=8, year_start="", year_end="", semester="")
    assert term_sort_key(dated) == (2024, 2025, 1)
    assert term_sort_key(later) == (2024, 2025, 2)
    assert term_sort_key(dated) < term_sort_key(later)
    assert term_sort_key(same) == term_sort_key(dated)
    assert term_sort_key(missing) == (8,)

    normalized = normalize_achievement({"课程目标2": "0.64", "课程目标1": 0.81, "总达成度": 0.7, "噪声": "无"})
    assert normalized == {"课程目标1": 0.81, "课程目标2": 0.64, "课程总目标": 0.7}
    assert normalize_achievement({"课程总目标": 0.4}) == {"课程总目标": 0.4}
    assert normalize_achievement(None) == {}

    loaded = _load_previous(achievement_to_previous_xlsx(normalized), tmp_path)
    assert loaded["课程目标1"] == pytest.approx(0.81)
    assert loaded["课程目标2"] == pytest.approx(0.64)
    assert loaded["课程总目标"] == pytest.approx(0.7)


def test_apply_term_fields_keeps_attainment_unless_replaced():
    term = Term(
        user_id=1,
        course_id=1,
        settings_json=json.dumps(
            {
                "mode": "forward",
                "last_achievement": {"课程目标1": 0.81, "课程总目标": 0.7},
                "previous_achievement": {"课程目标1": 0.5, "课程总目标": 0.4},
            },
            ensure_ascii=False,
        ),
    )
    apply_term_fields(
        term,
        {"mode": "reverse", "course_basic_info": {"major": "数字媒体"}, "course_open_info": {}},
    )
    blob = json.loads(term.settings_json)
    assert blob["mode"] == "reverse"
    assert blob["major"] == "数字媒体"
    assert blob["last_achievement"]["课程目标1"] == 0.81
    assert blob["previous_achievement"]["课程总目标"] == 0.4

    apply_term_fields(
        term,
        {"mode": "forward", "last_achievement": {"课程目标1": 0.2, "课程总目标": 0.2}},
    )
    blob = json.loads(term.settings_json)
    assert blob["last_achievement"]["课程目标1"] == 0.2
    assert blob["previous_achievement"]["课程总目标"] == 0.4


def test_find_prior_skips_self_same_semester_and_empty_terms(monkeypatch):
    older = _memory_term(1, "2023", "2024", "1", {"课程目标1": 0.4, "总达成度": 0.5})
    gap = _memory_term(2, "2024", "2025", "2", None)
    sibling = _memory_term(3, "2025", "2026", "1", {"课程目标1": 0.2, "课程总目标": 0.2})
    current = _memory_term(4, "2025", "2026", "1", {"课程目标1": 0.9, "课程总目标": 0.9})
    monkeypatch.setattr(
        "web_app.previous_attainment.list_terms",
        lambda _db, _course_id: [older, gap, sibling, current],
    )
    found, achievement = find_prior_term_with_achievement(None, SimpleNamespace(id=1), current)
    assert found.id == older.id
    assert achievement["课程目标1"] == pytest.approx(0.4)
    assert achievement["课程总目标"] == pytest.approx(0.5)

    later = _memory_term(5, "2026", "2027", "1", None)
    monkeypatch.setattr(
        "web_app.previous_attainment.list_terms",
        lambda _db, _course_id: [older, gap, sibling, current, later],
    )
    found, achievement = find_prior_term_with_achievement(None, SimpleNamespace(id=1), later)
    assert found.id == current.id
    assert achievement["课程总目标"] == pytest.approx(0.9)


def test_previous_hint_explains_autofill():
    html = (ROOT / "web_app/static/index.html").read_text(encoding="utf-8")
    script = (ROOT / "web_app/static/app.js").read_text(encoding="utf-8")
    assert HINT in html
    assert HINT in script
    assert "不导入则按 0 对比" not in script
    assert "不导入则不做同比" not in html


def test_first_term_without_history_stays_blank(client):
    course_id = _course(client)
    imported = _import_rows(client, course_id, _FORWARD_ROWS, TERM_A)
    assert imported["created_term"] is True
    assert imported["previous_filled"] is False
    assert imported["prior_term_id"] is None
    assert imported["previous_achievement"] == {}
    term_id = imported["term_id"]
    assert _previous_files(client, course_id) == []
    assert not _term_settings(term_id).get("previous_achievement")
    assert not _term_settings(term_id).get("last_achievement")

    calculated = _calculate(client, course_id)
    assert calculated["term_id"] == term_id
    expected = normalize_achievement(calculated["achievement"])
    assert set(expected) == {"课程目标1", "课程目标2", "课程总目标"}
    _assert_achievement(_term_settings(term_id)["last_achievement"], expected)
    assert not _term_settings(term_id).get("previous_achievement")
    assert _previous_files(client, course_id) == []
    per_obj, total = _previous_from_sheet(_eval_sheet(client, course_id))
    assert per_obj == {"课程目标1": PREVIOUS_MISSING, "课程目标2": PREVIOUS_MISSING}
    assert total == PREVIOUS_MISSING

    replaced = _import_rows(client, course_id, _FORWARD_ROWS, TERM_A, confirm=True, teacher="李老师")
    assert replaced["created_term"] is False
    assert replaced["previous_filled"] is False
    assert replaced["previous_achievement"] == {}
    assert replaced["term_id"] == term_id
    assert _previous_files(client, course_id) == []
    kept = _term_settings(term_id)
    _assert_achievement(kept["last_achievement"], expected)
    assert not kept.get("previous_achievement")
    assert client.get(f"/api/courses/{course_id}").json()["current_term"]["teacher"] == "李老师"


def test_later_term_autofills_each_objective_from_prior_calculation(client, tmp_path):
    course_id = _course(client)
    first = _import_rows(client, course_id, _FORWARD_ROWS, TERM_A)
    expected = normalize_achievement(_calculate(client, course_id)["achievement"])

    second = _import_rows(client, course_id, _score_rows(60), TERM_B)
    assert second["created_term"] is True
    assert second["term_id"] != first["term_id"]
    assert second["previous_filled"] is True
    assert second["prior_term_id"] == first["term_id"]
    _assert_achievement(second["previous_achievement"], expected)
    _assert_achievement(_term_settings(second["term_id"])["previous_achievement"], expected)
    assert not _term_settings(second["term_id"]).get("last_achievement")

    info, content = _download_latest_previous(client, course_id)
    assert info["original_name"] == AUTO_PREVIOUS_NAME
    loaded = _load_previous(content, tmp_path)
    _assert_achievement(loaded, expected)

    calculated = _calculate(client, course_id)
    assert calculated["term_id"] == second["term_id"]
    assert abs(calculated["achievement"]["总达成度"] - expected["课程总目标"]) > 0.05
    per_obj, total = _previous_from_sheet(_eval_sheet(client, course_id))
    assert per_obj["课程目标1"] == pytest.approx(expected["课程目标1"])
    assert per_obj["课程目标2"] == pytest.approx(expected["课程目标2"])
    assert total == pytest.approx(expected["课程总目标"])
    assert abs(float(total) - calculated["achievement"]["总达成度"]) > 0.05


def test_reimport_keeps_prior_term_and_not_this_term(client, tmp_path):
    course_id = _course(client)
    first = _import_rows(client, course_id, _FORWARD_ROWS, TERM_A)
    prior = normalize_achievement(_calculate(client, course_id)["achievement"])
    second = _import_rows(client, course_id, _score_rows(60), TERM_B)
    own = normalize_achievement(_calculate(client, course_id)["achievement"])
    assert abs(own["课程总目标"] - prior["课程总目标"]) > 0.05
    _assert_achievement(_term_settings(second["term_id"])["last_achievement"], own)

    replaced = _import_rows(client, course_id, _score_rows(60), TERM_B, confirm=True, teacher="李老师")
    assert replaced["created_term"] is False
    assert replaced["term_id"] == second["term_id"]
    assert replaced["previous_filled"] is True
    assert replaced["prior_term_id"] == first["term_id"]
    _assert_achievement(replaced["previous_achievement"], prior)
    assert abs(replaced["previous_achievement"]["课程总目标"] - own["课程总目标"]) > 0.05

    blob = _term_settings(second["term_id"])
    _assert_achievement(blob["previous_achievement"], prior)
    _assert_achievement(blob["last_achievement"], own)
    loaded = _load_previous(_download_latest_previous(client, course_id)[1], tmp_path)
    _assert_achievement(loaded, prior)

    again = _calculate(client, course_id)
    assert again["term_id"] == second["term_id"]
    per_obj, total = _previous_from_sheet(_eval_sheet(client, course_id))
    assert per_obj["课程目标1"] == pytest.approx(prior["课程目标1"])
    assert per_obj["课程目标2"] == pytest.approx(prior["课程目标2"])
    assert total == pytest.approx(prior["课程总目标"])
    assert abs(float(total) - own["课程总目标"]) > 0.05


def test_same_semester_and_uncalculated_gap_are_not_the_previous_round(client):
    course_id = _course(client)
    first = _import_rows(client, course_id, _FORWARD_ROWS, TERM_A)
    prior = normalize_achievement(_calculate(client, course_id)["achievement"])

    sibling = _import_rows(client, course_id, _score_rows(60, class_name="25数字媒体"), TERM_A)
    assert sibling["term_id"] != first["term_id"]
    assert sibling["previous_filled"] is False
    assert sibling["previous_achievement"] == {}
    assert _previous_files(client, course_id) == []

    middle = _import_rows(client, course_id, _score_rows(70, serial=41), TERM_B)
    assert middle["previous_filled"] is True
    assert middle["prior_term_id"] == first["term_id"]
    later = _import_rows(client, course_id, _score_rows(80, serial=51), TERM_C)
    assert later["previous_filled"] is True
    assert later["prior_term_id"] == first["term_id"]
    assert later["prior_term_id"] != middle["term_id"]
    _assert_achievement(later["previous_achievement"], prior)
    assert not _term_settings(middle["term_id"]).get("last_achievement")


def test_manual_previous_upload_overrides_autofill(client):
    course_id = _course(client)
    _import_rows(client, course_id, _FORWARD_ROWS, TERM_A)
    prior = normalize_achievement(_calculate(client, course_id)["achievement"])
    _import_rows(client, course_id, _score_rows(60), TERM_B)
    manual = {"课程目标1": 0.111, "课程目标2": 0.222, "总达成度": 0.15}
    uploaded = client.post(
        f"/api/courses/{course_id}/files",
        data={"kind": "previous"},
        files={"file": ("手填上一学年.xlsx", achievement_to_previous_xlsx(manual), _XLSX)},
    )
    assert uploaded.status_code == 200, uploaded.text
    info, _content = _download_latest_previous(client, course_id)
    assert info["original_name"] == "手填上一学年.xlsx"

    _calculate(client, course_id)
    per_obj, total = _previous_from_sheet(_eval_sheet(client, course_id))
    assert per_obj["课程目标1"] == pytest.approx(0.111)
    assert per_obj["课程目标2"] == pytest.approx(0.222)
    assert total == pytest.approx(0.15)
    assert abs(float(per_obj["课程目标1"]) - prior["课程目标1"]) > 0.05


def _course(client) -> int:
    created = client.post(
        "/api/courses/from-syllabus",
        json={
            "fields": {"course_name": "创意手作", "course_code": "BJ2230005"},
            "objectives": ["能完成作品", "能分析作品"],
            "description": "简介",
            "grad_req_map": [
                {"requirement": "毕业要求1", "indicator": "指标1", "strength": "H"},
                {"requirement": "毕业要求2", "indicator": "指标2", "strength": "M"},
            ],
            "relation_grid": _GRID,
        },
    )
    assert created.status_code == 200, created.text
    assert created.json()["terms"] == []
    return created.json()["id"]


def _score_rows(score, class_name="24数字媒体", serial=31):
    rows = []
    for offset, name in enumerate(("学员甲", "学员乙")):
        rows.append(
            [class_name, f"99{serial + offset:06d}", name, score, score, score, score, score, 0, score, score]
        )
    return rows


def _import_rows(client, course_id, rows, term, *, confirm=False, teacher="章老师"):
    payload = _register_xlsx(_FORWARD_HEADERS, rows, term=term, teacher=teacher)
    response = _post_register(client, course_id, payload, confirm=confirm)
    assert response.status_code == 200, response.text
    return response.json()


def _calculate(client, course_id) -> dict:
    response = client.post(f"/api/courses/{course_id}/calculate")
    assert response.status_code == 200, response.text
    return response.json()


def _term_settings(term_id: int) -> dict:
    from web_app.db import session_scope
    from web_app.terms import parse_settings

    with session_scope() as db:
        term = db.get(Term, int(term_id))
        assert term is not None
        return parse_settings(term.settings_json)


def _previous_files(client, course_id) -> list:
    course = client.get(f"/api/courses/{course_id}").json()
    return [item for item in course["files"] if item["kind"] == "previous"]


def _download_latest_previous(client, course_id):
    files = _previous_files(client, course_id)
    assert files
    response = client.get(f"/api/courses/{course_id}/files/{files[0]['id']}")
    assert response.status_code == 200, response.text
    return files[0], response.content


def _assert_achievement(actual, expected):
    assert actual["课程目标1"] == pytest.approx(expected["课程目标1"])
    assert actual["课程目标2"] == pytest.approx(expected["课程目标2"])
    assert actual["课程总目标"] == pytest.approx(expected["课程总目标"])


def _memory_term(term_id, year_start, year_end, semester, achievement):
    blob = {}
    if achievement is not None:
        blob["last_achievement"] = achievement
    return SimpleNamespace(
        id=term_id,
        year_start=year_start,
        year_end=year_end,
        semester=semester,
        settings_json=json.dumps(blob, ensure_ascii=False),
    )


def _load_previous(data: bytes, tmp_path) -> dict:
    from core_app.ai_report import AIReportMixin

    path = tmp_path / "previous.xlsx"
    path.write_bytes(data)

    class _Box(AIReportMixin):
        def __init__(self):
            self.relation_payload = {"objectives": ["能完成作品", "能分析作品"]}
            self.objective_requirements = ["能完成作品", "能分析作品"]
            self.status_label = None
            self.previous_achievement_data = None

    box = _Box()
    box.load_previous_achievement(str(path))
    return box.previous_achievement_data


def _eval_sheet(client, course_id):
    course = client.get(f"/api/courses/{course_id}").json()
    chosen = next(
        item
        for item in course["files"]
        if item["kind"] == "output"
        and item["original_name"].endswith(".xlsx")
        and "课程目标达成情况评价结果" in item["original_name"]
        and "逆向" not in item["original_name"]
    )
    downloaded = client.get(f"/api/courses/{course_id}/files/{chosen['id']}")
    assert downloaded.status_code == 200, downloaded.text
    return load_workbook(io.BytesIO(downloaded.content)).active


def _previous_from_sheet(sheet):
    header = [sheet.cell(1, column).value for column in range(1, 8)]
    assert header[6] == "上一轮教学分目标达成值"
    per_obj = {}
    total = None
    for row in sheet.iter_rows(min_row=2, values_only=True):
        label = str(row[0] or "").strip()
        if re.fullmatch(r"课程目标\d+", label):
            per_obj[label] = row[6]
        elif label == "上一轮教学课程目标达成值":
            total = row[5]
    return per_obj, total
