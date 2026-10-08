"""Windows严格控制台编码不应中断Word合并或使报告任务失败。"""
import io
import json
import zipfile

import pytest
from docx import Document

from core_app.report_builder import ReportBuilder
from tests.test_report_job import client, app, _prepare_course, _wait_job
from tests.test_ai_report_once import _Response
from utils import override_outputs_dir


@pytest.mark.parametrize("encoding", ["gbk", "ascii"])
def test_unencodable_logs_preserve_all_merged_documents(tmp_path, monkeypatch, encoding):
    output = tmp_path / "outputs"
    output.mkdir()
    template = tmp_path / "template.docx"
    document = Document()
    for number in range(1, 7):
        document.add_paragraph("{{INSERT_DOC_" + str(number) + "}}")
        source = Document()
        source.add_paragraph(f"合成第{number}部分，文档里的✓和emoji😀必须原样保留。")
        table = source.add_table(rows=1, cols=1)
        table.cell(0, 0).text = f"合成表{number}"
        source.save(output / f"{number}_合成😀.docx")
    document.save(template)
    buffer = io.BytesIO()
    with io.TextIOWrapper(buffer, encoding=encoding, errors="strict") as console:
        with monkeypatch.context() as patch:
            patch.setattr("sys.stdout", console)
            with override_outputs_dir(str(output)):
                filename = ReportBuilder(str(template), str(output)).build({"course_name": "合成课程"}, {}, {})
        console.flush()
        assert buffer.getvalue()
    merged = Document(filename)
    text = "\n".join(paragraph.text for paragraph in merged.paragraphs)
    for number in range(1, 7):
        assert f"合成第{number}部分，文档里的✓和emoji😀必须原样保留。" in text
        assert merged.tables[number - 1].cell(0, 0).text == f"合成表{number}"
    assert "文档插入失败" not in text
    assert "INSERT_DOC" not in text


def test_missing_source_keeps_diagnostic_without_gbk_log_failure(tmp_path, monkeypatch):
    template = tmp_path / "template.docx"
    document = Document()
    document.add_paragraph("{{INSERT_DOC_6}}")
    document.save(template)
    output = tmp_path / "outputs"
    output.mkdir()
    with io.TextIOWrapper(io.BytesIO(), encoding="gbk", errors="strict") as console:
        with monkeypatch.context() as patch:
            patch.setattr("sys.stdout", console)
            with override_outputs_dir(str(output)):
                filename = ReportBuilder(str(template), str(output)).build({}, {}, {})
    merged = Document(filename)
    assert "缺失文档" in merged.paragraphs[0].text


def test_report_job_finishes_under_strict_gbk_and_downloads_complete_report(client, monkeypatch):
    course_id = _prepare_course(client)
    calls = []
    def fake_post(*args, **kwargs):
        calls.append(True)
        return _Response(json.dumps({"overall": "合成总体分析", "objectives": [{"analysis": "合成分析", "improvement": "合成改进"}] * 2}))
    monkeypatch.setenv("DEEPSEEK_API_KEY", "sk-test-only")
    monkeypatch.setattr("core_app.ai_report.requests.post", fake_post)
    with io.TextIOWrapper(io.BytesIO(), encoding="gbk", errors="strict") as console:
        with monkeypatch.context() as patch:
            patch.setattr("sys.stdout", console)
            started = client.post(f"/api/courses/{course_id}/report-jobs")
            assert started.status_code == 200, started.text
            job_id = started.json()["job_id"]
            job, _ = _wait_job(client, job_id)
    assert job["done"] is True, job
    assert job["error"] == ""
    assert job["stage"] == "package"
    assert len(calls) == 1
    downloaded = client.get(f"/api/report-jobs/{job_id}/download")
    assert downloaded.status_code == 200, downloaded.text
    with zipfile.ZipFile(io.BytesIO(downloaded.content)) as archive:
        name = next(name for name in archive.namelist() if "达成度报告" in name and name.endswith(".docx"))
        report = Document(io.BytesIO(archive.read(name)))
    text = "\n".join(paragraph.text for paragraph in report.paragraphs)
    assert "文档插入失败" not in text
    assert "缺失文档" not in text
    assert "INSERT_DOC" not in text
    assert len(report.tables) >= 6
    contents = "\n".join(cell.text for table in report.tables for row in table.rows for cell in row.cells)
    assert "合成总体分析" in contents
    assert "合成改进" in contents
