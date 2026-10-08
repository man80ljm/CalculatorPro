"""报告提示词里的课程目标达成期望值必须和表5是同一次算出来的数。"""
import io
import json
import re
import zipfile

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook

from apply_noise import ScoreRng
from tests.test_deterministic_attainment import _settings
from web_app.service import run_ai_report, run_calculation, run_export

TABLE5_NAME = "5.基于考核结果的课程目标达成情况评价结果表.docx"
_ANSWER = {
    "overall": "总体一段",
    "objectives": [
        {"analysis": "分析1", "improvement": "改进1"},
        {"analysis": "分析2", "improvement": "改进2"},
    ],
}


def _grades(mode: str, score: float = 99) -> bytes:
    book = Workbook()
    sheet = book.active
    if mode == "forward":
        sheet.append(["姓名", "平时考核", "期末考核"])
        sheet.append(["姓名", "平时作业", "期末考试"])
    else:
        sheet.append(["姓名", "平时考核", "期末考核"])
    for index in range(3):
        sheet.append([f"学员{index}", score, score])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _previous(total: float = 0.85) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.append(["课程分目标", "分目标达成值"])
    sheet.append(["课程目标1", 0.84])
    sheet.append(["课程目标2", 0.86])
    sheet.append(["课程目标达成值", total])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _case_settings(mode: str) -> dict:
    settings = _settings(term_id=11)
    settings["mode"] = mode
    return settings


def _capture_prompts(monkeypatch):
    prompts = []

    class _Response:
        def raise_for_status(self):
            return None

        def json(self):
            return {
                "choices": [{"message": {"content": json.dumps(_ANSWER, ensure_ascii=False)}}],
                "usage": {},
            }

    def fake_post(url, headers=None, json=None, timeout=None):
        message = ""
        for item in (json or {}).get("messages") or []:
            if item.get("role") == "user":
                message = item.get("content") or ""
        prompts.append(message)
        return _Response()

    monkeypatch.setenv("DEEPSEEK_API_KEY", "test-key")
    monkeypatch.setattr("core_app.ai_report.requests.post", fake_post)
    return prompts


def _count_uniform(monkeypatch):
    count = {"n": 0}
    original = ScoreRng.uniform

    def wrapped(self, low, high):
        count["n"] += 1
        return original(self, low, high)

    monkeypatch.setattr(ScoreRng, "uniform", wrapped)
    return count


def _docx_bytes(blob: bytes) -> bytes:
    archive = zipfile.ZipFile(io.BytesIO(blob))
    name = next(item for item in archive.namelist() if item.endswith(TABLE5_NAME))
    return archive.read(name)


def _captured_docx(pairs) -> bytes:
    for name, data in pairs:
        if str(name).endswith(TABLE5_NAME):
            return data
    raise AssertionError([name for name, _data in pairs])


def _summary_text(docx_bytes: bytes, label: str) -> str:
    document = Document(io.BytesIO(docx_bytes))
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if not cells or cells[0] != label:
                continue
            for text in reversed(cells):
                if text and text != label:
                    return text
    raise AssertionError(label)


def _excel_expected(blob: bytes) -> str:
    archive = zipfile.ZipFile(io.BytesIO(blob))
    name = next(
        item
        for item in archive.namelist()
        if "课程目标达成情况评价结果" in item and item.endswith(".xlsx")
    )
    book = load_workbook(io.BytesIO(archive.read(name)), data_only=True)
    for sheet in book.worksheets:
        for row in sheet.iter_rows(values_only=True):
            if row and row[0] == "课程目标达成期望值":
                value = row[5]
                return f"{float(value):.3f}"
    raise AssertionError(name)


def _prompt_number(prompt: str, label: str) -> float:
    match = re.search(rf"{re.escape(label)}:\s*([0-9]+(?:\.[0-9]+)?)", prompt)
    assert match, prompt
    return round(float(match.group(1)), 3)


def _run_paths(monkeypatch, mode: str, previous: bytes | None):
    prompts = _capture_prompts(monkeypatch)
    uniform_calls = _count_uniform(monkeypatch)
    grades = _grades(mode)
    settings = _case_settings(mode)

    uniform_calls["n"] = 0
    captured = []
    run_calculation(grades, previous, dict(settings), captured)
    calculated = _summary_text(_captured_docx(captured), "课程目标达成期望值")
    calculated_calls = uniform_calls["n"]

    uniform_calls["n"] = 0
    _filename, exported, _summary = run_export(grades, previous, dict(settings))
    exported_text = _summary_text(_docx_bytes(exported), "课程目标达成期望值")
    exported_calls = uniform_calls["n"]

    uniform_calls["n"] = 0
    _filename, report = run_ai_report(grades, previous, dict(settings))
    report_text = _summary_text(_docx_bytes(report), "课程目标达成期望值")
    report_calls = uniform_calls["n"]
    assert prompts
    return {
        "prompt": prompts[0],
        "calculated": calculated,
        "exported": exported_text,
        "report": report_text,
        "excel": _excel_expected(report),
        "current": _summary_text(_docx_bytes(report), "课程目标达成值"),
        "previous": _summary_text(_docx_bytes(report), "上一轮教学课程目标达成值"),
        "calls": (calculated_calls, exported_calls, report_calls),
    }


@pytest.mark.parametrize("mode", ["reverse", "forward"])
def test_report_expected_matches_table5_with_previous_year(monkeypatch, mode):
    found = _run_paths(monkeypatch, mode, _previous(0.85))
    expected_text = found["report"]
    assert found["calculated"] == found["exported"] == expected_text == found["excel"]
    assert f"课程目标达成期望值: {expected_text}" in found["prompt"]
    assert found["calls"] == (1, 1, 1)

    expected = round(float(expected_text), 3)
    current = round(float(found["current"]), 3)
    previous = round(float(found["previous"]), 3)
    low, high = sorted((previous, current))
    assert low - 0.001 <= expected <= high + 0.001
    assert expected != 0.7
    assert abs(expected - 0.7) >= 0.05
    assert "无数据" not in found["prompt"]
    assert "学员0" not in found["prompt"]


@pytest.mark.parametrize("mode", ["reverse", "forward"])
def test_report_expected_matches_table5_without_previous_year(monkeypatch, mode):
    found = _run_paths(monkeypatch, mode, None)
    expected_text = found["report"]
    assert found["calculated"] == found["exported"] == expected_text == found["excel"]
    assert f"课程目标达成期望值: {expected_text}" in found["prompt"]
    assert found["calls"] == (0, 0, 0)

    expected = round(float(expected_text), 3)
    current = round(float(found["current"]), 3)
    prompt_current = _prompt_number(found["prompt"], "课程目标达成值（本学年）")
    assert expected == current == prompt_current
    assert expected != 0.7
    assert found["previous"] == "—"
    assert "无数据" in found["prompt"]
    assert "学员0" not in found["prompt"]
