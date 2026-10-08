"""混合班级整表导入：人数、缺失学期和计算结果。全部使用合成数据。"""
import io
import zipfile

import pytest
from docx import Document
from openpyxl import load_workbook

from tests.test_grade_register import (
    client, _craft_settings, _register_xlsx, _post_register, _grade_sheet,
    _FORWARD_HEADERS, _FORWARD_ROWS,
)
from tests.test_word_register import _table5
from web_app.grade_register import parse_register, detect_register_mode, forward_workbook, reverse_workbook
from web_app.service import run_export


def mixed_register(term="2024-2025学年第1学期"):
    rows = [list(row) for row in _FORWARD_ROWS]
    rows[1][0] = "24书法"
    content = _register_xlsx(_FORWARD_HEADERS, rows, term=term)
    book = load_workbook(io.BytesIO(content))
    book.active.cell(book.active.max_row, 1, "实考人数：1人 总人数2人")
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


@pytest.mark.parametrize("mode", ["forward", "reverse"])
def test_whole_roster_calculation_matches_original_conversion(client, mode, monkeypatch):
    # 隔离已有 XLSX 时间戳对逆向随机种子的影响，不改生产计算规则。
    monkeypatch.setattr("web_app.service.attainment_seed", lambda *_args: 12345)
    settings = _craft_settings()
    course_id = client.post("/api/courses", json={"name": "创意手作", "settings": settings}).json()["id"]
    content = mixed_register()
    parsed = parse_register("合成.xlsx", content)
    links = settings["relation_payload"]["links"]
    if mode == "forward":
        reference, _ = forward_workbook(parsed, links, detect_register_mode(parsed, links))
    else:
        reference, _ = reverse_workbook(parsed, links)
    imported = _post_register(client, course_id, content, mode=mode)
    assert imported.status_code == 200, imported.text
    assert imported.json()["student_count"] == 2
    assert imported.json()["exam_count"] == 1
    assert not any("请选择" in message for message in imported.json()["warnings"])
    course = client.get(f"/api/courses/{course_id}").json()
    assert course["current_term"]["class_name"] == "专业任选"
    assert course["current_term"]["major"] == ""
    assert course["current_term"]["exam_count"] == 1
    _grade, sheet = _grade_sheet(client, course_id)
    original = load_workbook(io.BytesIO(reference)).active
    assert list(sheet.values) == list(original.values)
    calculated = client.post(f"/api/courses/{course_id}/calculate")
    assert calculated.status_code == 200, calculated.text
    _, archive, expected = run_export(reference, None, course["settings"])
    for key in ("student_count", "average_score", "achievement"):
        assert calculated.json()[key] == expected[key]
    with zipfile.ZipFile(io.BytesIO(archive)) as files:
        name = next(name for name in files.namelist() if name.startswith("5.") and name.endswith(".docx"))
        document = Document(io.BytesIO(files.read(name)))
    expected_table = [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables]
    assert _table5(client, course_id) == expected_table


def test_missing_term_only_asks_for_term_and_keeps_all_classes(client):
    course_id = client.post("/api/courses", json={"name": "创意手作", "settings": _craft_settings()}).json()["id"]
    content = mixed_register(term="")
    blocked = _post_register(client, course_id, content)
    assert blocked.status_code == 409
    assert blocked.json()["needs"] == ["term"]
    imported = _post_register(client, course_id, content, year_start="2026", year_end="2027", semester="1")
    assert imported.status_code == 200, imported.text
    assert imported.json()["student_count"] == 2
    assert imported.json()["class_name"] == "专业任选"
