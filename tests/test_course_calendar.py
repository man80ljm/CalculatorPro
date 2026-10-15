"""首次读大纲保存明细、选课直下、内置空模板和大纲版本回归。"""
import copy
import io
import json
from pathlib import Path
from unittest.mock import Mock

import docx
import openpyxl
import pytest
import xlrd

from tests.test_web_flow import app, client, PASSWORD
from tests.test_teaching_calendar import outline_bytes
from tests.test_syllabus_rules import _ai
from web_app.syllabus import service as syllabus_service, extract_text
from web_app.teaching_calendar import conversion, engine, service
from web_app.teaching_calendar.source import capture_calendar_data

TEMPLATE = Path(service.__file__).with_name("templates") / "teaching_calendar.xls"


def calendar_data():
    return capture_calendar_data(docx.Document(io.BytesIO(outline_bytes())), "同一份大纲.docx")


def wizard_payload(calendar=None, *, name="直下验收课程", hours="4"):
    value = {"fields": {"course_name": name, "hours": hours, "college": "合成学院"},
             "objectives": ["能完成作品"], "grad_req_map": [],
             "relation_grid": [["考核环节", "占比", "考核方式", "课程目标1", "小计"],
                               ["平时考核", "100%", "作业", "100%", "100%"]]}
    if calendar is not None:
        value["teaching_calendar"] = calendar
    return value


def create_course(client, calendar=None, **kwargs):
    response = client.post("/api/courses/from-syllabus", json=wizard_payload(calendar, **kwargs))
    assert response.status_code == 200, response.text
    return response.json()


def download(client, course, output_format="xls", **kwargs):
    return client.post(f"/api/courses/{course['id']}/teaching-calendar?format={output_format}", **kwargs)


def workbook_from_download(response, directory):
    source = directory / "download.xls"
    source.write_bytes(response.content)
    return openpyxl.load_workbook(conversion.convert_file(source, "xlsx", directory))


@pytest.fixture(scope="module")
def converted_template(tmp_path_factory):
    try:
        conversion.soffice_path()
    except service.CalendarError:
        pytest.skip("LibreOffice is required to inspect the school's fixed XLS template")
    directory = tmp_path_factory.mktemp("school-calendar-template")
    source = directory / "original.xls"
    source.write_bytes(TEMPLATE.read_bytes())
    return conversion.convert_file(source, "xlsx", directory)


@pytest.fixture
def cached_template(monkeypatch, converted_template):
    """布局测试复用同一真实转换结果；最终 XLS 仍经过真实 LibreOffice 转换。"""
    original = service.convert_file
    def convert(source, output_format, directory):
        if Path(source).name == "template.xls" and output_format == "xlsx":
            return converted_template
        return original(source, output_format, directory)
    monkeypatch.setattr(service, "convert_file", convert)
    return converted_template


def test_syllabus_extract_reads_document_once_and_saves_calendar(client, monkeypatch):
    monkeypatch.setenv("DEEPSEEK_API_KEY", "sk-isolated-test")
    complete = Mock(return_value={"content": json.dumps(_ai(), ensure_ascii=False)})
    monkeypatch.setattr(syllabus_service, "complete", complete)
    opened = Mock(side_effect=docx.Document)
    monkeypatch.setattr(syllabus_service, "Document", opened)
    monkeypatch.setattr(extract_text, "Document", Mock(side_effect=AssertionError("DOCX opened twice")))
    response = client.post("/api/syllabus/extract", files={"syllabus": ("同一份大纲.docx", outline_bytes())})
    assert response.status_code == 200, response.text
    draft = response.json()
    assert opened.call_count == 1 and complete.call_count == 1
    calendar = draft["teaching_calendar"]
    assert calendar["lessons"][0]["教学内容"] == "知识1与实践"
    assert calendar["lessons"][0]["教学要求"] == "掌握知识1"
    assert calendar["教材"]["教材名称"] == "程序设计基础"
    created = create_course(client, calendar)
    stored = client.get(f"/api/courses/{created['id']}").json()
    assert stored["settings"]["teaching_calendar"] == calendar


def test_saved_calendar_survives_normal_course_edits_and_old_clients(client):
    course = create_course(client, calendar_data())
    edited = copy.deepcopy(course["settings"])
    edited.pop("teaching_calendar")
    edited["course_description"] = "修改课程简介"
    response = client.patch(f"/api/courses/{course['id']}", json={"settings": edited})
    assert response.status_code == 200, response.text
    assert response.json()["settings"]["teaching_calendar"] == course["settings"]["teaching_calendar"]


@pytest.mark.parametrize("count", [2, 8])
def test_builtin_blank_rows_resize_and_signature_follows_lessons(client, cached_template, monkeypatch, count, tmp_path):
    data = calendar_data()
    row = data["lessons"][0]
    data["lessons"] = [{**row, "课次": n} for n in range(1, count + 1)]
    course = create_course(client, data, hours=str(count * 2))
    monkeypatch.setattr(engine, "call_deepseek", Mock(side_effect=AssertionError("download called AI")))
    monkeypatch.setattr(service.docx, "Document", Mock(side_effect=AssertionError("download read DOCX")))
    response = download(client, course)
    assert response.status_code == 200, response.text[:1000]
    workbook = workbook_from_download(response, tmp_path)
    ws = workbook.active
    assert ws["C3"].value == "合成学院" and ws["C4"].value == "直下验收课程"
    assert ws["S4"].value == count * 2
    term = response.headers["X-Calendar-School-Year"]
    semester = response.headers["X-Calendar-Semester"]
    assert ws["A2"].value == f"{term} 学年第 {semester} 学期"
    assert response.headers["X-Calendar-Hours-Match"] == "true"
    merges = {str(area) for area in ws.merged_cells.ranges}
    last = engine.FIRST_ROW + count - 1
    for r in range(engine.FIRST_ROW, last + 1):
        assert ws.cell(r, 1).data_type == "s" and ws.cell(r, 1).number_format == "@"
        assert ws.cell(r, 2).value == 2 and ws.cell(r, 2).data_type == "n"
        assert ws.cell(r, 2).number_format == "General"
        assert f"D{r}:N{r}" in merges and f"O{r}:T{r}" in merges
    assert "签名" in ws.cell(last + 1, 1).value
    assert "签名" in ws.cell(last + 1, 12).value
    assert f"A{last+1}:K{last+4}" in merges and f"L{last+1}:U{last+4}" in merges
    assert ws.max_row == last + 4
    original = openpyxl.load_workbook(cached_template)
    assert engine.count_lesson_rows(original.active) == 4
    # XLS 导出后重读时行高会舍入；在中间工作簿中精确检查签字栏行高的移动。
    _name, editable, _summary = service.generate_for_course(course["name"], course["settings"], "xlsx")
    intermediate = openpyxl.load_workbook(io.BytesIO(editable))
    assert intermediate.active.row_dimensions[last+1].height == original.active.row_dimensions[12].height
    intermediate.close()
    if original.active.print_area:
        assert ws.print_area == f"'{ws.title}'!$A$1:$U${last+4}"
    else:
        assert not ws.print_area
    original.close()
    workbook.close()


def test_true_xls_download_keeps_numeric_hours_and_school_signatures(client, cached_template):
    data = calendar_data()
    data["lessons"][0]["学时"] = "1.5"
    data["lessons"][1]["学时"] = ""
    course = create_course(client, data, hours="1.5")
    response = client.post(f"/api/courses/{course['id']}/teaching-calendar")
    assert response.status_code == 200, response.text[:1000]
    assert response.content.startswith(service.OLE_MAGIC)
    book = xlrd.open_workbook(file_contents=response.content, formatting_info=True)
    ws = book.sheet_by_index(0)
    assert ws.cell(7, 0).ctype == xlrd.XL_CELL_TEXT and ws.cell_value(7, 0) == "1"
    assert ws.cell(7, 1).ctype == xlrd.XL_CELL_NUMBER and ws.cell_value(7, 1) == 1.5
    assert book.format_map[book.xf_list[ws.cell_xf_index(7, 1)].format_key].format_str == "General"
    assert ws.cell(8, 1).ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK)
    assert "签名" in ws.cell_value(9, 0) and "签名" in ws.cell_value(9, 11)
    book.release_resources()


def test_latest_syllabus_is_used_even_when_old_term_is_selected(client, cached_template, tmp_path):
    old = calendar_data()
    old["lessons"][0]["教学内容"] = "旧大纲内容"
    course = create_course(client, old)
    response = client.post(f"/api/courses/{course['id']}/terms", json={"year_start":"2025", "year_end":"2026", "semester":"2"})
    assert response.status_code == 200, response.text
    newer = copy.deepcopy(old)
    newer["lessons"][0]["教学内容"] = "新版大纲内容"
    response = client.post(f"/api/courses/{course['id']}/syllabus-version", json={**wizard_payload(newer), "expected_version":1})
    assert response.status_code == 200, response.text
    assert response.json()["settings"]["teaching_calendar"]["lessons"][0]["教学内容"] == "旧大纲内容"
    response = download(client, course)
    assert response.status_code == 200, response.text[:1000]
    workbook = workbook_from_download(response, tmp_path)
    assert workbook.active["D8"].value == "新版大纲内容"
    assert workbook.active["A2"].value.startswith(response.headers["X-Calendar-School-Year"])
    workbook.close()


def test_blank_course_or_missing_schedule_returns_actionable_error(client):
    course = create_course(client)
    response = download(client, course)
    assert response.status_code == 400
    assert response.json()["code"] == "calendar_source_missing"
    assert "从大纲建课" in response.json()["detail"]
    assert client.get("/healthz").status_code == 200


def test_missing_outline_table_does_not_block_original_course_creation(client):
    data = capture_calendar_data(docx.Document(io.BytesIO(outline_bytes(table=False))), "大纲.docx")
    assert "课次" in data["error"]
    course = create_course(client, data)
    response = download(client, course)
    assert response.status_code == 400
    assert response.json()["code"] == "calendar_source_unavailable"


def test_direct_download_requires_login_and_course_ownership(client, app):
    course = create_course(client, calendar_data())
    from fastapi.testclient import TestClient
    with TestClient(app) as anonymous:
        assert download(anonymous, course).status_code == 401
    assert client.post("/api/register", json={"invite_code":"TST123", "username":"other-calendar-user", "password":PASSWORD}).status_code == 200
    assert client.post("/api/login", json={"username":"other-calendar-user", "password":PASSWORD}).status_code == 200
    assert download(client, course).status_code == 404


def test_format_origin_and_numeric_validation(client):
    course = create_course(client, calendar_data())
    assert download(client, course, output_format="csv").status_code == 400
    xlsx = download(client, course, output_format="xlsx")
    assert xlsx.status_code == 400 and xlsx.json()["code"] == "invalid_format"
    assert download(client, course, headers={"Origin":"https://unrelated.example"}).status_code == 403
    bad = calendar_data()
    bad["lessons"][0]["学时"] = "4学时"
    response = client.post("/api/courses/from-syllabus", json=wizard_payload(bad, name="错误学时"))
    assert response.status_code == 400 and response.json()["code"] == "invalid_hours"


def test_hours_mismatch_is_returned_with_download(client, cached_template):
    course = create_course(client, calendar_data(), hours="48")
    response = download(client, course)
    assert response.status_code == 200, response.text[:1000]
    assert response.headers["X-Calendar-Hours-Match"] == "false"
    assert float(response.headers["X-Calendar-Total-Hours"]) == 4
    assert float(response.headers["X-Calendar-Course-Hours"]) == 48
