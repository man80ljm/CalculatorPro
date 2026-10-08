"""大纲草稿的程序规则，以及建课接口。模型调用全部打桩。"""
import io
import json

import pytest
from docx import Document
from openpyxl import Workbook
from fastapi.testclient import TestClient

from tests.test_web_flow import PASSWORD, _settings
from web_app.limiter import reset_limiters


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("SECRET_KEY", "test-secret-key")
    monkeypatch.setenv("ALLOW_SQLITE", "1")
    monkeypatch.setenv("DATABASE_URL", "sqlite:///" + (tmp_path / "app.db").as_posix())
    monkeypatch.setenv("UPLOAD_DIR", str(tmp_path / "uploads"))
    monkeypatch.setenv("COOKIE_SECURE", "false")
    monkeypatch.setenv("DEEPSEEK_API_KEY", "")
    reset_limiters()
    from web_app.app import create_app

    with TestClient(create_app()) as test_client:
        assert test_client.post("/api/register", json={"username": "teacher", "password": PASSWORD}).status_code == 200
        assert test_client.post("/api/login", json={"username": "teacher", "password": PASSWORD}).status_code == 200
        yield test_client
from web_app.syllabus.prefilter import prefilter, still_has_roster
from web_app.syllabus.rules import build_draft, validate_relation_grid
from web_app.syllabus.service import extract_draft


def _field(value="", status="需手填", source="", reason="", candidate=""):
    return {"value": value, "status": status, "source": source, "reason": reason, "candidate": candidate}


def _ai(**overrides):
    basic = {
        "course_name": _field("2D游戏引擎", "已填", "大纲"),
        "credits": _field("3", "已填"),
        "hours": _field("48", "已填"),
        "course_type": _field("专业限选课", "已填", "模型自作主张"),
        "course_code": _field("", "需手填"),
        "school_year_term": _field("4/6", "已填", "大纲"),
        "college": _field("美术学院", "已填"),
        "teacher": _field("张三", "已填", "大纲·课程负责人"),
        "major": _field("数字媒体", "已填", "从班级推断"),
        "class_name": _field("", "需手填"),
        "student_count": _field("", "需手填"),
        "exam_count": _field("", "需手填"),
    }
    data = {
        "course_basic_info": basic,
        "course_open_info": {},
        "course_description": _field("这是课程简介。", "已填"),
        "objectives_status": "已填",
        "objectives": [{"index": 1, "text": "能完成作品"}],
        "grad_req_map": [
            {
                "objective": "课程目标1",
                "indicator": _field("", "需手填", reason="无法区分"),
                "requirement": _field("理论知识", "已填"),
                "strength": "H",
            }
        ],
        "relation_table": {
            "objectives_count": 1,
            "links": [
                {
                    "name": "平时考核",
                    "ratio_text": "100%",
                    "methods": [{"name": "作业", "weights": {"课程目标1": "100%"}, "subtotal_text": "100%"}],
                }
            ],
        },
    }
    data.update(overrides)
    return data


SYLLABUS = "课程模块：专业拓展课程\n课程类别：选修\n课程性质：专业拓展课程/选修\n适用专业：数字媒体艺术\n课程负责人：张三\n开课学期：4/6\n"


def test_program_rules_override_the_model():
    register = {
        "course_type": "专业限选课",
        "course_code": "BJ2330188",
        "teacher": "黄老师",
        "school_year_term": "2025-2026学年第2学期",
        "year_start": "2025",
        "year_end": "2026",
        "semester": "2",
        "classes": ["24数字媒体"],
        "student_count": 31,
        "exam_count": 31,
        "college": "美术学院",
        "credits": "3",
        "course_name": "2D游戏引擎",
    }
    draft = build_draft(_ai(), SYLLABUS, "课程性质：专业限选课\n任课教师：黄老师", register)
    nature = draft["fields"]["course_type"]
    assert nature["status"] == "需手填"
    assert nature["value"] == ""
    labels = [item["label"] for item in nature["candidates"]]
    assert "大纲：专业拓展课程" in labels
    assert "大纲：选修" in labels
    assert "大纲：专业拓展课程/选修" in labels
    assert "成绩登记表：专业限选课" in labels
    for key in ("teacher", "major", "school_year_term", "class_name", "student_count", "exam_count", "year_start"):
        assert draft["fields"][key]["value"] == ""
        assert draft["fields"][key]["status"] == "需手填"
        assert draft["fields"][key]["reason"] == "学期导入时填写"
    assert "major" not in draft["groups"]["course"]
    assert "major" in draft["groups"]["term"]
    assert draft["fields"]["course_code"]["value"] == "BJ2330188"
    grad = draft["grad_req_map"][0]
    assert grad["indicator"]["value"] == ""
    assert grad["indicator"]["status"] == "需手填"
    assert grad["indicator"]["reason"] == "大纲没有课程目标与毕业要求的矩阵"
    assert grad["requirement"]["value"] == ""
    assert grad["requirement"]["status"] == "需手填"
    assert grad["requirement"]["reason"] == "大纲没有课程目标与毕业要求的矩阵"
    assert grad["strength"] == "H"
    assert draft["relation"]["ok"] is True


def test_table_cells_become_course_nature_candidates():
    from web_app.syllabus.rules import course_type_field, major_field

    syllabus = "| 课程编码 |  | 适用专业 | 数字媒体艺术 |  |\n| 课程模块 | 专业拓展课程 | 课程类别 | 选修 |  |"
    nature = course_type_field(syllabus, "", {"course_type": "专业限选课"})
    labels = [item["label"] for item in nature["candidates"]]
    assert labels == ["大纲：专业拓展课程", "大纲：选修", "大纲：专业拓展课程/选修", "成绩登记表：专业限选课"]
    assert nature["status"] == "需手填"
    major = major_field(syllabus, "", None, {})
    assert major["status"] == "已填"
    assert major["value"] == "数字媒体艺术"
    assert major["source"] == "大纲·适用专业"


def test_teacher_and_major_rules_stay_available_but_draft_defers_them():
    from web_app.syllabus.rules import major_field, school_year_field, teacher_field

    alone = build_draft(_ai(), SYLLABUS, "", None)
    assert alone["fields"]["teacher"]["reason"] == "学期导入时填写"
    assert alone["fields"]["school_year_term"]["reason"] == "学期导入时填写"
    teacher = teacher_field(SYLLABUS, "", None, _ai()["course_basic_info"]["teacher"])
    assert teacher["status"] == "需手填"
    assert teacher["value"] == ""
    assert teacher["candidates"][0]["value"] == "张三"
    year = school_year_field(SYLLABUS, None, _ai()["course_basic_info"]["school_year_term"])
    assert "4/6" in year["reason"] or year["candidates"]

    derived = major_field("", "", {"classes": ["24数字媒体", "24数字媒体2班"], "teacher": ""}, {})
    assert derived["status"] == "已填"
    assert derived["value"] == "数字媒体"
    assert derived["source"] == "成绩登记表·班级（推导）"
    conflict = major_field("", "", {"classes": ["23美术教育A", "24数字媒体"]}, {})
    assert conflict["status"] == "需手填"
    assert {item["value"] for item in conflict["candidates"]} == {"美术教育", "数字媒体"}


def test_need_fill_literal_is_dropped_and_bad_link_name_fails():
    ai = _ai()
    ai["course_basic_info"]["college"] = _field("需手填", "已填")
    draft = build_draft(ai, SYLLABUS, "", None)
    assert draft["fields"]["college"]["value"] == ""
    bad = validate_relation_grid(
        [
            ["考核环节", "占比", "考核方式", "课程目标1", "小计"],
            ["课堂表现", "100%", "发言", "100%", "100%"],
        ]
    )
    assert bad["ok"] is False
    assert any("平时" in item for item in bad["errors"])


def test_prefilter_drops_roster_names():
    document = Document()
    document.add_paragraph("一、基本信息")
    document.add_paragraph("适用专业：数字媒体艺术")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "学号"
    table.rows[0].cells[1].text = "姓名"
    table.rows[1].cells[0].text = "20240001"
    table.rows[1].cells[1].text = "测试生甲"
    document.add_paragraph("二、教学内容")
    document.add_paragraph("这里是不应送出的教学内容细节")
    buffer = io.BytesIO()
    document.save(buffer)
    from web_app.syllabus.extract_text import docx_to_markdown

    filtered = prefilter(docx_to_markdown(buffer.getvalue()))
    assert "测试生甲" not in filtered
    assert "20240001" not in filtered
    assert "共 1 行学生记录" in filtered
    assert "教学内容细节" not in filtered
    assert "数字媒体艺术" in filtered
    assert still_has_roster(filtered) is False


def test_doc_and_scan_do_not_call_model(monkeypatch):
    monkeypatch.setenv("DEEPSEEK_API_KEY", "sk-test")
    called = {"n": 0}

    def boom(*_args, **_kwargs):
        called["n"] += 1
        raise AssertionError("不应该调用模型")

    monkeypatch.setattr("web_app.syllabus.service.complete", boom)
    with pytest.raises(Exception, match="另存为 .docx"):
        extract_draft("大纲.doc", b"not-a-real-doc")
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer)
    pdf.showPage()
    pdf.save()
    with pytest.raises(Exception, match="扫描件"):
        extract_draft("scan.pdf", buffer.getvalue())
    assert called["n"] == 0


def test_extract_uses_filled_only_and_retries_once(monkeypatch):
    monkeypatch.setenv("DEEPSEEK_API_KEY", "sk-test")
    calls = {"n": 0}

    def fake_complete(_messages, **_kwargs):
        calls["n"] += 1
        if calls["n"] == 1:
            return {"content": "不是 JSON", "model": "deepseek-flash", "prompt_tokens": 3, "completion_tokens": 1}
        import json

        return {
            "content": json.dumps(_ai(), ensure_ascii=False),
            "model": "deepseek-flash",
            "prompt_tokens": 5,
            "completion_tokens": 4,
        }

    monkeypatch.setattr("web_app.syllabus.service.complete", fake_complete)
    document = Document()
    document.add_paragraph("课程模块：专业拓展课程")
    document.add_paragraph("课程类别：选修")
    buffer = io.BytesIO()
    document.save(buffer)
    draft = extract_draft("大纲.docx", buffer.getvalue())
    assert calls["n"] == 2
    assert draft["meta"]["calls"] == 2
    assert draft["fields"]["course_type"]["value"] == ""
    assert draft["fields"]["course_name"]["value"] == "2D游戏引擎"
    assert "sk-test" not in str(draft)


def _assert_course_has_no_term_identity(settings):
    basic = settings.get("course_basic_info") or {}
    opened = settings.get("course_open_info") or {}
    for key in ("teacher", "major", "class_name", "school_year_term", "student_count", "exam_count"):
        assert not basic.get(key)
    for key in ("year_start", "year_end", "semester", "term", "teacher"):
        assert not opened.get(key)
    assert "student_count" not in settings


def test_from_syllabus_allows_empty_grad_and_blocks_bad_relation(client):
    grid = [
        ["考核环节", "占比", "考核方式", "课程目标1", "小计"],
        ["平时考核", "100%", "作业", "100%", "100%"],
    ]
    bad = client.post(
        "/api/courses/from-syllabus",
        json={"fields": {"course_name": "空课"}, "objectives": ["目标"], "grad_req_map": [{"requirement": "", "indicator": "理论知识"}], "relation_grid": [["考核环节", "占比", "考核方式", "课程目标1"], ["随便考", "100%", "口试", "100%"]]},
    )
    assert bad.status_code == 400
    assert "平时" in bad.json()["detail"]
    created = client.post(
        "/api/courses/from-syllabus",
        json={
            "fields": {
                "course_name": "从大纲来的课",
                "course_type": "",
                "major": "数字媒体艺术",
                "teacher": "黄老师",
                "class_name": "24数字媒体",
                "school_year_term": "2025-2026学年第2学期",
                "year_start": "2025",
                "year_end": "2026",
                "semester": "2",
                "student_count": "31",
                "college": "美术学院",
            },
            "description": "简介",
            "objectives": ["能完成作品"],
            "grad_req_map": [{"requirement": "", "indicator": "理论知识", "strength": "H"}],
            "relation_grid": grid,
        },
    )
    assert created.status_code == 200, created.text
    body = created.json()
    assert body["name"] == "从大纲来的课"
    assert body["settings"]["grad_req_map"][0]["requirement"] == ""
    assert body["settings"]["course_basic_info"]["course_type"] == ""
    assert body["settings"]["course_basic_info"]["college"] == "美术学院"
    assert body["current_term"] is None
    assert body["terms"] == []
    _assert_course_has_no_term_identity(body["settings"])


def test_missing_key_is_explicit(client, monkeypatch):
    monkeypatch.setenv("DEEPSEEK_API_KEY", "")
    document = Document()
    document.add_paragraph("课程名称：测试")
    buffer = io.BytesIO()
    document.save(buffer)
    response = client.post("/api/syllabus/extract", files={"syllabus": ("a.docx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert response.status_code == 400
    assert "DEEPSEEK_API_KEY" in response.json()["detail"]


def test_second_term_keeps_its_own_files(client):
    course_id = client.post("/api/courses", json={"name": "学期课", "settings": _settings()}).json()["id"]
    first = client.get(f"/api/courses/{course_id}").json()
    uploaded = client.post(
        f"/api/courses/{course_id}/files",
        data={"kind": "grade"},
        files={"file": ("grades.xlsx", b"PK\x03\x04grades", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert uploaded.status_code == 200, uploaded.text
    created = client.post(f"/api/courses/{course_id}/terms", json={})
    assert created.status_code == 200, created.text
    assert len(created.json()["terms"]) == 2
    assert created.json()["files"] == []
    switched = client.post(f"/api/courses/{course_id}/terms/{first['current_term_id']}/select")
    assert switched.status_code == 200
    assert any(item["kind"] == "grade" for item in switched.json()["files"])
    blocked = client.delete(f"/api/courses/{course_id}/terms/{first['current_term_id']}")
    assert blocked.status_code == 400
    assert "成绩" in blocked.json()["detail"]
    removed = client.delete(f"/api/courses/{course_id}/terms/{created.json()['current_term_id']}?confirm=1")
    assert removed.status_code == 200, removed.text
    assert len(removed.json()["terms"]) == 1


def test_from_syllabus_keeps_description_string(client):
    grid = [
        ["考核环节", "占比", "考核方式", "课程目标1", "小计"],
        ["平时考核", "100%", "作业", "100%", "100%"],
    ]
    created = client.post(
        "/api/courses/from-syllabus",
        json={
            "fields": {
                "course_name": {"value": "对象课", "status": "已填", "source": "大纲"},
                "course_code": {"value": "BJ2330188", "status": "已填"},
                "student_count": {"value": "31", "status": "已填"},
            },
            "description": {
                "value": "这是课程简介原文。",
                "status": "已填",
                "source": "大纲",
                "reason": "",
                "candidate": "",
            },
            "objectives": [{"text": "能完成作品"}],
            "grad_req_map": [{"requirement": {"value": ""}, "indicator": {"value": "理论知识"}}],
            "relation_grid": grid,
        },
    )
    assert created.status_code == 200, created.text
    description = created.json()["settings"]["course_description"]
    assert description == "这是课程简介原文。"
    assert isinstance(description, str)
    assert description != "[object Object]"
    assert created.json()["settings"]["course_basic_info"]["course_name"] == "对象课"
    rejected = client.post(
        "/api/courses/from-syllabus",
        json={
            "fields": {"course_name": "坏简介"},
            "description": ["不是字符串"],
            "objectives": ["目标"],
            "grad_req_map": [{"requirement": "", "indicator": "理论知识"}],
            "relation_grid": grid,
        },
    )
    assert rejected.status_code == 400


def test_new_term_label_uses_year_after_it_is_filled(client):
    course_id = client.post("/api/courses", json={"name": "学期名", "settings": _settings()}).json()["id"]
    created = client.post(f"/api/courses/{course_id}/terms", json={})
    assert created.status_code == 200, created.text
    assert created.json()["current_term"]["label"] == "新学期（请填写学年学期）"
    updated = client.patch(
        f"/api/courses/{course_id}/terms/{created.json()['current_term_id']}",
        json={
            "year_start": "2025",
            "year_end": "2026",
            "semester": "2",
            "school_year_term": "",
            "class_name": "24数字媒体",
        },
    )
    assert updated.status_code == 200, updated.text
    term = updated.json()["current_term"]
    assert term["label"] == "2025-2026学年第2学期"
    assert term["class_name"] == "24数字媒体"


def _synthetic_register(rows) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet["A1"] = "广东第二师范学院2025-2026学年第2学期课程成绩登记表"
    sheet["A2"] = "课程名称：2D游戏引擎                           课程代码：BJ2330188                          课程性质：专业限选课"
    sheet["A3"] = "开课学院：美术学院                           任课教师：黄老师                               学分：3  考核方式：考查"
    sheet["A5"] = "班级"
    sheet["B5"] = "学号"
    sheet["D5"] = "姓名"
    sheet["E5"] = "平时(30%)"
    sheet["F5"] = "期中(0%)"
    sheet["G5"] = "期末(70%)"
    sheet["H5"] = "总评"
    sheet["I5"] = "备注"
    for index, (student_id, name) in enumerate(rows):
        row = 6 + index
        sheet.cell(row, 1, "24数字媒体")
        sheet.cell(row, 2, student_id)
        sheet.cell(row, 4, name)
        sheet.cell(row, 5, 80)
        sheet.cell(row, 6, 0)
        sheet.cell(row, 7, 90)
    sheet.cell(6 + len(rows) + 2, 1, f"实考 {len(rows)}人 总人数 {len(rows)}人")
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _syllabus_bytes() -> bytes:
    document = Document()
    document.add_paragraph("课程名称：2D游戏引擎")
    document.add_paragraph("适用专业：数字媒体艺术")
    document.add_paragraph("课程目标1：能完成作品")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_from_syllabus_imports_register_and_calculates(client, monkeypatch):
    monkeypatch.setenv("DEEPSEEK_API_KEY", "sk-test")
    roster = [("8800112201", "合成生甲"), ("8800112202", "合成生乙"), ("8800112203", "合成生丙")]
    register = _synthetic_register(roster)
    sent = []
    ai = _ai()
    ai["relation_table"] = {
        "objectives_count": 1,
        "links": [
            {
                "name": "平时考核",
                "ratio_text": "30%",
                "methods": [{"name": "作业", "weights": {"课程目标1": "100%"}, "subtotal_text": "100%"}],
            },
            {
                "name": "期末考核",
                "ratio_text": "70%",
                "methods": [{"name": "考查", "weights": {"课程目标1": "100%"}, "subtotal_text": "100%"}],
            },
        ],
    }

    def fake_complete(messages, **_kwargs):
        sent.append(messages)
        return {
            "content": json.dumps(ai, ensure_ascii=False),
            "model": "deepseek-flash",
            "prompt_tokens": 5,
            "completion_tokens": 4,
        }

    monkeypatch.setattr("web_app.syllabus.service.complete", fake_complete)
    extracted = client.post(
        "/api/syllabus/extract",
        files={
            "syllabus": ("大纲.docx", _syllabus_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
            "register": ("登记.xlsx", register, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        },
    )
    assert extracted.status_code == 200, extracted.text
    draft = extracted.json()
    assert len(sent) == 1
    posted = json.dumps(sent, ensure_ascii=False)
    for _student_id, name in roster:
        assert name not in posted
    assert "8800112201" not in posted
    assert "8800112202" not in posted
    assert "8800112203" not in posted
    assert "共 3 行学生记录" in posted
    assert draft["relation"]["ok"] is True
    fields = {key: item["value"] for key, item in draft["fields"].items()}
    created = client.post(
        "/api/courses/from-syllabus",
        data={
            "payload": json.dumps(
                {
                    "fields": fields,
                    "objectives": draft["objectives"],
                    "description": draft["course_description"]["value"],
                    "grad_req_map": draft["grad_req_map"],
                    "relation_grid": draft["relation"]["grid"],
                },
                ensure_ascii=False,
            )
        },
        files={"register": ("登记.xlsx", register, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert created.status_code == 200, created.text
    body = created.json()
    assert "grade_import_error" not in body
    assert body["grade_import"]["student_count"] == 3
    assert body["grade_import"]["mode"] == "forward"
    assert body["current_term"]["mode"] == "forward"
    assert any(item["kind"] == "grade" for item in body["files"])
    assert len(sent) == 1
    course_id = body["id"]
    calculated = client.post(f"/api/courses/{course_id}/calculate")
    assert calculated.status_code == 200, calculated.text
    assert calculated.json()["student_count"] == 3
    assert calculated.json()["mode"] == "forward"
    assert len(sent) == 1
    opened = client.post(f"/api/courses/{course_id}/terms", json={})
    assert opened.status_code == 200, opened.text
    assert opened.json()["files"] == []
    imported = client.post(
        f"/api/courses/{course_id}/grade-register",
        data={"confirm": "1"},
        files={"file": ("空白学期.xlsx", register, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert imported.status_code == 200, imported.text
    assert imported.json()["student_count"] == 3
    assert imported.json()["mode"] == "forward"
    assert len(sent) == 1


def test_bad_register_does_not_roll_back_the_new_course(client):
    grid = [
        ["考核环节", "占比", "考核方式", "课程目标1", "小计"],
        ["平时考核", "100%", "作业", "100%", "100%"],
    ]
    created = client.post(
        "/api/courses/from-syllabus",
        data={
            "payload": json.dumps(
                {
                    "fields": {"course_name": "导入失败仍保留"},
                    "objectives": ["能完成作品"],
                    "description": "简介",
                    "grad_req_map": [{"requirement": "", "indicator": "理论知识"}],
                    "relation_grid": grid,
                },
                ensure_ascii=False,
            )
        },
        files={"register": ("坏表.xlsx", b"this-is-not-a-workbook", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert created.status_code == 200, created.text
    body = created.json()
    assert body["name"] == "导入失败仍保留"
    assert "无法读取" in body["grade_import_error"]
    assert body["files"] == []
    listed = client.get(f"/api/courses/{body['id']}")
    assert listed.status_code == 200
    assert listed.json()["name"] == "导入失败仍保留"


def test_extract_hides_forward_register_scores(client, monkeypatch):
    monkeypatch.setenv("DEEPSEEK_API_KEY", "sk-test")
    book = Workbook()
    sheet = book.active
    sheet["A1"] = "课程名称：2D游戏引擎 课程代码：BJ2330188"
    sheet["A5"] = "班级"
    sheet["B5"] = "学号"
    sheet["C5"] = "姓名"
    sheet["D5"] = "课堂参与"
    sheet["E5"] = "作业"
    sheet["F5"] = "平时(30%)"
    sheet["G5"] = "期末(70%)"
    sheet["A6"] = "24数字媒体"
    sheet["B6"] = "8800112291"
    sheet["C6"] = "合成生甲"
    sheet["D6"] = 67.5
    sheet["E6"] = 91.25
    sheet["F6"] = 73.5
    sheet["G6"] = 64.25
    buffer = io.BytesIO()
    book.save(buffer)
    sent = []

    def fake_complete(messages, **_kwargs):
        sent.append(messages)
        return {
            "content": json.dumps(_ai(), ensure_ascii=False),
            "model": "deepseek-flash",
            "prompt_tokens": 3,
            "completion_tokens": 2,
        }

    monkeypatch.setattr("web_app.syllabus.service.complete", fake_complete)
    extracted = client.post(
        "/api/syllabus/extract",
        files={
            "syllabus": ("大纲.docx", _syllabus_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
            "register": ("正向.xlsx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        },
    )
    assert extracted.status_code == 200, extracted.text
    posted = json.dumps(sent, ensure_ascii=False)
    assert "合成生甲" not in posted
    assert "8800112291" not in posted
    for score in ("67.5", "91.25", "73.5", "64.25"):
        assert score not in posted
    assert "课堂参与" in posted
    assert "作业" in posted
