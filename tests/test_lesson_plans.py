"""教案接入、受保护源文件、原有重写逻辑、任务及无 API Word 导出。"""
import copy
import io
import json
import time
import threading
from datetime import timedelta
from pathlib import Path
from unittest.mock import Mock

import docx
from docx.oxml.ns import qn
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select
from sqlalchemy.orm import Session

from tests.test_web_flow import app, client, PASSWORD
from tests.test_teaching_calendar import outline_bytes
from web_app.db import LessonPlanSource, LessonPlan, LessonPlanTask, get_engine, utcnow
from web_app.lesson_plans import engine, source, service, jobs


@pytest.fixture(autouse=True)
def reset_lesson_limiter():
    from web_app.lesson_plans.routes import limiter
    limiter.reset()
    yield
    limiter.reset()


def syllabus():
    document=docx.Document(io.BytesIO(outline_bytes()))
    info=document.add_table(rows=2,cols=2)
    for row,values in zip(info.rows,[("课程中文名称","合成教案课程"),("课程负责人","合成教师")]):
        for cell,value in zip(row.cells,values): cell.text=value
    document.add_paragraph("六、参考资料")
    document.add_paragraph("（一）参考书目")
    document.add_paragraph("[1] 合成参考资料")
    document.add_paragraph("（三）网络资源")
    document.add_paragraph("合成学习网站：https://example.org")
    stream=io.BytesIO(); document.save(stream); return stream.getvalue()


def valid_ai():
    return {"教学目标":["目"*140],"基本内容":["内"*230],"教学重点":["重"],"教学难点":["难"],
            "教学手段与方法":["方"*120],"思考题":"题"*20,"讨论题":"题"*20,"作业":"题"*20}


def create_plan(client):
    response=client.post("/api/lesson-plans",files={"outline":("合成大纲.docx",syllabus())})
    assert response.status_code==200,response.text
    return response.json()


def wait_job(client,job,timeout=8):
    deadline=time.monotonic()+timeout
    while time.monotonic()<deadline:
        response=client.get(f"/api/lesson-plan-jobs/{job['job_id']}")
        assert response.status_code==200,response.text
        result=response.json()
        if result["status"] in {"success","fail"}: return result
        time.sleep(0.03)
    raise AssertionError("job timeout")


def test_original_limits_prompts_and_exactly_one_word_count_retry(monkeypatch):
    short={"教学目标":["短"],"基本内容":["短"],"教学手段与方法":["短"]}
    call=Mock(side_effect=[short,valid_ai()]); monkeypatch.setattr(engine,"deepseek",call)
    outline=source.parse_bytes("大纲.docx",syllabus())
    result=engine.gen_lesson(outline,outline["lessons"][0],"sk-offline")
    assert call.call_count==2 and "上一版字数不符合要求" in call.call_args.args[1]
    assert engine.LIMITS=={"教学目标":(120,180),"教学内容":(220,300),"教学手段与方法":(100,150),"思考讨论作业":(50,80)}
    assert all(low<=engine.hanzi_len(result[key])<=high for key,(low,high) in engine.LIMITS.items())
    call.reset_mock(side_effect=True); call.return_value=short
    engine.gen_lesson(outline,outline["lessons"][0],"sk-offline")
    assert call.call_count==2


def test_uploaded_raw_docx_is_preserved_and_reusable_with_course(client):
    course=client.post("/api/courses",json={"name":"关联教案课程"}).json()
    data=syllabus()
    saved=client.post(f"/api/courses/{course['id']}/lesson-plan-outline",files={"outline":("原大纲.docx",data)})
    assert saved.status_code==200,saved.text
    with Session(get_engine()) as db:
        row=db.get(LessonPlanSource,saved.json()["id"])
        assert source.source_path(row).read_bytes()==data
    available=client.get(f"/api/courses/{course['id']}/lesson-plan-outline").json()["source"]
    assert available["teacher"]=="合成教师" and available["lesson_count"]==2
    plan=client.post("/api/lesson-plans",json={"source_id":available["id"]}).json()
    assert plan["course_id"]==course["id"] and plan["context"]["授课教师"]=="" and plan["teaching_assignment"]["status"]=="missing"
    assert plan["context"]["学习资料"]==["合成参考资料"]
    different=docx.Document(io.BytesIO(data)); different.add_paragraph("新版大纲的说明")
    stream=io.BytesIO(); different.save(stream)
    next_upload=client.post(f"/api/courses/{course['id']}/lesson-plan-outline",files={"outline":("另一版.docx",stream.getvalue())})
    assert next_upload.status_code==200 and next_upload.json()["id"]!=available["id"]
    again=client.post(f"/api/courses/{course['id']}/lesson-plan-outline",files={"outline":("原大纲.docx",data)})
    assert again.status_code==200 and again.json()["id"]==available["id"]
    assert client.get(f"/api/courses/{course['id']}/lesson-plan-outline").json()["source"]["id"]==available["id"]


def test_old_course_requests_raw_upload_without_changing_course(client):
    course=client.post("/api/courses",json={"name":"旧课程"}).json()
    response=client.get(f"/api/courses/{course['id']}/lesson-plan-outline")
    assert response.json()["source"] is None and "上传" in response.json()["message"]
    assert client.get(f"/api/courses/{course['id']}").json()["settings"]==course["settings"]


@pytest.mark.parametrize("filename,data,status",[("bad.xls",b"anything",400),("bad.docx",b"broken",400),("big.docx",b"x"*(5*1024*1024+1),413)], ids=["wrong-extension", "damaged-docx", "too-large"])
def test_upload_validation(client,filename,data,status):
    assert client.post("/api/lesson-plans",files={"outline":(filename,data)}).status_code==status


def test_plan_editing_revision_and_readonly_source_fields(client):
    plan=create_plan(client); changed=copy.deepcopy(plan["context"])
    changed["lessons"][0]["教学目标"]=["手工编辑的目标"]
    response=client.patch(f"/api/lesson-plans/{plan['id']}",json={"revision":plan["revision"],"context":changed})
    assert response.status_code==200,response.text
    assert response.json()["context"]["lessons"][0]["教学目标"]==["手工编辑的目标"]
    assert not response.json()["counts"]["1"]["教学目标"]["in_range"]
    assert client.patch(f"/api/lesson-plans/{plan['id']}",json={"revision":1,"context":changed}).status_code==409
    changed["lessons"][0]["学时"]="999"
    assert client.patch(f"/api/lesson-plans/{plan['id']}",json={"revision":2,"context":changed}).status_code==400


def test_jobs_finish_generate_and_export_without_waiting_http(client,monkeypatch):
    plan=create_plan(client)
    monkeypatch.setattr("web_app.lesson_plans.client.api_key",lambda:"sk-offline-only")
    def answer(system,user,key,temperature=0.5):
        callback=getattr(__import__("web_app.lesson_plans.client",fromlist=["_local"])._local,"callback",None)
        if callback: callback()
        if system.startswith("你是高校教务老师"): return {"评价方式":"过程性与终结性相结合"}
        return valid_ai()
    monkeypatch.setattr(engine,"deepseek",answer)
    monkeypatch.setattr(service,"make_document",lambda context,check=True:(b"PK-synthetic-docx",{1:1,2:2},["服务器字体页数仅供参考"]))
    begun=client.post(f"/api/lesson-plans/{plan['id']}/generate",json={"revision":1})
    assert begun.status_code==202,begun.text
    done=wait_job(client,begun.json()); assert done["status"]=="success",done
    assert done["api_calls"]==3 and done["completed"]==2
    updated=client.get(f"/api/lesson-plans/{plan['id']}").json()
    assert updated["revision"]==2 and updated["pages"]=={"1":1,"2":2}
    assert updated["counts"]["1"]["教学目标"]["in_range"]
    monkeypatch.setattr(engine,"deepseek",Mock(side_effect=AssertionError("export must not call AI")))
    exported=client.post(f"/api/lesson-plans/{plan['id']}/export",json={"revision":2})
    assert wait_job(client,exported.json())["status"]=="success"
    download=client.get(f"/api/lesson-plans/{plan['id']}/download")
    assert download.status_code==200 and download.content==b"PK-synthetic-docx"
    assert "docx" in download.headers["content-disposition"]


def test_auth_ownership_and_origin_protect_all_records(client,app):
    plan=create_plan(client)
    with TestClient(app) as anonymous:
        assert anonymous.get(f"/api/lesson-plans/{plan['id']}").status_code==401
    assert client.post("/api/lesson-plans",files={"outline":("x.docx",syllabus())},headers={"Origin":"https://other.example"}).status_code==403
    assert client.post("/api/register",json={"invite_code":"TST123","username":"lesson-other","password":PASSWORD}).status_code==200
    assert client.post("/api/login",json={"username":"lesson-other","password":PASSWORD}).status_code==200
    assert client.get(f"/api/lesson-plans/{plan['id']}").status_code==404
    assert client.post("/api/lesson-plans",json={"source_id":plan["source_id"]}).status_code==404


def test_no_environment_key_is_actionable_and_manual_word_still_available(client,monkeypatch):
    plan=create_plan(client)
    assert client.post(f"/api/lesson-plans/{plan['id']}/generate",json={"revision":1}).status_code==503
    forbidden=Mock(side_effect=AssertionError("render must not call API")); monkeypatch.setattr(engine,"deepseek",forbidden)
    context=plan["context"]; content,pages,warnings=service.make_document(context,check=False)
    assert content.startswith(b"PK") and pages=={}
    document=docx.Document(io.BytesIO(content))
    texts="\n".join(p.text for p in document.paragraphs)+"\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert "{{" not in texts and "合成教案课程" in texts and "合成教师" not in texts
    forbidden.assert_not_called()


def test_rendered_lesson_table_uses_fangsong_small_four_without_changing_template(tmp_path):
    original_template = service.template_digest()
    context = service.initial_context(source.parse_bytes("大纲.docx", syllabus()))
    context["lessons"][0].update({"授课题目": "第二章 2D 游戏物理系统", "授课类型": "讲授、实践", "学时": "4"})
    context["lessons"][0]["教学目标"] = ["目标正文"]
    context["lessons"][0]["教学内容"] = ["教学正文"]
    context["lessons"][0]["教学手段与方法"] = ["方法正文"]
    context["lessons"][0]["思考讨论作业"] = ["作业正文"]
    output = tmp_path / "教案.docx"
    engine.render(engine.DEFAULT_TEMPLATE, context, str(output))
    document = docx.Document(output)
    lesson_tables = [table for table in document.tables if len(table.rows) == 6]
    assert len(lesson_tables) == len(context["lessons"])
    for table, lesson in zip(lesson_tables, context["lessons"]):
        assert table.rows[0].cells[2].text == lesson["授课题目"]
        assert table.rows[1].cells[1].text.strip() == f"第{lesson['课次']}课次"
        assert table.rows[1].cells[3].text == lesson["授课类型"]
        assert table.rows[1].cells[5].text == lesson["学时"]
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        if run.text.strip():
                            fonts = run._element.rPr.rFonts
                            for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
                                assert fonts.get(qn(f"w:{slot}")) == "仿宋_GB2312"
                            assert run.font.size.pt == 12
    cover_font = document.tables[0].rows[0].cells[1].paragraphs[0].runs[0]
    assert cover_font._element.rPr.rFonts.get(qn("w:eastAsia")) == "华文中宋"
    assert service.template_digest() == original_template


def test_under_minimum_even_inside_original_tolerance_rewrites_once(monkeypatch):
    answer=valid_ai(); answer["教学目标"]=["目"*119]
    call=Mock(side_effect=[answer,valid_ai()]); monkeypatch.setattr(engine,"deepseek",call)
    outline=source.parse_bytes("大纲.docx",syllabus())
    result=engine.gen_lesson(outline,outline["lessons"][0],"sk-offline")
    assert call.call_count==2 and engine.hanzi_len(result["教学目标"])==140


def test_save_original_does_not_run_any_outline_parser(client,monkeypatch):
    course=client.post("/api/courses",json={"name":"只保存原文件"}).json()
    forbidden=Mock(side_effect=AssertionError("saving source must not parse lessons"))
    monkeypatch.setattr(engine,"read_outline",forbidden)
    response=client.post(f"/api/courses/{course['id']}/lesson-plan-outline",files={"outline":("原始.docx",syllabus())})
    assert response.status_code==200,response.text
    forbidden.assert_not_called()


def test_duplicate_job_and_active_edit_protection(client,monkeypatch):
    plan=create_plan(client); hold=threading.Event()
    def document(context,check=True):
        assert hold.wait(6)
        return b"PK-synthetic",{1:1,2:1},[]
    monkeypatch.setattr(service,"make_document",document)
    first=client.post(f"/api/lesson-plans/{plan['id']}/check-pages",json={"revision":1})
    try:
        second=client.post(f"/api/lesson-plans/{plan['id']}/check-pages",json={"revision":1})
        assert first.status_code==second.status_code==202
        assert first.json()["job_id"]==second.json()["job_id"]
        assert client.patch(f"/api/lesson-plans/{plan['id']}",json={"revision":1,"context":plan["context"]}).status_code==409
        assert client.post(f"/api/lesson-plans/{plan['id']}/export",json={"revision":1}).status_code==409
    finally: hold.set()
    assert wait_job(client,first.json())["status"]=="success"


def test_expired_job_resumes_saved_lesson_and_evaluation(client,monkeypatch):
    plan=create_plan(client)
    monkeypatch.setattr("web_app.lesson_plans.client.api_key",lambda:"sk-offline")
    field=engine.ai_to_fields(valid_ai()); calls=[]
    def answer(system,user,key,temperature=0.5):
        calls.append(user)
        callback=getattr(__import__("web_app.lesson_plans.client",fromlist=["_local"])._local,"callback",None)
        if callback: callback()
        return valid_ai()
    monkeypatch.setattr(engine,"deepseek",answer)
    monkeypatch.setattr(service,"make_document",lambda context,check=True:(b"PK-synthetic",{1:1,2:1},[]))
    with Session(get_engine()) as db:
        record=db.get(LessonPlan,plan["id"])
        task=jobs.launch(db,record.user_id,record,"generate")
        row=db.get(LessonPlanTask,task["job_id"])
        row.status="running"; row.claim_token="abandoned"; row.lease_until=utcnow()-timedelta(seconds=1)
        row.checkpoint_json=json.dumps({"lessons":{"1":field},"evaluation":plan["context"]["评价"],"api_calls":2})
        db.commit()
    jobs.start_runtime()
    done=wait_job(client,task)
    assert done["status"]=="success" and len(calls)==1 and done["api_calls"]==3
    assert '"课次": "2"' in calls[0]


def test_failed_compression_preserves_edited_content(client,monkeypatch):
    plan=create_plan(client); edited=copy.deepcopy(plan["context"])
    edited["lessons"][0]["教学目标"]=["用户已编辑的目标，压缩失败必须保留"]
    saved=client.patch(f"/api/lesson-plans/{plan['id']}",json={"revision":1,"context":edited}).json()
    monkeypatch.setattr("web_app.lesson_plans.client.api_key",lambda:"sk-offline")
    monkeypatch.setattr(engine,"deepseek",Mock(side_effect=engine.LessonPlanError("API 调用超时")))
    monkeypatch.setattr(service,"make_document",lambda context,check=True:(b"PK-synthetic",{1:2,2:1},[]))
    started=client.post(f"/api/lesson-plans/{plan['id']}/lessons/1/compress",json={"revision":saved["revision"]})
    assert wait_job(client,started.json())["status"]=="success"
    result=client.get(f"/api/lesson-plans/{plan['id']}").json()
    assert result["context"]["lessons"][0]["教学目标"]==edited["lessons"][0]["教学目标"]
    assert any("压缩失败" in warning and "保留" in warning for warning in result["warnings"])


def test_auto_prepare_reuses_edited_draft_without_parsing_or_calling_ai(client,monkeypatch):
    course=client.post("/api/courses",json={"name":"自动准备教案"}).json()
    saved=client.post(f"/api/courses/{course['id']}/lesson-plan-outline",files={"outline":("原始.docx",syllabus())}).json()
    forbidden_ai=Mock(side_effect=AssertionError("opening plan must not call API"))
    monkeypatch.setattr(engine,"deepseek",forbidden_ai)
    first=client.post("/api/lesson-plans",json={"source_id":saved["id"],"reuse":True})
    assert first.status_code==200,first.text
    plan=first.json(); context=copy.deepcopy(plan["context"])
    context["lessons"][0]["教学目标"]=["教师已编辑的目标，应在再次打开时保留"]
    updated=client.patch(f"/api/lesson-plans/{plan['id']}",json={"revision":plan["revision"],"context":context})
    assert updated.status_code==200,updated.text
    forbidden_parse=Mock(side_effect=AssertionError("existing preview must be reused"))
    monkeypatch.setattr(source,"read_source",forbidden_parse)
    again=client.post("/api/lesson-plans",json={"source_id":saved["id"],"reuse":True})
    assert again.status_code==200 and again.json()["id"]==plan["id"]
    assert again.json()["revision"]==updated.json()["revision"]
    assert again.json()["context"]["lessons"][0]["教学目标"]==context["lessons"][0]["教学目标"]
    forbidden_parse.assert_not_called(); forbidden_ai.assert_not_called()


def test_lesson_client_uses_existing_ten_key_server_pool(monkeypatch,tmp_path):
    from web_app.lesson_plans import client as api_client
    from web_app.deepseek_pool import load_specs,reset_pool
    first=[f"fake-local-pool-a-{i}" for i in range(5)]
    second=[f"fake-local-pool-b-{i}" for i in range(5)]
    a=tmp_path/"keys-a.txt"; b=tmp_path/"keys-b.txt"
    a.write_text("\n".join(first),encoding="utf-8"); b.write_text("\n".join(second),encoding="utf-8")
    monkeypatch.setenv("DEEPSEEK_KEYS_A_FILE",str(a)); monkeypatch.setenv("DEEPSEEK_KEYS_B_FILE",str(b))
    monkeypatch.setenv("DEEPSEEK_API_KEY","fake-unused-single-key")
    reset_pool(); assert sum(len(keys) for _,keys in load_specs())==10
    used=[]
    class OfflineOpener:
        def open(self,request,timeout):
            used.append(request.get_header("Authorization").removeprefix("Bearer "))
            assert timeout==180
            assert json.loads(request.data)["thinking"]=={"type":"disabled"}
            return io.BytesIO(json.dumps({"choices":[{"message":{"content":json.dumps({"评价方式":"合成评价"},ensure_ascii=False)}}]}).encode())
    monkeypatch.setattr("urllib.request.build_opener",lambda *args:OfflineOpener())
    assert api_client.complete_json("你是高校教务老师。","合成验收内容")=={"评价方式":"合成评价"}
    api_client.complete_json("你是高校教务老师。","合成验收内容")
    assert used==[first[0],second[0]]


def test_api_truncation_is_reported_as_incomplete_content_not_timeout(monkeypatch):
    from web_app.lesson_plans import client as api_client
    monkeypatch.setenv("DEEPSEEK_API_KEY","fake-offline-key")
    class OfflineOpener:
        def open(self,request,timeout):
            return io.BytesIO(json.dumps({"choices":[{"finish_reason":"length","message":{"content":""}}]}).encode())
    monkeypatch.setattr("urllib.request.build_opener",lambda *args:OfflineOpener())
    with pytest.raises(engine.LessonPlanError,match="被截断") as error:
        api_client.complete_json("合成 JSON 测试","合成内容")
    assert error.value.status==502 and error.value.extra["code"]=="ai_truncated"


def test_all_ai_calls_failing_does_not_claim_ai_content_is_complete(client,monkeypatch):
    plan=create_plan(client)
    monkeypatch.setattr("web_app.lesson_plans.client.api_key",lambda:"fake-offline")
    monkeypatch.setattr(engine,"deepseek",Mock(side_effect=engine.LessonPlanError("API 返回的教案内容被截断")))
    monkeypatch.setattr(service,"make_document",lambda context,check=True:(b"PK-synthetic",{1:1,2:1},[]))
    started=client.post(f"/api/lesson-plans/{plan['id']}/generate",json={"revision":1})
    done=wait_job(client,started.json())
    assert done["status"]=="success" and "AI 生成未完成" in done["stage_label"]
