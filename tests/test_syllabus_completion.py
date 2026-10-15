"""Legacy format, bounded AI repair, durable imports and shared outline snapshots."""
import copy
import io
import json
import time
from unittest.mock import Mock

import docx
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select
from sqlalchemy.orm import Session
from tests.test_web_flow import app, client, PASSWORD
from web_app.db import Course, Term, LessonPlanSource, SyllabusImportTask, get_engine
from web_app.syllabus import compatibility, completion, client as ai_client, routes
from web_app.service import ServiceError
from web_app.limiter import RateLimiter


@pytest.fixture(autouse=True)
def limits(monkeypatch):
    monkeypatch.setattr(routes, "limiter", RateLimiter(30,60))


def legacy():
    document = docx.Document()
    document.add_paragraph("《合成造型》课程教学大纲")
    table = document.add_table(rows=0, cols=4)
    for row in [("课程中文名称","合成造型","课程类别","专业必修"),
                ("学分","1","学时","4"),("开课单位","合成学院","课程编号","TEST-01")]:
        for cell, value in zip(table.add_row().cells,row):
            cell.text=value
    document.add_paragraph("二、课程简介")
    document.add_paragraph("学习造型设计、材料使用和作品表达。")
    document.add_paragraph("三、课程目标")
    table=document.add_table(rows=0,cols=2)
    for row in [("课程目标","与毕业要求的对应关系"),("目标1：能够掌握基础知识","理论知识（H）"),
                ("目标2：能够分析材料性能","材料工艺（M）")]:
        for cell,value in zip(table.add_row().cells,row): cell.text=value
    table=document.add_table(rows=0,cols=3)
    for row in [("课程目标 毕业要求","理论知识","专业视野"),("支撑度","H","M"),
                ("课程目标1","√",""),("课程目标2","","√")]:
        for cell,value in zip(table.add_row().cells,row): cell.text=value
    document.add_paragraph("六、教学内容及学时分配")
    table=document.add_table(rows=0,cols=6)
    for row in [("课次","学时","章目名称","教学内容及教学目标","达成课程目标","教学方法"),
                ("1","2","第一章 基础知识","教学内容：基础造型。教学目标：掌握基础知识。","课程目标1","讲授法"),
                ("2","2","第二章 设计实践","教学内容：作品设计。教学目标：完成设计作品。","课程目标2 课程目标3","实践")]:
        for cell,value in zip(table.add_row().cells,row): cell.text=value
    document.add_paragraph("七、考核形式与成绩评定")
    table=document.add_table(rows=0,cols=2)
    for row in [("项目","说明"),("成绩评定","课堂参与度（15%）+课后作业成绩（15%）+期末考试成绩（70%）")]:
        for cell,value in zip(table.add_row().cells,row): cell.text=value
    document.add_paragraph("八、参考资料")
    document.add_paragraph("（一）参考书目")
    document.add_paragraph("合成作者.造型基础.合成出版社，2020.")
    document.add_paragraph("（二）网络资源")
    document.add_paragraph("示例资源：https://example.test")
    stream=io.BytesIO();document.save(stream)
    return stream.getvalue()


def answer(analysis):
    draft=analysis["draft"]
    return {"fields":{key:completion.value(draft["fields"].get(key,"")) for key in completion.COURSE_FIELDS},
        "course_description":completion.value(draft["course_description"]),
        "objectives":[goal or "能够完成并评价造型设计作品" for goal in draft["objectives"]],
        "grad_req_map":[{"requirement":completion.value(item["requirement"]) or "专业技能",
                        "indicator":completion.value(item["indicator"]) or "设计表达","strength":item["strength"] or "H"} for item in draft["grad_req_map"]],
        "relation_grid":draft["relation"]["grid"],
        "teaching_calendar":{"lessons":[{"课次":int(row["课次"]),"教学内容":"基础知识与造型设计实践","教学要求":"能够完成并评价设计作品"} for row in analysis["outline"]["lessons"]]},
        "changes":["补齐课程目标3及其对应关系","核对旧版表格的毕业要求矛盾"]}


def wait(client,job):
    deadline=time.monotonic()+10
    while time.monotonic()<deadline:
        response=client.get("/api/syllabus/imports/"+job["job_id"])
        assert response.status_code==200,response.text
        result=response.json()
        if result["done"]: return result
        time.sleep(.05)
    raise AssertionError("syllabus job timeout")


def payload(draft):
    return {"fields":{key:completion.value(value) for key,value in draft["fields"].items()},
        "description":completion.value(draft["course_description"]),"objectives":draft["objectives"],
        "grad_req_map":draft["grad_req_map"],"relation_grid":draft["relation"]["grid"],
        "teaching_calendar":draft["teaching_calendar"],"completion_job_id":draft["completion_job_id"]}


def test_legacy_matrix_and_missing_definition_recover_without_api():
    analysis=compatibility.load("旧版.docx",legacy())
    assert len(analysis["draft"]["objectives"])==3 and not analysis["draft"]["objectives"][-1]
    assert completion.value(analysis["draft"]["grad_req_map"][0]["requirement"])=="理论知识"
    assert any("前后不一致" in issue for issue in analysis["issues"])
    assert analysis["draft"]["relation"]["ratios"]=={"usual":.3,"midterm":0,"final":.7}
    assert len(analysis["outline"]["学习资料"])==len(analysis["outline"]["学习网站"])==1
    assert analysis["outline"]["lessons"][0]["教学目标及教学内容"].startswith("教学目标：")


def test_ai_rewrites_once_when_validation_fails(monkeypatch):
    analysis=compatibility.load("旧版.docx",legacy())
    correct=answer(analysis); bad=copy.deepcopy(correct); bad["objectives"][-1]=""
    call=Mock(side_effect=[{"content":json.dumps(bad,ensure_ascii=False)},{"content":json.dumps(correct,ensure_ascii=False)}])
    monkeypatch.setattr(ai_client,"complete",call)
    draft,outline=completion.run(analysis)
    assert call.call_count==2 and draft["completion"]["ok"]
    assert "objectives.2" in draft["completion"]["inferred_fields"]
    assert outline["考核"]["平时比例"]=="30.0"
    call=Mock(return_value={"content":json.dumps(bad,ensure_ascii=False)})
    monkeypatch.setattr(ai_client,"complete",call)
    draft,_=completion.run(analysis)
    assert call.call_count==2 and not draft["completion"]["ok"]
    assert any("仍有问题" in text for text in draft["completion"]["warnings"])


def test_ai_failure_preserves_editable_data_without_exposing_upstream_details(monkeypatch):
    analysis=compatibility.load("旧版.docx",legacy())
    monkeypatch.setattr(ai_client,"complete",Mock(side_effect=RuntimeError("private-api-key")))
    draft,_=completion.run(analysis)
    assert draft["objectives"][0] and draft["relation"]["grid"]
    assert draft["completion"]["warnings"]
    assert "private-api-key" not in json.dumps(draft)
    exhausted=Mock(side_effect=AssertionError("recovered task must not exceed its persisted API budget"))
    monkeypatch.setattr(ai_client,"complete",exhausted)
    draft,_=completion.run(analysis,call_offset=2)
    assert exhausted.call_count==0 and draft["meta"]["calls"]==2


def test_valid_source_numbers_cannot_be_silently_changed():
    analysis=compatibility.load("旧版.docx",legacy()); data=answer(analysis)
    data["fields"]["hours"]="6"
    with pytest.raises(ServiceError,match="课程信息"): completion.merge(analysis,data)
    data=answer(analysis); data["relation_grid"]=copy.deepcopy(data["relation_grid"]); data["relation_grid"][1][1]="40%"; data["relation_grid"][-1][1]="60%"
    with pytest.raises(ServiceError): completion.merge(analysis,data)


def test_background_import_creates_course_and_reuses_completed_outline(client,monkeypatch):
    analysis=compatibility.load("旧版.docx",legacy())
    call=Mock(return_value={"content":json.dumps(answer(analysis),ensure_ascii=False)})
    monkeypatch.setattr(ai_client,"complete",call)
    started=client.post("/api/syllabus/imports",files={"syllabus":("旧版.docx",legacy())})
    assert started.status_code==202,started.text
    job=wait(client,started.json()); assert job["status"]=="success",job
    draft=job["draft"]; assert draft["completion"]["ok"]
    checked=client.post("/api/syllabus/validate-draft",json=payload(draft))
    assert checked.status_code==200 and checked.json()["ok"],checked.text
    created=client.post("/api/courses/from-syllabus",json=payload(draft))
    assert created.status_code==200,created.text
    course=created.json()
    before=client.get(f"/api/courses/{course['id']}").json()
    old_browser_settings=copy.deepcopy(before["edit_settings"])
    old_browser_settings.pop("teaching_outline",None);old_browser_settings.pop("syllabus_completion",None)
    saved_course=client.patch(f"/api/courses/{course['id']}",json={"settings":old_browser_settings})
    assert saved_course.status_code==200,saved_course.text
    with Session(get_engine()) as db:
        assert json.loads(db.get(Course,course["id"]).settings_json)["teaching_outline"]["completion_id"]==draft["completion_job_id"]
    available=client.get(f"/api/courses/{course['id']}/lesson-plan-outline").json()["source"]
    assert available["outline_version"]==draft["completion_job_id"] and available["lesson_count"]==2
    plan=client.post("/api/lesson-plans",json={"source_id":available["id"]}).json()
    assert len(plan["context"]["_outline"]["课程目标"])==3
    assert plan["context"]["学习网站"]==["示例资源：https://example.test"]
    # Saving cannot replace the server's validated syllabus snapshot.
    context=copy.deepcopy(plan["context"]);context["_outline"]["课程目标"]=["伪造源数据"]
    saved=client.patch("/api/lesson-plans/"+plan["id"],json={"revision":plan["revision"],"context":context}).json()
    assert len(saved["context"]["_outline"]["课程目标"])==3
    with Session(get_engine()) as db:
        db.add(Term(user_id=1,course_id=course["id"],year_start="2025",year_end="2026",semester="2",teacher="成绩单教师",settings_json="{}"))
        db.commit()
    changed_answer=answer(analysis); changed_answer["objectives"][0]="新版目标：能够理解并评价造型原理"
    call.return_value={"content":json.dumps(changed_answer,ensure_ascii=False)}
    second=client.post("/api/syllabus/imports",json={"course_id":course["id"]})
    assert second.status_code==202,second.text
    newer=wait(client,second.json())["draft"]
    assert newer["expected_version"]==1 and newer["source_course_id"]==course["id"]
    updated=client.post(f"/api/courses/{course['id']}/syllabus-version",json={**payload(newer),"expected_version":1})
    assert updated.status_code==200,updated.text
    different=client.post("/api/lesson-plans",json={"source_id":available["id"],"reuse":True}).json()
    assert different["id"]!=plan["id"] and different["outline_version"]==newer["completion_job_id"]
    assert client.get("/api/lesson-plans/"+plan["id"]).json()["outline_version"]==draft["completion_job_id"]
    assert call.call_count==2
    with Session(get_engine()) as db:
        old_term=db.scalar(select(Term).where(Term.course_id==course["id"]))
        frozen=json.loads(old_term.settings_json)["syllabus_settings"]
        assert frozen["objective_requirements"][0]==draft["objectives"][0]
        assert old_term.teacher=="成绩单教师"
    malformed={**payload(newer),"objectives":[{"unexpected":"value"}]}
    response=client.post("/api/syllabus/validate-draft",json=malformed)
    assert response.status_code==200 and not response.json()["ok"]


def test_import_owner_origin_type_and_size_protection(client,app):
    assert client.post("/api/syllabus/imports",files={"syllabus":("x.docx",legacy())},headers={"Origin":"https://other.example"}).status_code==403
    assert client.post("/api/syllabus/imports",files={"syllabus":("x.xls",b"x")}).status_code==400
    assert client.post("/api/syllabus/imports",files={"syllabus":("x.docx",b"x"*(5*1024*1024+1))}).status_code==413
    with TestClient(app) as anonymous:
        assert anonymous.get("/api/syllabus/imports/unknown").status_code==401
    assert client.get("/api/syllabus/imports/unknown").status_code==404
