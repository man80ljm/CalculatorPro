"""课程性质原文来源：大纲课程类别，成绩单课程性质优先。"""
import io
import json

import pytest
from docx import Document
from openpyxl import load_workbook

from tests.test_grade_register import client, _craft_settings, _register_xlsx, _post_register, _FORWARD_HEADERS, _FORWARD_ROWS
from tests.test_syllabus_rules import _ai
from tests.test_word_register import word_from_excel
from web_app.syllabus.rules import build_draft, course_type_field
from web_app.syllabus.service import extract_draft


def syllabus_bytes():
    document = Document()
    document.add_paragraph("《合成课程》课程教学大纲")
    document.add_paragraph("一、课程基本信息")
    table = document.add_table(rows=2, cols=4)
    rows = [["课程中文名称", "合成课程", "适用专业", "合成专业"],
            ["课程模块", "专业拓展课程", "课 程 类 别", "选修"]]
    for row, values in zip(table.rows, rows):
        for cell, value in zip(row.cells, values):
            cell.text = value
    document.add_paragraph("二、课程简介")
    document.add_paragraph("课程类别：正文中的其他类别")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_extract_uses_opening_table_category_and_ignores_model(monkeypatch):
    monkeypatch.setenv("DEEPSEEK_API_KEY", "sk-test")
    monkeypatch.setattr("web_app.syllabus.service.complete", lambda *_args, **_kwargs: {"content": json.dumps(_ai())})
    draft = extract_draft("合成大纲.docx", syllabus_bytes())
    assert draft["fields"]["course_type"]["value"] == "选修"
    assert draft["fields"]["course_type"]["status"] == "已填"
    assert draft["fields"]["course_type"]["source"] == "大纲·课程类别"
    assert not any(issue.get("field") == "course_type" for issue in draft.get("issues", []))


def test_full_original_table_is_used_even_when_excerpt_lacks_category():
    draft = build_draft(_ai(), "课程简介：合成文字", "", None,
                        matrix_text="一、课程基本信息\n| 课程模块 | 专业拓展课程 | 课程类别 | 选修 |\n二、课程简介\n课程类别：其他")
    assert draft["fields"]["course_type"]["value"] == "选修"


@pytest.mark.parametrize("text", ["课程模块：专业拓展课程", "课程性质：专业拓展课程/选修", "一、课程基本信息\n二、课程简介\n课程类别：选修"])
def test_no_category_in_opening_information_stays_manual(text):
    nature = course_type_field(text, "", None)
    assert nature["value"] == ""
    assert nature["status"] == "需手填"
    assert "课程类别" in nature["reason"]


@pytest.mark.parametrize("format", ["xlsx", "docx"])
@pytest.mark.parametrize("nature", ["专业任选课", "", None])
def test_import_register_overrides_only_nonempty_nature(client, format, nature):
    settings = _craft_settings()
    settings["course_basic_info"]["course_type"] = "选修"
    course_id = client.post("/api/courses", json={"name": "创意手作", "settings": settings}).json()["id"]
    book = load_workbook(io.BytesIO(_register_xlsx(_FORWARD_HEADERS, _FORWARD_ROWS)))
    suffix = "" if nature is None else " 课程性质：" + nature
    book.active["A2"] = "课程名称：创意手作 课程代码：BJ2230005" + suffix
    buffer = io.BytesIO()
    book.save(buffer)
    content = buffer.getvalue()
    if format == "docx":
        content = word_from_excel(content)
    response = _post_register(client, course_id, content, filename="合成登记." + format)
    assert response.status_code == 200, response.text
    basic = client.get(f"/api/courses/{course_id}").json()["settings"]["course_basic_info"]
    assert basic["course_type"] == (nature or "选修")
    assert ("课程性质" in response.json()["filled"]) is bool(nature)
    # 再导入未写性质的成绩单，仍保留刚刚确认的值。
    book.active["A2"] = "课程名称：创意手作 课程代码：BJ2230005"
    buffer = io.BytesIO()
    book.save(buffer)
    repeated = _post_register(client, course_id, buffer.getvalue(), confirm=True)
    assert repeated.status_code == 200, repeated.text
    assert client.get(f"/api/courses/{course_id}").json()["settings"]["course_basic_info"]["course_type"] == (nature or "选修")


def test_word_register_nature_wins_in_combined_syllabus_draft(monkeypatch):
    monkeypatch.setenv("DEEPSEEK_API_KEY", "sk-test")
    monkeypatch.setattr("web_app.syllabus.service.complete", lambda *_args, **_kwargs: {"content": json.dumps(_ai())})
    register = word_from_excel(_register_xlsx(_FORWARD_HEADERS, _FORWARD_ROWS))
    draft = extract_draft("合成大纲.docx", syllabus_bytes(), "合成登记.docx", register)
    nature = draft["fields"]["course_type"]
    assert nature["value"] == "专业任选课"
    assert nature["source"] == "成绩登记表·课程性质"
