"""Word 与 Excel 同数据识别一致；相同随机种子下计算和表 5 一致。"""
import io
import zipfile

import pytest
from docx import Document
from openpyxl import load_workbook

from tests.test_web_flow import app, client
from tests.test_previous_term_attainment import _course
from tests.test_grade_register import _FORWARD_HEADERS, _FORWARD_ROWS, _post_register, _register_xlsx
from tests.test_register_headcounts import HEADERS, STUDENTS, register_bytes
from web_app.grade_register import parse_register
from web_app.service import ServiceError


def word_from_excel(content, split=False):
    book = load_workbook(io.BytesIO(content), data_only=True)
    document = Document()
    table = None
    width = book.active.max_column
    for row in book.active.values:
        values = [str(value) if value is not None else "" for value in row]
        if table is None and "姓名" not in values:
            if values[0]:
                document.add_paragraph(values[0])
            continue
        if table is None or (split and values[0] == "成绩统计"):
            table = document.add_table(rows=0, cols=width)
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.mark.parametrize("content", [
    register_bytes(STUDENTS, ["实考人数：2人 总人数3人"]),
    _register_xlsx(_FORWARD_HEADERS, _FORWARD_ROWS),
])
def test_word_and_excel_parse_identically(content):
    assert parse_register("合成登记.DOCX", word_from_excel(content)) == parse_register("合成登记.xlsx", content)


def test_word_merged_class_and_repeated_tables_keep_students():
    document = Document()
    document.add_paragraph("2026-2027学年第1学期 课程名称：合成课程")
    first = document.add_table(rows=3, cols=len(HEADERS))
    for cells, values in zip((row.cells for row in first.rows), [HEADERS] + STUDENTS[:2]):
        for cell, value in zip(cells, values):
            cell.text = str(value)
    first.cell(1, 0).merge(first.cell(2, 0)).text = "合成班"
    document.add_paragraph("续表")
    second = document.add_table(rows=3, cols=len(HEADERS))
    for cells, values in zip((row.cells for row in second.rows), [HEADERS, [""] * len(HEADERS), STUDENTS[2]]):
        for cell, value in zip(cells, values):
            cell.text = str(value)
    document.add_paragraph("成绩统计")
    document.add_paragraph("考核人数：2人")
    buffer = io.BytesIO()
    document.save(buffer)
    parsed = parse_register("合成.docx", buffer.getvalue())
    assert parsed["student_count"] == 3 and parsed["exam_count"] == 2
    assert parsed["classes"] == ["合成班"]
    assert [student["name"] for student in parsed["students"]] == ["合成甲", "合成乙", "合成丙"]
    assert parsed["school_year_term"] == "2026-2027学年第1学期"


def test_word_spaced_headers_and_fullwidth_percent_match_excel():
    content = register_bytes(STUDENTS, ["实考人数：2人"])
    document = Document(io.BytesIO(word_from_excel(content)))
    spaced = ["班 级", "学 号", "姓 名", "平 时（30％）", "期 末（70％）", "总 评", "备注"]
    for cell, text in zip(document.tables[0].rows[0].cells, spaced):
        cell.text = text
    buffer = io.BytesIO()
    document.save(buffer)
    assert parse_register("空格表头.docx", buffer.getvalue()) == parse_register("合成.xlsx", content)


@pytest.mark.parametrize("filename, content, message", [
    ("坏文件.docx", b"not a word document", "无法读取这份 Word"),
    ("旧格式.doc", b"legacy doc", "另存为 .docx"),
])
def test_bad_or_legacy_word_has_clear_error(filename, content, message):
    with pytest.raises(ServiceError, match=message):
        parse_register(filename, content)


def test_word_without_table_has_clear_error():
    document = Document()
    document.add_paragraph("这是合成文字，没有学生成绩表格。")
    buffer = io.BytesIO()
    document.save(buffer)
    with pytest.raises(ServiceError, match="没有找到表格"):
        parse_register("纯文字.docx", buffer.getvalue())


def _table5(client, course_id):
    response = client.post(f"/api/courses/{course_id}/export")
    assert response.status_code == 200, response.text
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        filename = next(name for name in archive.namelist() if name.startswith("5.") and name.endswith(".docx"))
        doc = Document(io.BytesIO(archive.read(filename)))
    return [[[cell.text for cell in row.cells] for row in table.rows] for table in doc.tables]


@pytest.mark.parametrize("mode", ["forward", "reverse"])
def test_word_import_matches_excel_calculation_and_table5(client, mode, monkeypatch):
    if mode == "reverse":
        # 现有种子包含 XLSX 原始字节，重写时的 ZIP 时间戳会变化。
        # 固定两次的随机种子，隔离文件格式之外的已有随机行为；不修改生产计算。
        monkeypatch.setattr("web_app.service.attainment_seed", lambda _path, _settings: 12345)
    course_id = _course(client)
    excel = _register_xlsx(_FORWARD_HEADERS, _FORWARD_ROWS)
    imported = _post_register(client, course_id, word_from_excel(excel), filename="合成登记.docx", mode=mode, confirm=True)
    assert imported.status_code == 200, imported.text
    calculated = client.post(f"/api/courses/{course_id}/calculate")
    assert calculated.status_code == 200, calculated.text
    before = calculated.json()
    word_table = _table5(client, course_id)
    replaced = _post_register(client, course_id, excel, mode=mode, confirm=True)
    assert replaced.status_code == 200, replaced.text
    calculated = client.post(f"/api/courses/{course_id}/calculate")
    assert calculated.status_code == 200, calculated.text
    after = calculated.json()
    assert before["term_id"] == after["term_id"]
    for key in ("student_count", "average_score", "achievement"):
        assert before[key] == after[key]
    assert word_table == _table5(client, course_id)
