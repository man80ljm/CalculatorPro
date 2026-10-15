"""教学日历：规则兼容、可编辑数据、登录保护、真实格式转换及文件清理。"""
import io
import json
import tempfile
from datetime import date, datetime, timezone
from pathlib import Path
from unittest.mock import Mock

import docx
import openpyxl
import pytest
import requests
import xlrd
from fastapi.testclient import TestClient
from openpyxl.styles import Alignment, Border, Font, Side

from tests.test_web_flow import app, client
from web_app.teaching_calendar import conversion, engine, service
from web_app.teaching_calendar.errors import CalendarError
from web_app.teaching_calendar.academic_term import current_academic_term


def outline_bytes(*, table=True):
    document = docx.Document()
    book = document.add_table(rows=1, cols=2)
    book.cell(0, 0).text = "教材选用"
    book.cell(0, 1).text = "程序设计基础，第三版，张三，高等教育出版社，2024"
    if table:
        lessons = document.add_table(rows=3, cols=5)
        headers = ["课次", "学时", "章目名称", "教学目标及教学内容", "教学方式"]
        for cell, value in zip(lessons.rows[0].cells, headers):
            cell.text = value
        for number, row in enumerate(lessons.rows[1:], 1):
            values = [str(number), "2", f"第{number}章", f"教学目标：掌握知识{number}；\n教学内容：知识{number}与实践。", "讲授"]
            for cell, value in zip(row.cells, values):
                cell.text = value
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def template_bytes(count=3, total=6):
    workbook = openpyxl.Workbook()
    ws = workbook.active
    ws.title = "教学日历"
    ws.merge_cells("A1:U1")
    ws["A1"] = "教务系统教学日历"
    ws.merge_cells("A2:U2")
    ws["A2"] = "2024-2025 学年第 2 学期"
    ws["A2"].font = Font(name="Microsoft YaHei", size=12)
    ws["A2"].alignment = Alignment(horizontal="center")
    ws["S4"] = total
    for area in ("C5:J5", "K5:P5", "Q5:S5", "T5:U5"):
        ws.merge_cells(area)
    for ref in engine.BOOK_CELLS.values():
        ws[ref] = "原模板教材信息"
    for name, col in engine.COL.items():
        ws.cell(7, col, name)
    thin = Side(style="thin", color="333333")
    for index in range(count):
        row = engine.FIRST_ROW + index
        ws.cell(row, 1, index + 1)
        ws.cell(row, 2, "2")
        ws.row_dimensions[row].height = 24
        for col in range(1, 22):
            cell = ws.cell(row, col)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.font = Font(name="Microsoft YaHei", size=10)
            cell.alignment = Alignment(vertical="center")
        ws.merge_cells(f"D{row}:N{row}")
        ws.merge_cells(f"O{row}:T{row}")
    footer = engine.FIRST_ROW + count + 1
    ws.merge_cells(f"A{footer}:C{footer}")
    ws.merge_cells(f"D{footer}:N{footer}")
    ws.merge_cells(f"O{footer}:U{footer}")
    ws.cell(footer, 1, "任课教师签字")
    ws.cell(footer, 4, "教研室主任签字")
    ws.cell(footer, 15, "学院审核")
    ws.row_dimensions[footer].height = 42
    ws.row_dimensions[footer].hidden = False
    ws.merge_cells(start_row=footer + 1, start_column=1, end_row=footer + 2, end_column=21)
    ws.cell(footer + 1, 1, "签字日期：")
    ws.row_dimensions[footer + 1].height = 32
    ws.print_area = f"A1:U{footer + 2}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.column_dimensions["A"].width = 9
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def preview_request(client, *, outline=None, template=None, use_ai=False):
    return client.post("/api/teaching-calendar/preview", files={
        "outline": ("大纲.docx", outline if outline is not None else outline_bytes()),
        "template": ("模板.xlsx", template if template is not None else template_bytes()),
    }, data={"use_ai": str(use_ai).lower()})


def payload(count=2):
    return {"教材": {"教材名称": "程序设计基础", "版次": "第3版", "编者": "张三", "出版社": "高等教育出版社"},
            "lessons": [{"课次": number, "学时": "2", "教材章节": f"第{number}章",
                         "教学内容": "知识与实践\n第二行", "教学要求": "掌握基本知识", "备注": "讲授"}
                        for number in range(1, count + 1)]}


def generate_request(client, data=None, output_format="xlsx", template=None):
    return client.post("/api/teaching-calendar/generate", files={
        "template": ("模板.xlsx", template if template is not None else template_bytes()),
    }, data={"payload": json.dumps(data or payload(), ensure_ascii=False), "format": output_format})


def test_pages_and_api_require_teacher_login(app):
    with TestClient(app) as anonymous:
        response = anonymous.get("/teaching-calendar", follow_redirects=False)
        assert response.status_code == 303
        assert response.headers["location"] == "/login"
        assert preview_request(anonymous).status_code == 401
        assert generate_request(anonymous).status_code == 401


def test_page_entry_styles_and_asset_versions(client):
    home = client.get("/").text
    assert home.index('id="newCourseBtn"') < home.index('id="calendarBtn"') < home.index('id="deleteCourseBtn"')
    assert 'id="who"' not in home
    response = client.get("/teaching-calendar", follow_redirects=False)
    assert response.status_code == 303
    assert response.headers["location"] == "/"
    assert 'id="calendarFeedback"' not in home and 'id="calendarXlsxBtn"' not in home
    assert 'id="calendarError"' in home
    assert "教学日历已下载" not in home and "下载 .xlsx" not in home
    script = client.get("/static/app.js").text
    assert 'location.href = "/teaching-calendar"' not in script
    assert "downloadTeachingCalendar" in script


def test_rule_preview_opens_docx_once_and_reports_mismatch(client, monkeypatch):
    document_bytes = outline_bytes()
    original = docx.Document
    opened = Mock(side_effect=original)
    monkeypatch.setattr(service.docx, "Document", opened)
    response = preview_request(client, outline=document_bytes)
    assert response.status_code == 200, response.text
    data = response.json()
    assert opened.call_count == 1
    assert data["mode"] == "rule"
    assert data["教材"] == payload()["教材"]
    assert data["lessons"][0] == {"课次": 1, "学时": "2", "教材章节": "第1章",
                                   "教学内容": "知识1与实践", "教学要求": "掌握知识1", "备注": "讲授"}
    assert data["total_hours"] == 4 and data["template_total_hours"] == 6
    assert data["hours_match"] is False and data["missing_merges"] == []
    assert any("不一致" in warning for warning in data["warnings"])
    assert "_原文" not in response.text


def test_missing_outline_table_returns_error_and_service_stays_alive(client):
    response = preview_request(client, outline=outline_bytes(table=False))
    assert response.status_code == 400
    assert response.json()["code"] == "outline_table_missing"
    assert client.get("/healthz").status_code == 200
    assert preview_request(client).status_code == 200


def test_ai_missing_key_can_fall_back(client):
    response = preview_request(client, use_ai=True)
    assert response.status_code == 503
    assert response.json()["fallback_available"] is True
    assert response.json()["code"] == "ai_not_configured"
    assert preview_request(client).status_code == 200


def test_ai_success_only_rewrites_supported_fields(client, monkeypatch):
    monkeypatch.setenv("DEEPSEEK_API_KEY", "sk-calendar-test-secret")
    response = Mock()
    response.json.return_value = {"choices": [{"message": {"content": json.dumps({
        "教材": {"教材名称": "AI教材"},
        "lessons": [{"课次": 1, "教学内容": "1.知识", "教学要求": "1.掌握", "学时": "99"}],
    }, ensure_ascii=False)}}]}
    call = Mock(return_value=response)
    monkeypatch.setattr(engine.requests, "post", call)
    result = preview_request(client, use_ai=True)
    assert result.status_code == 200, result.text
    data = result.json()
    assert data["教材"]["教材名称"] == "AI教材"
    assert data["教材"]["出版社"] == "高等教育出版社"
    assert data["lessons"][0]["教学内容"] == "1.知识"
    assert data["lessons"][0]["学时"] == "2"
    assert data["lessons"][1]["教学内容"] == "知识2与实践"
    assert any("AI 没有返回" in message for message in data["warnings"])
    assert call.call_args.kwargs["timeout"] == (10, 120)
    assert call.call_args.kwargs["json"]["response_format"] == {"type": "json_object"}
    assert call.call_args.kwargs["headers"]["Authorization"] == "Bearer sk-calendar-test-secret"
    assert "sk-calendar-test-secret" not in result.text


@pytest.mark.parametrize("failure,code,status", [
    (requests.Timeout("secret upstream"), "ai_timeout", 504),
    (requests.HTTPError("Bearer sk-private-key"), "ai_failed", 502),
])
def test_ai_failure_is_sanitized_and_rule_retry_works(client, monkeypatch, failure, code, status):
    monkeypatch.setenv("DEEPSEEK_API_KEY", "sk-private-key")
    monkeypatch.setattr(engine.requests, "post", Mock(side_effect=failure))
    response = preview_request(client, use_ai=True)
    assert response.status_code == status
    assert response.json()["code"] == code
    assert response.json()["fallback_available"] is True
    assert "sk-private-key" not in response.text and "secret upstream" not in response.text
    assert preview_request(client).status_code == 200


def test_malformed_ai_response_is_business_error(client, monkeypatch):
    monkeypatch.setenv("DEEPSEEK_API_KEY", "sk-test")
    upstream = Mock()
    upstream.json.return_value = {"choices": [{"message": {"content": "not JSON"}}]}
    monkeypatch.setattr(engine.requests, "post", Mock(return_value=upstream))
    response = preview_request(client, use_ai=True)
    assert response.status_code == 502 and response.json()["code"] == "ai_failed"


def test_ai_rate_limit_does_not_block_rules(client):
    for _ in range(3):
        assert preview_request(client, use_ai=True).status_code == 503
    response = preview_request(client, use_ai=True)
    assert response.status_code == 429 and response.json()["fallback_available"] is True
    assert preview_request(client).status_code == 200


def test_five_mb_limit_independent_of_other_module_limit(client, monkeypatch):
    monkeypatch.setenv("MAX_UPLOAD_MB", "50")
    response = preview_request(client, outline=b"a" * (service.MAX_FILE_BYTES + 1))
    assert response.status_code == 413 and "5MB" in response.json()["detail"]


@pytest.mark.parametrize("name,content", [("模板.xls", b"PKfake"), ("模板.xlsx", outline_bytes()), ("模板.xlsm", b"fake")],
                         ids=["fake-xls", "docx-as-xlsx", "unsupported-extension"])
def test_extension_and_content_checks(client, name, content):
    response = client.post("/api/teaching-calendar/preview", files={
        "outline": ("大纲.docx", outline_bytes()), "template": (name, content),
    }, data={"use_ai": "false"})
    assert response.status_code == 400


def test_wrong_template_merges_return_error(client):
    workbook = openpyxl.load_workbook(io.BytesIO(template_bytes()))
    workbook.active.unmerge_cells("D8:N8")
    output = io.BytesIO()
    workbook.save(output)
    response = preview_request(client, template=output.getvalue())
    assert response.status_code == 400 and response.json()["code"] == "invalid_template"


@pytest.mark.parametrize("count", [1, 3, 5])
def test_edited_export_preserves_cells_merges_signature_and_print_area(client, monkeypatch, count):
    monkeypatch.setattr(engine, "call_deepseek", Mock(side_effect=AssertionError("download must not call AI")))
    data = payload(count)
    data["教材"]["版次"] = ""
    data["lessons"][0]["教学内容"] = "=这是一段教学文本\n修改后的内容"
    data["lessons"][0]["学时"] = "1.5"
    response = generate_request(client, data)
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "filename*=UTF-8''" in response.headers["content-disposition"]
    assert response.headers["cache-control"] == "no-store"
    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    ws = wb.active
    assert ws["C5"].value == "程序设计基础" and ws["K5"].value is None
    assert ws["Q5"].value == "张三" and ws["T5"].value == "高等教育出版社"
    assert ws["B8"].value == 1.5 and ws["S4"].value == 6
    assert ws["B8"].data_type == "n" and ws["B8"].number_format == "General"
    assert ws["D8"].value == data["lessons"][0]["教学内容"] and ws["D8"].data_type == "s"
    assert ws["D8"].alignment.wrap_text is True
    assert ws["D8"].font.name == "Microsoft YaHei"
    for index in range(count):
        row = engine.FIRST_ROW + index
        assert ws.cell(row, 1).value == str(index + 1)
        assert ws.cell(row, 1).data_type == "s" and ws.cell(row, 1).number_format == "@"
        assert ws.cell(row, 2).data_type == "n" and ws.cell(row, 2).number_format == "General"
        assert f"D{row}:N{row}" in ws.merged_cells
        assert f"O{row}:T{row}" in ws.merged_cells
        assert ws.cell(row, 4).border.bottom.style == "thin"
    footer = engine.FIRST_ROW + count + 1
    assert ws.cell(footer, 1).value == "任课教师签字"
    assert ws.cell(footer, 4).value == "教研室主任签字"
    assert ws.cell(footer, 15).value == "学院审核"
    assert f"A{footer}:C{footer}" in ws.merged_cells
    assert f"A{footer + 1}:U{footer + 2}" in ws.merged_cells
    assert ws.row_dimensions[footer].height == 42
    assert ws.row_dimensions[footer + 1].height == 32
    assert f"$U${footer + 2}" in ws.print_area
    assert ws.page_setup.orientation == "landscape" and ws.column_dimensions["A"].width == 9


def test_template_total_formula_is_preserved(client):
    response = generate_request(client, template=template_bytes(total="=2*3"))
    assert response.status_code == 200, response.text
    ws = openpyxl.load_workbook(io.BytesIO(response.content)).active
    assert ws["S4"].value == "=2*3" and ws["S4"].data_type == "f"


@pytest.mark.parametrize("field,value,code", [("学时", "NaN", "invalid_hours"), ("学时", "-2", "invalid_hours"),
                                            ("学时", "two", "invalid_hours"), ("学时", "0x10", "invalid_hours"),
                                            ("学时", "1_000", "invalid_hours"), ("课次", "0", "invalid_lesson_number")])
def test_edited_numeric_fields_are_validated(client, field, value, code):
    data = payload()
    data["lessons"][0][field] = value
    response = generate_request(client, data)
    assert response.status_code == 400 and response.json()["code"] == code


def test_empty_lessons_duplicate_numbers_and_blank_hours(client):
    data = payload()
    data["lessons"] = []
    assert generate_request(client, data).json()["code"] == "empty_lessons"
    data = payload()
    data["lessons"][1]["课次"] = 1
    assert generate_request(client, data).json()["code"] == "invalid_lesson_number"
    data = payload()
    data["lessons"][0]["学时"] = " "
    assert generate_request(client, data).status_code == 200


def test_cross_origin_upload_is_rejected(client):
    response = client.post("/api/teaching-calendar/generate", files={"template": ("模板.xlsx", template_bytes())},
                           data={"payload": json.dumps(payload()), "format": "xlsx"},
                           headers={"Origin": "https://another-site.example"})
    assert response.status_code == 403 and response.json()["code"] == "invalid_origin"


def test_request_body_limit(client):
    response = client.post("/api/teaching-calendar/preview", content=b"", headers={
        "Content-Type": "multipart/form-data; boundary=123", "Content-Length": str(12 * 1024 * 1024),
    })
    assert response.status_code == 413


def test_temporary_directories_removed_after_success_and_failure(client, monkeypatch, tmp_path):
    created = []
    original = tempfile.TemporaryDirectory
    def temporary(**options):
        result = original(dir=tmp_path, **options)
        created.append(Path(result.name))
        return result
    monkeypatch.setattr(service, "TemporaryDirectory", temporary)
    assert preview_request(client).status_code == 200
    assert generate_request(client).status_code == 200
    assert preview_request(client, outline=outline_bytes(table=False)).status_code == 400
    assert len(created) == 3 and all(not path.exists() for path in created)


def test_conversion_timeout_is_business_error_and_cleans_profile(monkeypatch, tmp_path):
    source = tmp_path / "template.xls"
    source.write_bytes(b"fixture")
    monkeypatch.setattr(conversion, "soffice_path", lambda: "soffice")
    process = Mock()
    process.communicate.side_effect = conversion.subprocess.TimeoutExpired("soffice", 60)
    monkeypatch.setattr(conversion.subprocess, "Popen", Mock(return_value=process))
    stopped = Mock()
    monkeypatch.setattr(conversion, "_stop", stopped)
    with pytest.raises(CalendarError) as error:
        conversion.convert_file(source, "xlsx", tmp_path)
    assert error.value.status == 504 and error.value.extra["code"] == "conversion_timeout"
    stopped.assert_called_once_with(process)
    assert list(tmp_path.iterdir()) == [source]


def test_unavailable_converter_returns_error_and_xlsx_remains_available(client, monkeypatch):
    monkeypatch.setenv("LIBREOFFICE_BIN", "C:/nonexistent/soffice")
    response = generate_request(client, output_format="xls")
    assert response.status_code == 503 and response.json()["code"] == "converter_unavailable"
    assert generate_request(client).status_code == 200


def test_real_libreoffice_xls_input_and_export_preserve_layout(client, tmp_path):
    try:
        conversion.soffice_path()
    except CalendarError:
        pytest.skip("LibreOffice not installed; actual conversion requires deployment/runtime verification")
    source = tmp_path / "source.xlsx"
    source.write_bytes(template_bytes())
    xls = conversion.convert_file(source, "xls", tmp_path).read_bytes()
    assert xls.startswith(service.OLE_MAGIC)
    preview = client.post("/api/teaching-calendar/preview", files={
        "outline": ("大纲.docx", outline_bytes()), "template": ("模板.xls", xls),
    }, data={"use_ai": "false"})
    assert preview.status_code == 200, preview.text
    assert preview.json()["total_hours"] == 4
    for count in (1, 5):
        response = generate_request(client, payload(count), output_format="xls")
        assert response.status_code == 200, response.text
        book = xlrd.open_workbook(file_contents=response.content, formatting_info=True)
        ws = book.sheet_by_index(0)
        assert ws.cell_value(4, 2) == "程序设计基础"
        assert ws.cell_value(1, 0) == current_academic_term()["template_text"]
        assert (1, 2, 0, 21) in ws.merged_cells
        assert ws.cell_value(7, 3) == "知识与实践\n第二行"
        footer = engine.FIRST_ROW + count + 1
        assert ws.cell_value(footer - 1, 0) == "任课教师签字"
        assert ws.cell_value(footer - 1, 3) == "教研室主任签字"
        assert (footer - 1, footer, 0, 3) in ws.merged_cells
        for index in range(count):
            row = 7 + index
            assert ws.cell_type(row, 0) == xlrd.XL_CELL_TEXT
            assert ws.cell_type(row, 1) == xlrd.XL_CELL_NUMBER
            assert ws.cell_value(row, 1) == 2
            cell_format = book.xf_list[ws.cell_xf_index(row, 1)]
            assert book.format_map[cell_format.format_key].format_str == "General"
            assert (row, row + 1, 3, 14) in ws.merged_cells
            assert (row, row + 1, 14, 20) in ws.merged_cells


@pytest.mark.parametrize("hours,expected,value_type", [("4", 4, int), (4, 4, int), ("1.5", 1.5, float),
                                                       ("4.0", 4.0, float), ("2e0", 2, int),
                                                       ("2e-1", 0.2, float), ("", None, type(None)),
                                                       (" ", None, type(None)), (None, None, type(None))])
def test_fill_writes_hours_as_numbers_and_keeps_other_fields_text(hours, expected, value_type):
    ws = openpyxl.Workbook().active
    rows = payload(1)["lessons"]
    rows[0]["学时"] = hours
    engine.fill(ws, rows)
    assert ws["B8"].value == expected and type(ws["B8"].value) is value_type
    assert ws["B8"].data_type == "n" and ws["B8"].number_format == "General"
    for field, column in engine.COL.items():
        if field != "学时":
            cell = ws.cell(8, column)
            assert cell.value == str(rows[0][field]) and cell.data_type == "s"
            assert cell.number_format == "@"


def test_generate_preserves_numeric_integer_decimal_and_empty_hours(client):
    data = payload(3)
    for row, hours in zip(data["lessons"], ("4", "1.5", "")):
        row["学时"] = hours
    response = generate_request(client, data)
    assert response.status_code == 200, response.text
    ws = openpyxl.load_workbook(io.BytesIO(response.content)).active
    assert ws["B8"].value == 4 and type(ws["B8"].value) is int
    assert ws["B9"].value == 1.5 and type(ws["B9"].value) is float
    assert ws["B10"].value is None
    for row in range(8, 11):
        assert ws.cell(row, 2).data_type == "n" and ws.cell(row, 2).number_format == "General"
        assert ws.cell(row, 1).value == str(row - 7) and ws.cell(row, 1).data_type == "s"


def test_real_xls_integer_decimal_and_empty_hours(client):
    try:
        conversion.soffice_path()
    except CalendarError:
        pytest.skip("LibreOffice not installed")
    data = payload(3)
    for row, hours in zip(data["lessons"], ("4", "1.5", "")):
        row["学时"] = hours
    response = generate_request(client, data, output_format="xls")
    assert response.status_code == 200, response.text
    book = xlrd.open_workbook(file_contents=response.content, formatting_info=True)
    sheet = book.sheet_by_index(0)
    assert sheet.cell_type(7, 1) == xlrd.XL_CELL_NUMBER and sheet.cell_value(7, 1) == 4
    assert sheet.cell_type(8, 1) == xlrd.XL_CELL_NUMBER and sheet.cell_value(8, 1) == 1.5
    assert sheet.cell_type(9, 1) in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}
    for row in range(7, 10):
        assert sheet.cell_type(row, 0) == xlrd.XL_CELL_TEXT
        cell_format = book.xf_list[sheet.cell_xf_index(row, 1)]
        assert book.format_map[cell_format.format_key].format_str == "General"


@pytest.mark.parametrize("today,start,semester", [
    (date(2026, 1, 31), "2025", "1"),
    (date(2026, 2, 1), "2025", "2"),
    (date(2026, 8, 31), "2025", "2"),
    (date(2026, 9, 1), "2026", "1"),
    (date(2026, 10, 12), "2026", "1"),
    (date(2027, 1, 1), "2026", "1"),
    (date(2027, 2, 1), "2026", "2"),
])
def test_current_academic_term_boundaries(today, start, semester):
    term = current_academic_term(today)
    assert term["year_start"] == start
    assert term["year_end"] == str(int(start) + 1)
    assert term["semester"] == semester


def test_academic_term_uses_china_time_at_boundary():
    term = current_academic_term(datetime(2026, 8, 31, 20, 0, tzinfo=timezone.utc))
    assert term["year_start"] == "2026" and term["semester"] == "1"


def test_download_recomputes_term_after_preview_and_keeps_template_layout(client, monkeypatch):
    monkeypatch.setattr(service, "current_academic_term", lambda: current_academic_term(date(2026, 8, 31)))
    preview = preview_request(client)
    assert preview.status_code == 200, preview.text
    data = preview.json()
    assert data["academic_term"]["label"] == "2025-2026学年第2学期"
    monkeypatch.setattr(service, "current_academic_term", lambda: current_academic_term(date(2026, 9, 1)))
    response = generate_request(client, data)
    assert response.status_code == 200, response.text
    ws = openpyxl.load_workbook(io.BytesIO(response.content)).active
    assert ws["A2"].value == "2026-2027 学年第 1 学期"
    assert "A2:U2" in ws.merged_cells
    assert ws["A2"].font.name == "Microsoft YaHei" and ws["A2"].font.sz == 12
    assert ws["A2"].alignment.horizontal == "center"
    assert ws["A1"].value == "教务系统教学日历" and ws["S4"].value == 6
    assert ws["B8"].data_type == "n" and ws["B8"].number_format == "General"
    assert response.headers["X-Calendar-School-Year"] == "2026-2027"
    assert response.headers["X-Calendar-Semester"] == "1"
    assert "2026-2027" in response.headers["Content-Disposition"]
