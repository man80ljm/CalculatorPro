"""同一学期、同一成绩、同一组设置下，计算、导出、报告的达成度必须相同。"""
import io
import random
import zipfile

import numpy as np
from docx import Document
from openpyxl import Workbook

from web_app.service import run_ai_report, run_calculation, run_export

RELATION = {
    "objectives_count": 2,
    "links": [
        {
            "name": "平时考核",
            "ratio": 0.3,
            "methods": [{"name": "平时作业", "supports": {"课程目标1": 0.5, "课程目标2": 0.5}, "subtotal": 1.0}],
        },
        {
            "name": "期末考核",
            "ratio": 0.7,
            "methods": [{"name": "期末考试", "supports": {"课程目标1": 0.4, "课程目标2": 0.6}, "subtotal": 1.0}],
        },
    ],
}


def _settings(term_id=9):
    data = {
        "mode": "reverse",
        "term_id": term_id,
        "student_count": 3,
        "course_open_info": {"course_name": "种子课", "semester": "2", "term": "2"},
        "course_basic_info": {"course_name": "种子课", "course_code": "CS9"},
        "ratios": {"usual": 0.3, "midterm": 0, "final": 0.7},
        "relation_payload": RELATION,
        "course_description": "用于检查拆分是否可重复。",
        "objective_requirements": ["会做", "会讲"],
        "spread_mode": "中跨度（7-13分）",
        "distribution": "标准正态",
        "noise_config": {
            "noise_ratio": 0.4,
            "severity_mode": "random",
            "allowed_items": ["平时作业", "期末考试"],
        },
        "report_style": "专业",
        "word_limit": 40,
    }
    return data


def _grades(usual: float, final: float) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.append(["姓名", "平时考核", "期末考核"])
    for index in range(3):
        sheet.append([f"学员{index}", usual, final])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _objectives(achievement: dict) -> dict:
    return {key: round(float(value), 3) for key, value in achievement.items() if "总" not in key}


def _table5(blob: bytes) -> dict:
    archive = zipfile.ZipFile(io.BytesIO(blob))
    name = next(item for item in archive.namelist() if item.endswith("5.基于考核结果的课程目标达成情况评价结果表.docx"))
    document = Document(io.BytesIO(archive.read(name)))
    found = {}
    for table in document.tables:
        if not table.rows:
            continue
        header = "".join(cell.text for cell in table.rows[0].cells)
        if "分目标达成值" not in header:
            continue
        for row in table.rows[1:]:
            cells = [cell.text.strip() for cell in row.cells]
            label = cells[0]
            if label.startswith("课程目标") and label[4:5].isdigit():
                found[label] = round(float(cells[5]), 3)
            elif label == "课程目标达成值":
                found["总达成度"] = round(float(cells[6]), 3)
    return found


def test_calculation_export_and_report_share_attainment(monkeypatch):
    monkeypatch.setenv("DEEPSEEK_API_KEY", "test-key")
    monkeypatch.setattr("web_app.service.generate_answers", lambda *args, **kwargs: ["总体"] + ["分析", "改进"] * 2)
    grades = _grades(88, 76)
    settings = _settings()
    py_state = random.getstate()
    np_state = np.random.get_state()

    first = run_calculation(grades, None, settings)
    second = run_calculation(grades, None, dict(settings))
    exported = run_export(grades, None, dict(settings))
    _filename, report = run_ai_report(grades, None, dict(settings))

    assert random.getstate() == py_state
    assert np.array_equal(np.random.get_state()[1], np_state[1])
    left = _objectives(first["achievement"])
    assert left == _objectives(second["achievement"])
    assert left == _objectives(exported[2]["achievement"])
    table = _table5(report)
    for key, value in left.items():
        assert table[key] == value
    assert table["总达成度"] == round(float(first["achievement"]["总达成度"]), 3)

    other = run_calculation(_grades(55, 42), None, _settings())
    assert _objectives(other["achievement"]) != left


def test_same_seed_is_stable_and_another_seed_differs():
    from apply_noise import GradeReverseEngine, ScoreRng

    one = GradeReverseEngine(ScoreRng(20261005))
    two = GradeReverseEngine(ScoreRng(20261005))
    other = GradeReverseEngine(ScoreRng(7))
    assert one.dist_normal(80) == two.dist_normal(80)
    assert one.dist_bimodal(75) != other.dist_bimodal(75)
